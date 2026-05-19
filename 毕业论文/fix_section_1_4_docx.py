from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


BASE_DIR = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文")
DOC_PATH = BASE_DIR / "欧阳照林-毕业论文.docx"
FIG_PATH = BASE_DIR / "generated_figures" / "fig1_3_thesis_structure_academic.png"
BACKUP_PATH = BASE_DIR / "欧阳照林-毕业论文_备份_20260511_1.4修复前.docx"

SECTION_TITLE = "1.4　论文组织结构"
NEXT_CHAPTER_TITLE = "第2章　相关理论与技术基础"

INTRO = (
    "为便于把握全文研究内容的展开路径以及各章节之间的逻辑衔接关系，本文的整体组织结构与研究主线如图1-3所示。"
    "全文按照“问题提出与理论奠基—量化误差与分布统计分析—group-wise旋转算法设计与系统实现—实验验证与总结展望”的思路逐层推进，"
    "各章内容相互支撑、层层递进。"
)

CAPTION = "图1-3　本文整体组织结构与研究主线"

BODY_PARAGRAPHS = [
    "第1章为前言。本章首先从大语言模型低比特量化与 FP4/microscaling 部署需求出发，说明开展旋转优化研究的背景与意义；随后系统综述通用 PTQ 方法、旋转辅助量化方法以及 NVFP4/MXFP4 格式特化研究的国内外进展；在此基础上给出本文的主要研究内容、创新点与整体结构安排，为全文研究奠定问题背景。",
    "第2章介绍相关理论与技术基础。本章围绕后续方法设计所需的知识基础展开，系统梳理大语言模型尤其是 decoder-only Transformer 的基本结构、典型线性层功能差异、后训练量化的基本映射方式与误差来源，并进一步说明 NVFP4/MXFP4 微缩放量化格式及旋转辅助量化的等价变换原理，为后续分析与算法设计提供理论支撑。",
    "第3章围绕量化误差来源与权重/激活分布统计展开分析。本章面向 group-wise 场景，对低比特浮点量化中的典型误差模式进行归纳，结合权重与激活在不同局部 group 中的统计特征，分析影响量化精度的关键因素，并讨论这些统计规律对旋转策略设计所带来的启示，从而建立后续方法构建的分析依据。",
    "第4章给出面向低比特量化的 group-wise 旋转算法与系统实现。本章在前述误差分析和统计观察基础上，进一步介绍 group-wise 旋转搜索空间的构建、目标函数的设计与选择策略，以及量化导出、等价部署、推理接入和自动化评测链路的实现方式，形成从算法建模到工程落地的完整方法框架。",
    "第5章展示实验结果与分析。本章围绕所提出方法的有效性与适用性开展实验评估，从整体量化效果、不同目标函数对比、旋转策略与 group 粒度设置的消融分析、不同量化格式差异以及部署开销等多个角度展开讨论，对方法性能、优势与局限进行系统评估。",
    "第6章对全文工作进行总结与展望。本章对本文的研究内容、主要结论与工程实现价值进行归纳，进一步指出当前工作仍有待完善的问题，并结合大模型低比特量化与格式特化优化的发展趋势，对后续可能的研究方向进行展望。",
]


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def main() -> None:
    if not DOC_PATH.exists():
        raise FileNotFoundError(DOC_PATH)
    if not FIG_PATH.exists():
        raise FileNotFoundError(FIG_PATH)
    if not BACKUP_PATH.exists():
        copy2(DOC_PATH, BACKUP_PATH)

    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    section_idx = None
    next_chapter_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == SECTION_TITLE:
            section_idx = idx
        elif section_idx is not None and text == NEXT_CHAPTER_TITLE:
            next_chapter_idx = idx
            break

    if section_idx is None or next_chapter_idx is None or next_chapter_idx <= section_idx:
        raise RuntimeError("Failed to locate section 1.4 boundaries.")

    for idx in range(next_chapter_idx - 1, section_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[section_idx + 1]

    intro_para = anchor.insert_paragraph_before(INTRO)
    intro_para.style = doc.styles["Normal"]

    image_para = anchor.insert_paragraph_before("")
    image_para.style = doc.styles["Normal"]
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.paragraph_format.first_line_indent = Cm(0)
    image_para.paragraph_format.left_indent = Cm(0)
    image_para.paragraph_format.right_indent = Cm(0)
    image_run = image_para.add_run()
    image_run.add_picture(str(FIG_PATH), width=Cm(15.2))

    caption_para = anchor.insert_paragraph_before(CAPTION)
    caption_para.style = doc.styles["Caption"]
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for text in BODY_PARAGRAPHS:
        para = anchor.insert_paragraph_before(text)
        para.style = doc.styles["Normal"]

    doc.save(str(DOC_PATH))


if __name__ == "__main__":
    main()
