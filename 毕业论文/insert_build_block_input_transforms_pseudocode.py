from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
PLACEHOLDER = "【伪代码占位：展示 build_block_input_transforms 如何围绕 qkv、o、gate_up、down 四个槽位构造 group-wise 搜索任务，并将搜索结果组装为 MixedGroupTransform。】"


TITLE_LINE = "算法4-1 build_block_input_transforms的group-wise搜索任务构造"
NEXT_HEADING = "4.2.2 Identity 变换"


PSEUDOCODE_LINES = [
    "Input: block, args, quantizer settings, slot statistics",
    "Output: qkv_in_transform, o_in_transform, gate_up_in_transform, down_in_transform",
    "1  if transform_search = False then",
    "2      build one input transform for qkv, o, gate_up, and down, respectively",
    "3      return qkv_in_transform, o_in_transform, gate_up_in_transform, down_in_transform",
    "4  end if",
    "5  resolve search configuration and validate group-wise quantization constraints",
    "6  slot_tasks <- {qkv:{q_proj,k_proj,v_proj}, o:{o_proj}, gate_up:{gate_proj,up_proj}, down:{down_proj}}",
    "7  for each slot_name in {qkv, o, gate_up, down} do",
    "8      extract current-slot weights, and load covariance / activation statistics when required",
    "9      slot_transform[slot_name] <- search_best_group_transform(current_slot, candidates, group_size, objective, ...)",
    "10     // inside search_best_group_transform: choose the best candidate for each group",
    "11     // and stack the selected forward / backward matrices into MixedGroupTransform",
    "12 end for",
    "13 return slot_transform[qkv], slot_transform[o], slot_transform[gate_up], slot_transform[down]",
]


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def insert_paragraph_after(paragraph, text: str, style_name: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.add_run(text)
    return new_para


def remove_between_title_and_heading(doc):
    paragraphs = list(doc.paragraphs)
    title_para = None
    next_heading_para = None
    for paragraph in paragraphs:
        if paragraph.text.strip() == TITLE_LINE:
            title_para = paragraph
        elif title_para is not None and paragraph.text.strip() == NEXT_HEADING:
            next_heading_para = paragraph
            break
    if title_para is None or next_heading_para is None:
        raise RuntimeError("Unable to locate pseudocode range.")

    cursor = title_para._element.getnext()
    while cursor is not None and cursor is not next_heading_para._element:
        next_cursor = cursor.getnext()
        cursor.getparent().remove(cursor)
        cursor = next_cursor


def main():
    doc = Document(str(DOCX_PATH))
    target = None
    next_heading_para = None
    start_idx = None
    end_idx = None
    for idx, paragraph in enumerate(doc.paragraphs):
        if PLACEHOLDER in paragraph.text or TITLE_LINE in paragraph.text:
            target = paragraph
            start_idx = idx
        if NEXT_HEADING in paragraph.text:
            next_heading_para = paragraph
            end_idx = idx
        if target is not None and next_heading_para is not None:
            break

    if target is None:
        raise RuntimeError("Target pseudocode placeholder not found.")
    if next_heading_para is None:
        raise RuntimeError("Next heading not found.")
    if start_idx is None or end_idx is None:
        raise RuntimeError("Unable to resolve pseudocode paragraph range.")

    style = target.style
    clear_paragraph(target)
    target.style = style

    title_run = target.add_run(TITLE_LINE)
    title_run.bold = True

    remove_between_title_and_heading(doc)

    current = target
    for line in PSEUDOCODE_LINES:
        current = insert_paragraph_after(current, line, style.name if style is not None else "")

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
