from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
TEMP_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文_第三章局部更新待替换.docx")
CHAPTER_TITLE = "第3章 权重/激活分布统计与量化误差来源分析"
STOP_TITLE = "3.4 group 粒度下的量化误差来源分析"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    sub_node = OxmlElement("m:sub")
    append_expr(sub_node, sub)
    sup_node = OxmlElement("m:sup")
    append_expr(sup_node, sup)
    node.append(e)
    node.append(sub_node)
    node.append(sup_node)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_center_placeholder_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Normal"
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def main() -> None:
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    stop_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == CHAPTER_TITLE:
            start_idx = idx
        elif start_idx is not None and text == STOP_TITLE:
            stop_idx = idx
            break

    if start_idx is None or stop_idx is None or stop_idx <= start_idx:
        raise RuntimeError("Failed to locate Chapter 3 pre-3.4 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    stop_p = doc.paragraphs[stop_idx]._p
    remove_tables_between(start_p, stop_p)

    for idx in range(stop_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]

    insert_text_paragraph_before(
        anchor,
        "本章围绕权重与激活在低比特浮点量化场景下的局部分布特征及其误差来源展开分析。"
        "本文重点关注 group 列块这一局部尺度：一方面统计不同模型、不同层中 outlier 通道的分布情况以及普通通道的相对均匀性，"
        "说明旋转并非对所有列块都同样有益；另一方面比较全局旋转与 group 粒度分块旋转对局部统计结构的影响，"
        "论证在 NVFP4/MXFP4 场景下采用 group-wise 旋转、并避免单一固定旋转矩阵的必要性。"
        "在此基础上，后文将进一步结合量化误差统计与分位点贡献分析，讨论不同局部模式对应的误差来源。",
        "Normal",
    )

    insert_heading_before(anchor, "3.2 权重/激活数据分布统计", "Heading 2")

    insert_mixed_paragraph_before(
        anchor,
        [
            "在统计分析中，本文将通道维按照大小为 ",
            math_segment(["m"]),
            " 的列块划分为若干 group，并假设不同 token 在相同位置上的 group 共享同一局部变换。"
            "若单层通道维度记为 ",
            math_segment(["d"]),
            "，则 group 数量 ",
            math_segment(["G"]),
            " 可表示为式（3-1）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("A", "l"),
            " = [",
            m_sub("A", "l,1"),
            ", ",
            m_sub("A", "l,2"),
            ", …, ",
            m_sub("A", "l,G"),
            "],   G = ",
            m_frac("d", "m"),
        ],
        "（3-1）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "多模型、多层的统计结果表明，outlier 通道并非局部或偶然现象，而是在大语言模型中广泛存在。"
            "无论是注意力模块中的投影层，还是前馈模块中的 ",
            math_segment(["up_proj"]),
            "、",
            math_segment(["down_proj"]),
            " 等线性层，都能够观察到少量通道在幅值上显著高于其余通道。"
            "为了刻画这种局部峰值主导现象，本文对每个 group 统计最大幅值和归一化幅值分布，其定义如式（3-2）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("M", "g"),
            " = ",
            m_sub("max", "1≤i≤m"),
            "|",
            m_sub("x", "g,i"),
            "|,   ",
            m_sub("u", "g,i"),
            " = ",
            m_frac(["|", m_sub("x", "g,i"), "|"], m_sub("M", "g")),
        ],
        "（3-2）",
        doc,
    )

    insert_text_paragraph_before(
        anchor,
        "与此同时，统计结果也显示，大多数通道和大多数局部列块的分布其实相对平滑。"
        "也就是说，outlier 的存在虽然为旋转提供了潜在收益空间，但这种收益并不是均匀分布在所有 group 上的。"
        "对于本身已经较为均匀的列块，若盲目施加旋转，反而可能打破原有的局部平衡，使峰值集中度上升并恶化共享尺度匹配。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "这一现象在 NVFP4/MXFP4 场景下尤为重要。由于这类格式以 group 为基本共享尺度单元，"
        "量化精度直接受局部列块内部统计形态影响，因此是否旋转、以及在何种粒度上施加旋转，都应建立在 group 级分布统计的基础上。"
        "从这一意义上说，outlier 通道的广泛存在说明旋转具有必要性，而大多数列块的相对均匀性以及 microscaling 的格式约束，则进一步说明旋转应当在 group 粒度上有选择地进行。",
        "Normal",
    )

    insert_caption_before(anchor, "表3-1 分布统计观测对象与设置汇总（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入参与统计的模型、层类型、观测对象、采样 token 数与 group size 设置汇总表】")
    insert_center_placeholder_before(anchor, "【插图占位：此处插入不同模型不同层中 outlier 通道广泛存在、同时多数通道相对均匀的联合统计图】")
    insert_caption_before(anchor, "图3-1 不同模型不同层中 outlier 通道与均匀区间的联合统计结果（占位）")
    insert_center_placeholder_before(anchor, "【插图占位：此处插入典型层中局部 group 分布形态示意图，用于说明均匀列块与异常列块的差异】")
    insert_caption_before(anchor, "图3-2 典型层中局部 group 分布形态示意（占位）")

    insert_heading_before(anchor, "3.3 旋转粒度对局部统计分布的影响", "Heading 2")
    insert_heading_before(anchor, "3.3.1 全局旋转与 group 粒度分块旋转的对比", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "为了比较不同旋转粒度对 outlier 扩散程度的影响，本文采用层级大值比例与 block 污染比例两个统计量。"
            "前者描述某层中超过给定阈值的元素占比，后者描述受较大幅值影响的 group 比例，分别定义为式（3-3）与式（3-4）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("ρ", "l"),
            "(",
            "τ",
            ") = ",
            m_frac("1", m_sub("N", "l")),
            m_subsup("Σ", "j=1", m_sub("N", "l")),
            "I(",
            "|",
            m_sub("x", "l,j"),
            "| ≥ τ",
            m_sub("M", "l"),
            ")",
        ],
        "（3-3）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("η", "l"),
            "(",
            "τ",
            ") = ",
            m_frac("1", m_sub("G", "l")),
            m_subsup("Σ", "g=1", m_sub("G", "l")),
            "I(",
            m_sub("M", "g"),
            " ≥ τ",
            m_sub("M", "l"),
            ")",
        ],
        "（3-4）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若采用全局统一旋转，则同一层所有通道共同参与同一个变换；若采用 group 粒度分块旋转，则不同列块只在各自局部范围内完成线性混合。"
            "两种作用方式可统一写成式（3-5）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sup(m_sub("A", "l"), "global"),
            " = ",
            m_sub("A", "l"),
            "R,   ",
            m_sup(m_sub("A", "l"), "group"),
            " = ",
            m_sub("A", "l"),
            "blkdiag(",
            m_sub("R", "1"),
            ", ",
            m_sub("R", "2"),
            ", …, ",
            m_sub("R", "G"),
            ")",
        ],
        "（3-5）",
        doc,
    )

    insert_text_paragraph_before(
        anchor,
        "二者差异的关键在于能量传播范围不同。全局旋转会把原本集中在少数通道上的大值信息投影到更大范围的坐标方向上，"
        "在 tensor-wise 或 channel-wise 场景中，这种扩散有时有助于削弱单一坐标主导效应；"
        "但在 NVFP4/MXFP4 这类 group-wise microscaling 场景中，它往往会同步抬高层级大值比例和 block 污染比例，"
        "将原本局限在少量 block 中的异常能量扩散到大量普通 block 之中。"
        "相比之下，group 粒度分块旋转将线性混合限制在局部列块内部，跨 group 的能量传播被显式阻断，"
        "因而更有利于在缓解局部 outlier 的同时保留原有的尺度隔离结构。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "从部署含义上看，group 粒度分块旋转还允许同一位置上的列块在不同 token 间共享同一变换，"
        "从而兼顾统计针对性和实现一致性。"
        "因此，对于依赖局部共享尺度的 NVFP4/MXFP4 格式而言，真正需要比较的并不是“是否做旋转”，"
        "而是“是否应当将旋转范围收缩到 group 列块内部”。",
        "Normal",
    )

    insert_center_placeholder_before(anchor, "【插图占位：此处插入全局旋转与 group 粒度分块旋转前后大值比例和 block 污染比例的对比图】")
    insert_caption_before(anchor, "图3-3 全局旋转与 group 粒度分块旋转对 outlier 扩散影响的对比（占位）")
    insert_caption_before(anchor, "表3-2 全局旋转与 group 粒度分块旋转比较（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入不同旋转粒度下的大值比例、block 污染比例与误差变化比较表】")

    insert_heading_before(anchor, "3.3.2 单一固定旋转矩阵的局限性", "Heading 3")

    insert_text_paragraph_before(
        anchor,
        "即便将旋转粒度限制在局部列块内部，单一固定旋转矩阵仍然难以直接覆盖所有 group 和所有量化格式。"
        "一方面，不同模型、不同层中 outlier 的强度和局部分布形态差异明显，同一固定变换未必能够同时兼顾所有列块；"
        "另一方面，对于原本已经较为平滑的 group，固定模式混合可能反而引入新的不均匀性。"
        "尤其是在 NVFP4 这类对局部共享尺度极为敏感的格式中，分块 Hadamard 等固定结构旋转并不总是带来正收益，"
        "其副作用恰恰说明旋转矩阵需要针对局部统计状态和量化格式特性进行定制。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "因此，问题的关键并不只是“是否选择分块旋转”，还在于“对什么样的 group 采用什么样的变换”。"
        "对于 outlier 挤压显著的列块，更重要的是削弱峰值主导；对于本身分布均衡的列块，则更应避免无谓扰动；"
        "而对于不同量化格式，还需要考虑其尺度表示方式和局部数值密度特征的差异。"
        "为了在后续误差分析中识别不同数值区间的贡献，本文采用式（3-6）定义的分位点误差贡献比例。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("c", "g,k"),
            " = ",
            m_frac(
                [
                    m_sub("Σ", [m_sub("u", "g,i"), "∈", m_sub("B", "k")]),
                    "|",
                    m_sub("e", "g,i"),
                    "|",
                ],
                [
                    m_subsup("Σ", "i=1", "m"),
                    "|",
                    m_sub("e", "g,i"),
                    "|",
                ],
            ),
            ",   ",
            m_sub("e", "g,i"),
            " = ",
            m_sub("x̂", "g,i"),
            " - ",
            m_sub("x", "g,i"),
        ],
        "（3-6）",
        doc,
    )

    insert_center_placeholder_before(anchor, "【插图占位：此处插入单一固定旋转矩阵在 NVFP4/MXFP4 场景下可能引入副作用的示意图】")
    insert_caption_before(anchor, "图3-4 单一固定旋转矩阵在局部列块上的潜在副作用（占位）")

    try:
        doc.save(str(DOC_PATH))
        print(f"SAVED:{DOC_PATH}")
    except PermissionError:
        doc.save(str(TEMP_PATH))
        print(f"TEMP_SAVED:{TEMP_PATH}")


if __name__ == "__main__":
    main()
