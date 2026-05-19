from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
TEMP_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文_2.3局部更新待替换.docx")
START_TITLE = "2.3.2　旋转改善量化误差的作用机理"
NEXT_SECTION_TITLE = "2.5　本章小结"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_center_placeholder_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Normal"
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def insert_table_before(anchor, doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    anchor._p.addprevious(table._tbl)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def rebuild_subsections():
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == START_TITLE:
            start_idx = idx
        elif start_idx is not None and text == NEXT_SECTION_TITLE:
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate subsection 2.3.2 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx - 1, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx]

    insert_heading_before(anchor, "2.3.2　旋转改善量化误差的作用机理", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "在 ",
            math_segment(["NVFP4/MXFP4"]),
            " 一类以 group 为单位共享尺度的低比特量化中，量化误差并不只由码字位宽决定，"
            "还与同一 group 内数据的动态范围是否均衡密切相关。若某一局部列块中存在幅值显著高于其他元素的 "
            "outlier，则共享尺度往往会被该极大值主导，进而使大量普通元素只能落在较粗的量化步长上。"
            "因此，分析旋转改善误差的机理，关键在于说明 ",
            math_segment([m_sub("s", "g")]),
            " 如何被 outlier 放大，以及旋转如何通过重分配能量降低这一放大量。"
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "记第 ",
            math_segment(["g"]),
            " 个 group 中的输入或权重向量为 ",
            math_segment([m_sub("x", "g"), " = [", m_sub("x", "g,1"), ", …, ", m_sub("x", "g,m"), "]"]),
            "。若采用按 group 共享尺度的对称量化，则量化与反量化过程可写为式（2-29），"
            "对应的共享尺度估计可写为式（2-30）。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("q", "g,i"),
            " = clip(round(",
            m_frac(m_sub("x", "g,i"), m_sub("s", "g")),
            "), -",
            m_sub("q", "max"),
            ", ",
            m_sub("q", "max"),
            "),   ",
            m_sub("x̂", "g,i"),
            " = ",
            m_sub("s", "g"),
            m_sub("q", "g,i"),
        ],
        "（2-29）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("s", "g"),
            " = ",
            m_frac([m_sub("max", "1≤i≤m"), "|", m_sub("x", "g,i"), "|"], m_sub("q", "max")),
        ],
        "（2-30）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "在忽略截断并采用就近舍入的近似条件下，式（2-29）对应的单元素绝对误差上界可表示为式（2-31）。"
            "这说明，只要 group 内存在一个很大的峰值，所有元素的误差上界都会被同步抬高。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            "|",
            m_sub("e", "g,i"),
            "| = |",
            m_sub("x̂", "g,i"),
            " - ",
            m_sub("x", "g,i"),
            "| ≤ ",
            m_frac(m_sub("s", "g"), "2"),
            " = ",
            m_frac([m_sub("max", "1≤j≤m"), "|", m_sub("x", "g,j"), "|"], ["2", m_sub("q", "max")]),
        ],
        "（2-31）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "进一步地，若将 group 内的最大幅值记为 ",
            math_segment([m_sub("M", "g"), " = ", m_sub("max", "1≤j≤m"), "|", m_sub("x", "g,j"), "|"]),
            "，则对于任意非 outlier 元素，其相对误差上界可写为式（2-32）。"
            "可以看到，当普通元素幅值远小于 ",
            math_segment([m_sub("M", "g")]),
            " 时，相对误差会迅速增大。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_frac(["|", m_sub("e", "g,i"), "|"], ["|", m_sub("x", "g,i"), "|"]),
            " ≤ ",
            m_frac(m_sub("M", "g"), ["2", m_sub("q", "max"), "|", m_sub("x", "g,i"), "|"]),
        ],
        "（2-32）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若从有效分辨率的角度观察，同一普通元素在其局部幅值区间内可利用的离散级数可近似表示为式（2-33）。"
            "因此，outlier 的存在本质上会压缩普通元素可用的码字分辨能力，这也是低比特条件下精度明显退化的直接原因。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("L", "g,i"),
            " ≈ ",
            m_frac(["2|", m_sub("x", "g,i"), "|"], m_sub("s", "g")),
            " = ",
            m_frac(["2", m_sub("q", "max"), "|", m_sub("x", "g,i"), "|"], m_sub("M", "g")),
        ],
        "（2-33）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "旋转之所以能够缓解这类误差，并不是因为它改变了模型的线性映射，而是因为它在保持等价计算的前提下改变了坐标系中的能量分布。"
            "记旋转后的 group 表示为 ",
            math_segment([m_sub("x′", "g"), " = ", m_sub("x", "g"), "R"]),
            "，并记旋转后 group 内最大幅值为 ",
            math_segment([m_sub("M′", "g")]),
            "。若所选旋转能够削弱单坐标主导效应，则有式（2-34）。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("x′", "g"),
            " = ",
            m_sub("x", "g"),
            "R,   ",
            m_sub("M′", "g"),
            " = ",
            m_sub("max", "1≤i≤m"),
            "|",
            m_sub("x′", "g,i"),
            "| < ",
            m_sub("M", "g"),
        ],
        "（2-34）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "此时，共享尺度会同步收缩为式（2-35），相应的量化误差上界也随之下降。"
            "换言之，只要旋转后峰值主导程度减弱，低比特量化器便能将更多码字分辨率分配给原先被压缩的普通元素。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("s′", "g"),
            " = ",
            m_frac(m_sub("M′", "g"), m_sub("q", "max")),
            ",   |",
            m_sub("e′", "g,i"),
            "| ≤ ",
            m_frac(m_sub("s′", "g"), "2"),
            " = ",
            m_frac(m_sub("M′", "g"), ["2", m_sub("q", "max")]),
        ],
        "（2-35）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "对于正交旋转，向量的 ",
            math_segment(["l₂"]),
            " 范数保持不变，因此旋转所做的只是重新分配各坐标上的能量，而不是凭空改变总体能量。"
            "这一性质可表示为式（2-36）。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            "‖",
            m_sub("x′", "g"),
            "‖₂ = ‖",
            m_sub("x", "g"),
            "R‖₂ = ‖",
            m_sub("x", "g"),
            "‖₂",
        ],
        "（2-36）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "又由于任意 ",
            math_segment(["m"]),
            " 维向量都满足 ",
            math_segment(["l∞"]),
            " 范数与 ",
            math_segment(["l₂"]),
            " 范数之间的基本不等式，故旋转后最大坐标幅值的理论下界可写为式（2-37）。"
            "这意味着，理想旋转并不是无限制地压低峰值，而是尽可能在保持局部结构可部署的前提下，使能量分布向更均匀的状态靠近。"
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("M′", "g"),
            " = ",
            m_sub("max", "1≤i≤m"),
            "|",
            m_sub("x′", "g,i"),
            "| ≥ ",
            m_frac(["‖", m_sub("x", "g"), "‖₂"], "√m"),
        ],
        "（2-37）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "综合式（2-31）至式（2-37）可知，旋转改善量化误差的本质在于：通过降低 group 内的峰值集中度，缩小共享尺度，"
            "从而同时减小绝对误差上界并提高普通元素的有效分辨率。"
            "不过，这一收益并非无条件成立。若旋转虽然降低了局部最大值，却破坏了原有的 group 结构，使高幅值信息跨 group 扩散，"
            "则也可能削弱 microscaling 对局部尺度隔离的优势。因而，面向 NVFP4/MXFP4 场景的旋转设计，本质上是在 "
            "“峰值抑制能力” 与 “局部结构保持能力” 之间寻求平衡。"
        ],
        "Normal",
    )

    insert_heading_before(anchor, "2.3.3　常见旋转量化方法及适用场景", "Heading 3")

    insert_text_paragraph_before(
        anchor,
        "从方法来源和实现方式看，旋转辅助量化并非单一算法，而是一类围绕候选变换构造、误差适配方式与部署开销展开设计的方法集合。"
        "对于大语言模型低比特量化而言，常见方法大致可分为不旋转基线、固定结构旋转、训练自由的固定旋转算法，以及带有可学习或结构增强特征的旋转算法几类。",
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "其中，",
            math_segment(["Identity"]),
            " 可视为不施加额外旋转的对照基线。该策略不引入附加变换与额外部署开销，适合用于判断某一层或某一 group 是否确有旋转需求；"
            "但它无法主动缓解 outlier 主导、重尾分布或局部尺度失配等问题，因此通常只作为比较参考，而不是提升量化精度的主要手段。"
        ],
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "Hadamard 类变换属于典型的固定结构正交旋转。该类方法具有结构规则、计算开销低、硬件实现友好等优点，常被用作训练自由的基础候选。"
        "由于其不依赖额外参数学习，适合大规模模型的快速部署，也适合作为后续更复杂旋转策略的初始候选或对照对象。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入 Hadamard 类变换示意图】")
    insert_caption_before(anchor, "图2-3　Hadamard 类变换示意（占位）")

    insert_text_paragraph_before(
        anchor,
        "QuaRot[5] 可以看作固定旋转路线中的代表性方法。其核心思想是在训练自由条件下利用结构化正交变换削弱隐藏状态中的异常峰值，"
        "并将旋转扩展到权重、激活乃至 KV cache 等多个环节，从而支持更激进的低比特量化。"
        "该方法适合强调端到端低比特部署、同时希望避免额外训练与复杂搜索开销的场景；其局限在于旋转形式较为固定，对不同层和不同局部区域的适配能力相对有限。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入 QuaRot 方法示意图】")
    insert_caption_before(anchor, "图2-4　QuaRot 方法示意（占位）")

    insert_text_paragraph_before(
        anchor,
        "SpinQuant[6] 则代表了可学习旋转的思路。该方法借助少量校准数据对旋转矩阵进行优化，使旋转结果能够更贴合模型实际分布与量化目标，"
        "因而在难量化模型、极低位宽设置或对精度较为敏感的场景下通常具有更好的效果。与固定旋转相比，SpinQuant 的优势在于适配性更强；"
        "但其代价是需要额外的离线优化流程，部署前准备成本也相对更高。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入 SpinQuant 方法示意图】")
    insert_caption_before(anchor, "图2-5　SpinQuant 方法示意（占位）")

    insert_text_paragraph_before(
        anchor,
        "GSR[7] 体现了结构增强型旋转的设计思路。该方法在 Walsh-Hadamard 变换基础上进一步引入更细粒度的结构组织，使 outlier 的影响能够在局部范围内得到更有针对性的隔离与扩散，"
        "从而改善极低位宽下的量化稳定性。相较于标准固定旋转，GSR 在局部适配能力上更进一步；但其收益通常依赖于具体的分组机制和结构设计，分析与实现复杂度也相应更高。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入 GSR 方法示意图】")
    insert_caption_before(anchor, "图2-6　GSR 方法示意（占位）")

    insert_text_paragraph_before(
        anchor,
        "为便于比较，表2-3 对上述几类常见旋转量化方法的主要性质、适用场景与典型局限进行了归纳。",
        "Normal",
    )
    insert_caption_before(anchor, "表2-3　常见旋转量化方法的性质比较")
    insert_table_before(
        anchor,
        doc,
        [
            ["方法类别", "代表方法", "主要性质", "适用场景", "典型局限"],
            ["不旋转基线", "Identity", "无额外变换开销，适合作为对照基线", "用于判断旋转是否必要，或作为消融比较", "无法主动缓解 outlier 与共享尺度失配"],
            ["固定结构旋转", "Hadamard 类变换", "结构规则、实现简单、计算开销低", "强调快速部署与大规模覆盖的场景", "适配性有限，难以针对局部差异细化调整"],
            ["固定旋转算法", "QuaRot", "训练自由，可兼顾权重、激活与 KV cache 的低比特量化", "端到端低比特部署、关注整体实现成本的场景", "旋转形式较固定，对层间和局部异质性的利用有限"],
            ["可学习旋转", "SpinQuant", "旋转矩阵可根据校准数据优化，分布适配能力更强", "极低位宽、难量化模型、精度敏感场景", "需要额外离线优化，准备成本较高"],
            ["结构增强旋转", "GSR", "在固定结构基础上强化局部适配能力，改善极低位宽稳定性", "希望兼顾较低部署开销与更强局部误差控制的场景", "结构设计与收益依赖具体分组机制，实现分析更复杂"],
        ],
    )

    insert_text_paragraph_before(
        anchor,
        "总体而言，现有旋转量化方法之间并不存在对所有模型、所有层和所有量化格式都普遍最优的单一方案。"
        "固定结构方法的优势在于成本低、易部署，可学习方法的优势在于适配性强，结构增强方法则试图在二者之间取得折中。"
        "因此，更合理的研究路径不是预设某一种旋转必然优于其他候选，而是结合量化格式、层级角色、局部统计特征与部署约束，对候选旋转进行有针对性的选择与设计。",
        "Normal",
    )

    try:
        doc.save(str(DOC_PATH))
        print(f"SAVED:{DOC_PATH}")
    except PermissionError:
        doc.save(str(TEMP_PATH))
        print(f"TEMP_SAVED:{TEMP_PATH}")


if __name__ == "__main__":
    rebuild_subsections()
