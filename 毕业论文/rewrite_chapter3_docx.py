from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
TEMP_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文_第三章更新待替换.docx")
CHAPTER_TITLE = "第3章 权重/激活分布统计与量化误差来源分析"
NEXT_CHAPTER_TITLE = "第4章 面向低比特量化的 group-wise 旋转算法与系统实现"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def m_subsup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    sub_node = OxmlElement("m:sub")
    append_expr(sub_node, sub)
    sup_node = OxmlElement("m:sup")
    append_expr(sup_node, sup)
    node.append(e)
    node.append(sub_node)
    node.append(sup_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_center_placeholder_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Normal"
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def build_chapter_3():
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == CHAPTER_TITLE:
            start_idx = idx
        elif start_idx is not None and text == NEXT_CHAPTER_TITLE:
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate Chapter 3 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]

    insert_text_paragraph_before(
        anchor,
        "本章围绕权重与激活在低比特浮点量化场景下的统计分布特征及其误差来源展开分析。"
        "本文将关注重点放在 group 粒度的局部列块上，一方面统计不同模型、不同层中 outlier 通道的分布情况以及普通通道的相对均匀性，"
        "另一方面结合 group 内量化误差、最大值统计和分位点误差贡献，归纳典型误差模式，并进一步讨论不同旋转粒度对局部统计结构的影响。"
        "通过这些分析，本文旨在论证面向 NVFP4/MXFP4 场景采用 group-wise 分块旋转、并避免使用单一固定旋转矩阵的必要性，"
        "为下一章算法设计提供直接依据。",
        "Normal",
    )

    insert_heading_before(anchor, "3.1 分析对象与统计指标", "Heading 2")

    insert_mixed_paragraph_before(
        anchor,
        [
            "为了统一描述后续统计过程，本文将典型线性层的激活矩阵记为 ",
            math_segment([m_sub("A", "l")]),
            "，权重矩阵记为 ",
            math_segment([m_sub("W", "l")]),
            "。其中，激活沿通道维按大小为 ",
            math_segment(["m"]),
            " 的列块划分为若干 group，同一位置上的 group 在不同 token 间共享相同的局部变换。"
            "若记单层通道维度为 ",
            math_segment(["d"]),
            "，则 group 数量为 ",
            math_segment(["G = d/m"]),
            "，可写成式（3-1）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("A", "l"),
            " = [",
            m_sub("A", "l,1"),
            ", ",
            m_sub("A", "l,2"),
            ", …, ",
            m_sub("A", "l,G"),
            "],   ",
            m_sub("W", "l"),
            " = [",
            m_sub("W", "l,1"),
            ", ",
            m_sub("W", "l,2"),
            ", …, ",
            m_sub("W", "l,G"),
            "],   G = ",
            m_frac("d", "m"),
        ],
        "（3-1）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "对任一 group 内的一维向量样本 ",
            math_segment([m_sub("x", "g"), " = [", m_sub("x", "g,1"), ", …, ", m_sub("x", "g,m"), "]"]),
            "，本文首先统计其最大幅值与归一化位置，以刻画局部峰值集中程度及组内相对分布形态，相关定义如式（3-2）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("M", "g"),
            " = ",
            m_sub("max", "1≤i≤m"),
            "|",
            m_sub("x", "g,i"),
            "|,   ",
            m_sub("u", "g,i"),
            " = ",
            m_frac(["|", m_sub("x", "g,i"), "|"], m_sub("M", "g")),
        ],
        "（3-2）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "进一步地，若将量化与反量化后的 group 记为 ",
            math_segment([m_sub("x̂", "g")]),
            "，则本文以 group 级重建误差 ",
            math_segment([m_sub("E", "g")]),
            " 作为局部量化损失的基本度量，其表达式见式（3-3）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("E", "g"),
            " = ",
            m_frac("1", "m"),
            m_subsup("Σ", "i=1", "m"),
            "|",
            m_sub("x̂", "g,i"),
            " - ",
            m_sub("x", "g,i"),
            "|",
            m_sup("", "2"),
        ],
        "（3-3）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "为了比较不同旋转策略对大值扩散的影响，本文还定义层级大值比例与 block 污染比例。"
            "前者描述某层中超过阈值的元素占比，后者描述受较大幅值影响的 group 比例，分别如式（3-4）与式（3-5）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("ρ", "l"),
            "(",
            "τ",
            ") = ",
            m_frac("1", m_sub("N", "l")),
            m_subsup("Σ", "j=1", m_sub("N", "l")),
            "I(",
            "|",
            m_sub("x", "l,j"),
            "| ≥ τ",
            m_sub("M", "l"),
            ")",
        ],
        "（3-4）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("η", "l"),
            "(",
            "τ",
            ") = ",
            m_frac("1", m_sub("G", "l")),
            m_subsup("Σ", "g=1", m_sub("G", "l")),
            "I(",
            m_sub("M", "g"),
            " ≥ τ",
            m_sub("M", "l"),
            ")",
        ],
        "（3-5）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "最后，为了分析组内不同幅值区间对量化误差的贡献，本文将归一化后的 ",
            math_segment([m_sub("u", "g,i")]),
            " 按区间划分为若干分位点 bin ",
            math_segment([m_sub("B", "k")]),
            "，并定义对应误差贡献比例，如式（3-6）所示。"
            "该指标可用于判断主要误差究竟来自接近零的小值区域，还是来自靠近峰值的大值区域。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("c", "g,k"),
            " = ",
            m_frac(
                [
                    m_sub("Σ", [m_sub("u", "g,i"), "∈", m_sub("B", "k")]),
                    "|",
                    m_sub("e", "g,i"),
                    "|",
                ],
                [
                    m_subsup("Σ", "i=1", "m"),
                    "|",
                    m_sub("e", "g,i"),
                    "|",
                ],
            ),
            ",   ",
            m_sub("e", "g,i"),
            " = ",
            m_sub("x̂", "g,i"),
            " - ",
            m_sub("x", "g,i"),
        ],
        "（3-6）",
        doc,
    )

    insert_text_paragraph_before(
        anchor,
        "上述指标分别对应局部峰值幅度、重建误差、大值扩散范围与误差贡献区间四个观察视角，"
        "后续分布统计与误差分析均围绕这些量展开。",
        "Normal",
    )
    insert_caption_before(anchor, "表3-1 统计对象与设置汇总（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入参与统计的模型、层类型、观测对象、采样 token 数与 group size 设置汇总表】")

    insert_heading_before(anchor, "3.2 权重/激活数据分布统计", "Heading 2")

    insert_heading_before(anchor, "3.2.1 outlier 通道的普遍性", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "多模型、多层的统计结果表明，outlier 通道并非局部或偶然现象，而是在大语言模型中广泛存在。"
        "无论是注意力模块中的投影层，还是前馈模块中的 up_proj、down_proj 等线性层，都可以观察到少量通道在幅值上显著高于其余通道。"
        "不同模型之间的差异主要体现在 outlier 的强度、持续层数以及其在不同层类型中的分布位置，而不是其是否存在。"
        "这说明，在低比特量化场景下，仅依赖“整体分布较稳定”的平均判断是不够的，局部极值通道对量化尺度的主导作用必须被单独考察。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "从方法设计角度看，这一观察意味着旋转变换的主要收益来源之一，确实可能来自对少数异常方向的重新分配。"
        "但与此同时，outlier 的层间分布并不一致，不同模型和不同层的局部统计异质性较强，"
        "因此后续旋转策略不宜仅依据全局经验进行统一处理，而需要建立在更细粒度的统计分析之上。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入不同模型不同层中 outlier 通道分布统计图】")
    insert_caption_before(anchor, "图3-1 不同模型不同层中 outlier 通道分布统计（占位）")

    insert_heading_before(anchor, "3.2.2 大多数通道的相对均匀性与选择性旋转动机", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "尽管少数 outlier 通道较为突出，但对大多数通道而言，其局部分布往往相对平滑，组内幅值差异并不剧烈。"
        "这意味着，并非所有列块都天然适合施加旋转。对于本身已经较为均匀的 group，若盲目施加变换，"
        "反而可能破坏原有的局部平衡，使得峰值集中度上升、共享尺度变差，进而引入额外量化误差。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "因此，旋转是否有益并不是一个对所有 group 都成立的统一命题，而更像一个依赖局部统计状态的选择问题。"
        "当某些列块存在显著峰值主导效应时，旋转可能有效缓解 outlier 挤压；而当列块内部原本分布均衡时，保持其结构不变往往更加稳妥。"
        "这一点直接构成了后续“选择性旋转”与“按 group 决策”两项设计动机。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入典型层中多数通道相对均匀、少数通道异常突出的分布示意图】")
    insert_caption_before(anchor, "图3-2 典型层中通道分布的均匀区间与异常峰值示意（占位）")

    insert_heading_before(anchor, "3.3 旋转粒度对局部统计分布的影响", "Heading 2")

    insert_heading_before(anchor, "3.3.1 全层旋转的 outlier 扩散效应", "Heading 3")
    insert_mixed_paragraph_before(
        anchor,
        [
            "若采用全层统一旋转，则同一层所有通道将共同参与同一个变换，其作用方式可抽象为式（3-7）。",
            "在这种情形下，原本集中在少数通道上的大值信息会被投影到更大范围的坐标方向上，"
            "从而有可能将 outlier 能量扩散到大量原本较为普通的 group 中。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sup(m_sub("A", "l"), "global"),
            " = ",
            m_sub("A", "l"),
            "R,   ",
            m_sup(m_sub("W", "l"), "global"),
            " = ",
            m_sub("W", "l"),
            "R",
        ],
        "（3-7）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "对于 tensor-wise 或 channel-wise 量化而言，这种扩散有时有助于削弱单一坐标上的极值主导效应。"
        "但在 NVFP4/MXFP4 这类 group-wise microscaling 场景中，全层旋转会直接改变大量局部列块的内部统计，"
        "导致层级大值比例和 block 污染比例同时上升。换言之，原本只集中在少量 block 中的异常能量，可能在旋转后渗入更多普通 block，"
        "使后续共享尺度的匹配质量整体恶化。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入全层旋转前后大值比例与 block 污染程度对比图】")
    insert_caption_before(anchor, "图3-3 全层旋转导致 outlier 能量扩散的示意（占位）")

    insert_heading_before(anchor, "3.3.2 group 粒度分块旋转的局部性优势", "Heading 3")
    insert_mixed_paragraph_before(
        anchor,
        [
            "与全层统一旋转相比，group 粒度分块旋转只在局部列块内部进行线性混合。"
            "若记不同 group 上的局部变换分别为 ",
            math_segment([m_sub("R", "1"), ", ", m_sub("R", "2"), ", …, ", m_sub("R", "G")]),
            "，则其作用方式可写为式（3-8）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sup(m_sub("A", "l"), "group"),
            " = ",
            m_sub("A", "l"),
            "blkdiag(",
            m_sub("R", "1"),
            ", ",
            m_sub("R", "2"),
            ", …, ",
            m_sub("R", "G"),
            "),   ",
            m_sup(m_sub("W", "l"), "group"),
            " = ",
            m_sub("W", "l"),
            "blkdiag(",
            m_sub("R", "1"),
            ", ",
            m_sub("R", "2"),
            ", …, ",
            m_sub("R", "G"),
            ")",
        ],
        "（3-8）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "这种设计的关键优势在于局部性：一方面，同一位置上的 group 在不同 token 间共享同一变换，因此能够保持部署形式上的一致；"
        "另一方面，跨 group 的能量传播被显式阻断，异常大值只能在局部列块内被重分配，而不会扩散到整层的大量普通 block 中。"
        "在此基础上，后续算法还可以进一步只对统计上“确有必要”的 group 施加旋转，而对原本均衡的 group 保持不变，从而兼顾收益与风险。",
        "Normal",
    )
    insert_caption_before(anchor, "表3-2 全层旋转与 group 粒度分块旋转比较（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入不同旋转粒度下的大值比例、block 污染比例与误差变化比较表】")

    insert_heading_before(anchor, "3.3.3 单一固定旋转矩阵的局限性", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "已有研究与本文观察共同表明，单一固定旋转矩阵并不适合直接覆盖所有 group 和所有量化格式。"
        "一方面，全局旋转在 NVFP4/MXFP4 场景下容易把 outlier 能量扩散到大量普通 block 中，造成误差放大；"
        "另一方面，即便采用分块 Hadamard 这类结构化固定旋转，在 NVFP4 等格式上仍可能出现副作用，"
        "因为某些原本较平滑的局部分布会在固定模式混合后变得更不均匀，从而破坏共享尺度的适配效果。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "这说明，问题的关键并不只是“是否旋转”，而在于“对什么样的 group 采用什么样的变换”。"
        "对于 outlier 挤压显著的列块，需要更强调峰值缓解；对于原本已经均衡的列块，则更应避免无谓扰动；"
        "而对不同量化格式而言，旋转矩阵还应与其尺度表示方式和局部数值密度特征相匹配。"
        "因此，后续方法不能停留在单一固定矩阵层面，而应走向 group-wise、格式特化的旋转设计路线。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入固定旋转矩阵在 NVFP4/MXFP4 场景下副作用示意图】")
    insert_caption_before(anchor, "图3-4 单一固定旋转矩阵在局部列块上的潜在副作用（占位）")

    insert_heading_before(anchor, "3.4 group 粒度下的量化误差来源分析", "Heading 2")

    insert_heading_before(anchor, "3.4.1 组内量化误差与最大值统计", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "在 group-wise 场景下，最直接的观察方式是同时统计每个 group 的量化误差与其最大值幅度。"
        "若经典的 outlier 挤压机制占主导，则最大值相对于普通值越大，组内重建误差通常也越大。"
        "这一现象在部分层上表现得较为明显，说明异常峰值确实会通过共享尺度放大普通值的误差；"
        "但也有一些层并不严格服从这一单调关系，表明组内误差来源并不只有单一的 outlier 挤压机制。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "因此，仅从“最大值是否很大”来判断旋转需求仍然不够。"
        "某些 group 尽管峰值突出，但主要误差确实来自小值被挤压；另一些 group 的主要误差则可能来自多个较大值之间的共同作用。"
        "这也是本文进一步引入分位点贡献分析的原因。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入每个 group 的量化误差与最大值统计对比图】")
    insert_caption_before(anchor, "图3-5 组内量化误差与最大值统计关系（占位）")

    insert_heading_before(anchor, "3.4.2 分位点视角下的误差贡献分解", "Heading 3")
    insert_mixed_paragraph_before(
        anchor,
        [
            "为了进一步识别误差真正来自哪些数值区间，本文对误差较大的 group 进行细化分析。"
            "具体做法是先将组内数据按最大值归一化到 ",
            "区间 ",
            math_segment(["[0, 1]"]),
            "，再依据式（3-6）统计不同分位点 bin 对总误差的贡献比例。"
            "若低分位区间的误差贡献显著偏高，则说明主要问题来自接近零的小值被压缩；"
            "若高分位区间贡献更大，则说明误差更多来自大值附近的表示稀疏或局部尺度不匹配。",
        ],
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "这一分析视角比仅观察最大值更细致，因为它能够区分“峰值存在”与“峰值是否真正决定主要误差来源”之间的差别。"
        "从结果解释上看，分位点贡献分解为后续旋转设计提供了更直接的线索：不同的误差分布模式，对应不同的变换目标与优先级。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入高误差 group 的分位点误差贡献分解图】")
    insert_caption_before(anchor, "图3-6 分位点视角下的组内误差贡献分解（占位）")

    insert_heading_before(anchor, "3.4.3 典型误差模式归纳", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "结合组内最大值统计与分位点误差贡献，可以将 group-wise 场景下的量化误差大致归纳为三类。"
        "第一类是离群值挤压型误差：组内只有极少数 outlier，而大部分普通值集中在较小幅值区间，最终低分位点贡献了主要误差。"
        "第二类是大值主导型误差：组内除了最大值之外还存在多个较大值，主要误差不再集中于接近零的区域，而更多来自高分位区间。"
        "第三类则是混合型误差：低分位与高分位区间同时贡献显著误差，说明该 group 的局部结构更为复杂，单一解释已不足以覆盖全部现象。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "三类模式的区别在于主要误差的来源区间不同，因此对应的旋转设计目标也不相同。"
        "离群值挤压型更强调削弱峰值主导；大值主导型更强调重新组织多个大值之间的局部关系；"
        "而混合型则往往需要在峰值缓解与局部结构保持之间做更细致的权衡。",
        "Normal",
    )
    insert_caption_before(anchor, "表3-3 典型量化误差模式归纳（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入离群值挤压型、大值主导型与混合型误差模式的对比归纳表】")

    insert_heading_before(anchor, "3.4.4 对 group-wise 旋转设计的启示", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "上述分布统计与误差分析共同说明，group-wise 低比特量化场景下的旋转设计不能建立在单一经验假设之上。"
        "一方面，outlier 通道在不同模型和不同层中广泛存在，的确为旋转提供了应用空间；"
        "另一方面，大多数局部列块又保持着相对均匀的分布，如果对所有 group 统一施加固定旋转，"
        "则不仅难以保证收益，还可能破坏原有的局部平衡。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "进一步地，全层旋转与局部旋转对统计结构的影响并不相同，NVFP4/MXFP4 等 microscaling 格式对这种差异尤其敏感。"
        "再结合离群值挤压型、大值主导型和混合型等多种误差模式，可以得到一个更明确的结论："
        "后续算法需要以 group 为基本决策单元，根据局部统计特征和量化误差模式选择性地施加旋转，并允许不同 group、不同格式对应不同候选变换。"
        "这也是下一章构建 group-wise 旋转搜索空间与目标函数的直接出发点。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入第三章统计分析到第四章算法设计的衔接示意图】")
    insert_caption_before(anchor, "图3-7 统计分析对 group-wise 旋转设计的启示（占位）")

    insert_heading_before(anchor, "3.5 本章小结", "Heading 2")
    insert_text_paragraph_before(
        anchor,
        "本章围绕权重与激活的局部分布统计以及 group-wise 量化误差来源展开分析。"
        "首先，通过对多模型、多层的统计观察，指出 outlier 通道在大语言模型中广泛存在，但大多数通道和局部列块又保持着相对均匀的分布；"
        "这意味着旋转并非对所有 group 都同样有益，而应建立在选择性应用的前提上。"
        "其次，通过比较全层旋转与 group 粒度分块旋转对大值比例和 block 污染程度的影响，"
        "论证了在 NVFP4/MXFP4 场景下采用局部旋转、并避免单一固定旋转矩阵的必要性。"
        "最后，结合组内量化误差、最大值统计和分位点误差贡献，归纳了离群值挤压型、大值主导型和混合型三类典型误差模式，"
        "从而为后续按 group 选择候选变换、构建目标函数和设计格式特化旋转策略提供了分析基础。",
        "Normal",
    )

    try:
        doc.save(str(DOC_PATH))
        print(f"SAVED:{DOC_PATH}")
    except PermissionError:
        doc.save(str(TEMP_PATH))
        print(f"TEMP_SAVED:{TEMP_PATH}")


if __name__ == "__main__":
    build_chapter_3()
