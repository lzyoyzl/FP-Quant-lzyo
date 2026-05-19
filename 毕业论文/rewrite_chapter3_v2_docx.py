from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
TEMP_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文_第三章重写待替换.docx")
CHAPTER_TITLE = "第3章 权重/激活分布统计与量化误差来源分析"
NEXT_CHAPTER_TITLE = "第4章 面向低比特量化的 group-wise 旋转算法与系统实现"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    sub_node = OxmlElement("m:sub")
    append_expr(sub_node, sub)
    sup_node = OxmlElement("m:sup")
    append_expr(sup_node, sup)
    node.append(e)
    node.append(sub_node)
    node.append(sup_node)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_center_placeholder_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Normal"
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def build_chapter_3():
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == CHAPTER_TITLE:
            start_idx = idx
        elif start_idx is not None and text == NEXT_CHAPTER_TITLE:
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate Chapter 3 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]

    insert_text_paragraph_before(
        anchor,
        "本章围绕权重与激活的局部分布统计以及 group-wise 场景下的量化误差来源展开分析。"
        "与前一章侧重理论基础不同，本章更关注直接支撑后续算法设计的经验事实："
        "首先比较全层旋转与 group 粒度分块旋转对 outlier 扩散与缓解的影响，以论证分块旋转的必要性；"
        "随后结合局部列块的分布统计，说明并非所有 group 都适合统一施加旋转，而应当根据局部形态进行有选择的处理；"
        "最后围绕三类典型的 group-wise 量化误差来源展开分析，进一步说明不同 group 需要定制化选择旋转矩阵，"
        "从而为下一章的方法设计与搜索空间构建提供依据。",
        "Normal",
    )

    insert_heading_before(anchor, "3.1 全层旋转与 group 粒度分块旋转的影响对比", "Heading 2")

    insert_mixed_paragraph_before(
        anchor,
        [
            "为了比较不同旋转粒度对 outlier 扩散程度的影响，本文以层级大值比例和 block 污染比例作为两个基本统计量。"
            "前者描述某层中超过给定阈值的元素占比，后者描述受较大幅值影响的 group 比例，定义分别如式（3-1）和式（3-2）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("ρ", "l"),
            "(",
            "τ",
            ") = ",
            m_frac("1", m_sub("N", "l")),
            m_subsup("Σ", "j=1", m_sub("N", "l")),
            "I(",
            "|",
            m_sub("x", "l,j"),
            "| ≥ τ",
            m_sub("M", "l"),
            ")",
        ],
        "（3-1）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("η", "l"),
            "(",
            "τ",
            ") = ",
            m_frac("1", m_sub("G", "l")),
            m_subsup("Σ", "g=1", m_sub("G", "l")),
            "I(",
            m_sub("M", "g"),
            " ≥ τ",
            m_sub("M", "l"),
            ")",
        ],
        "（3-2）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若采用全层统一旋转，则同一层所有通道共同参与同一个变换；若采用 group 粒度分块旋转，则不同列块只在各自局部范围内完成线性混合。"
            "两种作用方式可统一写成式（3-3）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sup(m_sub("A", "l"), "global"),
            " = ",
            m_sub("A", "l"),
            "R,   ",
            m_sup(m_sub("A", "l"), "group"),
            " = ",
            m_sub("A", "l"),
            "blkdiag(",
            m_sub("R", "1"),
            ", ",
            m_sub("R", "2"),
            ", …, ",
            m_sub("R", "G"),
            ")",
        ],
        "（3-3）",
        doc,
    )

    insert_text_paragraph_before(
        anchor,
        "比较结果表明，二者差异的关键在于能量传播范围不同。全层旋转会将原本集中在少数通道上的大值信息投影到更广的坐标方向上，"
        "这在传统 tensor-wise 或 channel-wise 场景中有时有助于削弱单轴主导效应；"
        "但在 NVFP4/MXFP4 这类以局部 block 为共享尺度单元的 microscaling 场景中，这种扩散往往会同步抬高层级大值比例和 block 污染比例，"
        "使原本局限在少量列块中的异常能量渗入大量普通 block。"
        "已有工作《Block Rotation is All You Need for MXFP4 Quantization》也指出，全局旋转与 MXFP4 的 block scaling 机制之间存在明显冲突，"
        "其根本原因正是 outlier 能量在旋转后被扩散到过多局部块中。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "相比之下，group 粒度分块旋转将线性混合限制在局部列块内部，跨 group 的能量传播被显式阻断。"
        "因此，它更有可能在缓解局部 outlier 的同时保持原有的尺度隔离结构。"
        "从部署角度看，这种粒度还允许不同 token 在相同位置上的 group 共享同一局部变换，从而兼顾统计针对性与实现一致性。"
        "这说明，在 NVFP4/MXFP4 场景下，真正需要比较的不是是否旋转，而是是否应当将旋转范围收缩到 group 列块内部。",
        "Normal",
    )

    insert_center_placeholder_before(anchor, "【插图占位：此处插入全层旋转与 group 粒度分块旋转前后较大值比例对比图】")
    insert_caption_before(anchor, "图3-1 全层旋转与 group 粒度分块旋转对较大值比例影响的对比（占位）")
    insert_caption_before(anchor, "表3-1 全层旋转与 group 粒度分块旋转统计结果比较（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入不同旋转粒度下的大值比例、block 污染比例与误差变化比较表】")

    insert_heading_before(anchor, "3.2 局部列块分布统计与选择性旋转动机", "Heading 2")

    insert_mixed_paragraph_before(
        anchor,
        [
            "在 group 粒度下考察局部列块分布时，可以用最大幅值及其归一化形态刻画组内峰值集中程度。"
            "若记第 ",
            math_segment(["g"]),
            " 个 group 的一维样本为 ",
            math_segment([m_sub("x", "g"), " = [", m_sub("x", "g,1"), ", …, ", m_sub("x", "g,m"), "]"]),
            "，则其最大幅值与归一化表示如式（3-4）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("M", "g"),
            " = ",
            m_sub("max", "1≤i≤m"),
            "|",
            m_sub("x", "g,i"),
            "|,   ",
            m_sub("u", "g,i"),
            " = ",
            m_frac(["|", m_sub("x", "g,i"), "|"], m_sub("M", "g")),
        ],
        "（3-4）",
        doc,
    )

    insert_text_paragraph_before(
        anchor,
        "多模型、多层的统计结果表明，outlier 通道在大语言模型中广泛存在。"
        "无论是注意力模块中的投影层，还是前馈模块中的 up_proj、down_proj 等线性层，都能够观察到少量通道在幅值上显著高于其余通道。"
        "这一现象说明，旋转确实具有潜在的应用空间，因为局部异常方向往往会主导共享尺度，从而成为量化误差的重要来源之一。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "但同时也必须看到，大多数列块的局部分布其实相对均匀。"
        "这意味着旋转收益并不是均匀分布在所有 group 上的。"
        "对于原本已经较为平滑的列块，若统一施加旋转，反而可能打破其局部平衡，使峰值集中度上升、共享尺度变差，从而引入额外误差。"
        "因此，outlier 的广泛存在并不能推出“所有列块都应旋转”的结论；更合理的推论是，旋转应当建立在局部分布统计基础上，有选择地施加于确有必要的 group。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "这一点在 NVFP4/MXFP4 场景下尤为关键。由于这类格式以 group 为共享尺度单元，"
        "量化精度直接依赖于局部列块内部的统计形态。"
        "从这一意义上说，outlier 通道的广泛存在说明旋转具有必要性，而大多数列块的相对均匀性则进一步说明旋转不能采取“一刀切”的方式，"
        "而应当转向 group 粒度下的选择性施加。",
        "Normal",
    )

    insert_center_placeholder_before(anchor, "【插图占位：此处插入不同模型不同层中 outlier 通道广泛存在的统计图】")
    insert_caption_before(anchor, "图3-2 不同模型不同层中 outlier 通道分布统计（占位）")
    insert_center_placeholder_before(anchor, "【插图占位：此处插入典型层中多数列块相对均匀、少数列块异常突出的分布示意图】")
    insert_caption_before(anchor, "图3-3 典型层中局部列块分布形态示意（占位）")

    insert_heading_before(anchor, "3.3 group 粒度下的量化误差来源分析", "Heading 2")

    insert_mixed_paragraph_before(
        anchor,
        [
            "group-wise 量化的一个重要特点在于，outlier 被约束在更小的局部块中，从而减弱其对其他 block 的影响；"
            "但与此同时，group 内部的相对分布形态也会直接决定误差来源。"
            "若对第 ",
            math_segment(["g"]),
            " 个 group 采用共享尺度量化，则量化与反量化过程可写为式（3-5），相应的重建误差可由式（3-6）度量。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("q", "g,i"),
            " = clip(round(",
            m_frac(m_sub("x", "g,i"), m_sub("s", "g")),
            "), -",
            m_sub("q", "max"),
            ", ",
            m_sub("q", "max"),
            "),   ",
            m_sub("x̂", "g,i"),
            " = ",
            m_sub("s", "g"),
            m_sub("q", "g,i"),
            ",   ",
            m_sub("s", "g"),
            " = ",
            m_frac(m_sub("M", "g"), m_sub("q", "max")),
        ],
        "（3-5）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("E", "g"),
            " = ",
            m_frac("1", "m"),
            m_subsup("Σ", "i=1", "m"),
            "|",
            m_sub("x̂", "g,i"),
            " - ",
            m_sub("x", "g,i"),
            "|",
            m_sup("", "2"),
        ],
        "（3-6）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "为了进一步识别误差主要来自哪些数值区间，本文将组内元素按归一化幅值划分为若干分位点区间 "
            "bin，并用式（3-7）定义各区间的误差贡献比例。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("c", "g,k"),
            " = ",
            m_frac(
                [
                    m_sub("Σ", [m_sub("u", "g,i"), "∈", m_sub("B", "k")]),
                    "|",
                    m_sub("e", "g,i"),
                    "|",
                ],
                [
                    m_subsup("Σ", "i=1", "m"),
                    "|",
                    m_sub("e", "g,i"),
                    "|",
                ],
            ),
            ",   ",
            m_sub("e", "g,i"),
            " = ",
            m_sub("x̂", "g,i"),
            " - ",
            m_sub("x", "g,i"),
        ],
        "（3-7）",
        doc,
    )

    insert_heading_before(anchor, "3.3.1 离群值挤压型误差", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "第一类典型模式是离群值挤压型误差。此时，一个 group 中通常只包含极少数 outlier，而大部分普通值集中在较小幅值区间。"
        "在共享尺度机制下，最大值决定了整个 group 的量化步长，于是接近零附近的普通值被迫共享较粗的分辨率，"
        "最终低分位区间贡献了主要误差。"
        "这类情形对应 ppt 中“1 个 outlier + 若干 normal value”的模式，其本质仍然是经典 outlier 挤压现象在 group 粒度下的局部化表现。",
        "Normal",
    )

    insert_heading_before(anchor, "3.3.2 大值主导型误差", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "第二类典型模式是大值主导型误差。与前一种情形不同，这一类 group 中除了最大值之外，往往还存在多个较大值。"
        "此时主要误差不再集中于接近零的小值区域，而更多来自高分位区间本身。"
        "换言之，误差的主导因素不只是“一个极端 outlier 压缩了大量普通值”，而是“多个较大值共同决定了局部尺度与表示稀疏性”。"
        "这对应 ppt 中“1 个 maximum + 若干大值”的模式，也说明在某些层上仅靠削弱单一峰值并不足以显著改善量化结果。",
        "Normal",
    )

    insert_heading_before(anchor, "3.3.3 混合型误差", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "第三类是混合型误差。此时低分位和高分位区间同时贡献显著误差，既存在普通值被挤压的现象，也存在多个较大值之间共同作用的问题。"
        "这种模式表明 group 内部的局部结构更为复杂，单一解释已不足以覆盖全部现象。"
        "从统计意义上看，混合型误差往往出现在局部列块内部既有明显峰值，又存在一定规模的大值簇时，"
        "因而也是最难通过单一固定变换直接处理的一类情况。",
        "Normal",
    )

    insert_center_placeholder_before(anchor, "【插图占位：此处插入高误差 group 的量化误差与最大值统计图】")
    insert_caption_before(anchor, "图3-4 组内量化误差与最大值统计关系（占位）")
    insert_center_placeholder_before(anchor, "【插图占位：此处插入不同分位点区间的误差贡献分解图】")
    insert_caption_before(anchor, "图3-5 分位点视角下的组内误差贡献分解（占位）")
    insert_caption_before(anchor, "表3-2 三类 group-wise 量化误差来源归纳（占位）")
    insert_center_placeholder_before(anchor, "【表格占位：此处插入离群值挤压型、大值主导型与混合型误差模式的对比归纳表】")

    insert_heading_before(anchor, "3.3.4 对定制化旋转矩阵选择的启示", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "上述三类误差来源说明，group 粒度下的旋转设计不能建立在“单一矩阵对所有列块都有效”的假设之上。"
        "对于离群值挤压型误差，更重要的是削弱单一峰值对共享尺度的主导；"
        "对于大值主导型误差，更重要的是重新组织多个较大值之间的局部关系；"
        "而对于混合型误差，则需要在峰值缓解与局部结构保持之间取得平衡。"
        "这意味着，不同 group 的旋转目标并不一致，相应地，所选旋转矩阵也不应完全相同。",
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "因此，从第三章的统计分析可以得到一个直接结论：后续方法需要以 group 列块为基本决策单元，"
        "依据局部数据分布和量化误差来源对候选旋转进行定制化选择。"
        "这正是下一章构建 group-wise 旋转搜索空间、目标函数与格式特化部署策略的出发点。",
        "Normal",
    )
    insert_center_placeholder_before(anchor, "【插图占位：此处插入第三章统计分析到第四章旋转搜索设计的衔接示意图】")
    insert_caption_before(anchor, "图3-6 统计分析对定制化 group-wise 旋转设计的启示（占位）")

    insert_heading_before(anchor, "3.4 本章小结", "Heading 2")
    insert_text_paragraph_before(
        anchor,
        "本章围绕权重与激活的局部分布统计以及 group-wise 场景下的量化误差来源展开分析。"
        "首先，通过比较全层旋转与 group 粒度分块旋转对较大值比例和 block 污染比例的影响，"
        "说明在 NVFP4/MXFP4 场景下采用局部旋转的必要性。"
        "其次，通过局部列块分布统计指出，虽然 outlier 通道在不同模型和不同层中广泛存在，但大多数列块本身保持相对均匀，"
        "因此旋转不能简单地统一施加，而应建立在选择性应用的前提上。"
        "最后，围绕离群值挤压型、大值主导型和混合型三类 group-wise 量化误差来源进行分析，"
        "进一步说明不同 group 的旋转目标并不一致，因而需要在后续方法中对旋转矩阵进行定制化选择。"
        "这些结论共同构成了下一章 group-wise 旋转搜索方法设计的直接依据。",
        "Normal",
    )

    try:
        doc.save(str(DOC_PATH))
        print(f"SAVED:{DOC_PATH}")
    except PermissionError:
        doc.save(str(TEMP_PATH))
        print(f"TEMP_SAVED:{TEMP_PATH}")


if __name__ == "__main__":
    build_chapter_3()
