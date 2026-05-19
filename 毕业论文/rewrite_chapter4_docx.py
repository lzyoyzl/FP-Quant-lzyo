import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DEFAULT_DOC_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
CH4_TITLE = "第4章 面向低比特量化的 group-wise 旋转算法实现"
CH5_TITLE_PREFIX = "第5章"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    sub_node = OxmlElement("m:sub")
    append_expr(sub_node, sub)
    sup_node = OxmlElement("m:sup")
    append_expr(sup_node, sup)
    node.append(e)
    node.append(sub_node)
    node.append(sup_node)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def m_rad(expr):
    node = OxmlElement("m:rad")
    deg_hide = OxmlElement("m:degHide")
    deg_hide.set(qn("m:val"), "1")
    node.append(deg_hide)
    e = OxmlElement("m:e")
    append_expr(e, expr)
    node.append(e)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_table_before(anchor, doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    anchor._p.addprevious(table._tbl)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def build_chapter4(doc: Document, anchor):
    insert_heading_before(anchor, "4.1 问题定义与总体思路", "Heading 2")
    insert_text_paragraph_before(
        anchor,
        "在完成第三章对局部列块分布统计与量化误差来源的分析之后，本章进一步转入 group-wise 旋转算法的具体实现。"
        "本章关注的问题不再是“是否旋转”，而是如何在低比特量化约束下，围绕 Transformer block 内不同输入槽位的 group 列块，"
        "从候选变换集合中搜索更合适的局部旋转矩阵，并将搜索结果贯穿量化、导出与推理部署全流程。"
        "相较于整层统一施加单一旋转，本实现将搜索粒度固定在 weight group 上，允许同一槽位内不同列块采用不同变换。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("x", "g"), "'", " = ", m_sub("x", "g"), " ", m_sub("T", "g"),
            ",    ",
            m_sub("W", "g"), "'", " = ", m_sub("W", "g"), " ", m_sup(m_sub("T", "g"), "-T"),
        ],
        "（4-1）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "式（4-1）给出了实现中统一采用的等价变换关系。前向阶段对激活侧施加 group-wise 输入变换，"
        "权重侧则在量化前折叠对应的逆转置矩阵，从而保持线性映射不变。代码层面并未对每个线性层完全独立搜索，"
        "而是按照功能相近的输入槽位共享搜索：q_proj、k_proj 与 v_proj 共享 qkv 变换；gate_proj 与 up_proj 共享 gate_up 变换；"
        "o_proj 与 down_proj 则分别独立处理。这样既保留了局部搜索的灵活性，也控制了实际搜索与导出的复杂度。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("T", "g", "*"), " = arg min_{T ∈ C} ",
            m_sub("S", "g"),
            "(",
            m_sub("T", "g"),
            ")",
        ],
        "（4-2）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "对每个 group 列块，搜索过程都可写成式（4-2）的形式，其中候选集合 C 来自预定义的局部变换矩阵库，"
        "评分函数 S_g(·) 则由量化目标决定。后续各节将依次介绍候选变换、目标函数以及实现链路。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "【插图占位：此处插入第4章总体流程图，展示校准数据收集、group-wise 旋转搜索、量化导出与推理接入的完整链路。】",
        "Normal",
    )
    insert_caption_before(anchor, "图4-1 第4章 group-wise 旋转算法总体实现流程图（占位）")

    insert_heading_before(anchor, "4.2 group-wise 旋转搜索空间设计", "Heading 2")
    insert_heading_before(anchor, "4.2.1 变换粒度与作用范围", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "从代码实现看，搜索粒度首先受量化格式约束。主程序在解析参数后会对 NVFP4 与 MXFP4 的 group 大小进行统一约束："
        "NVFP4 对应 w_group_size=a_group_size=16，MXFP4 对应 w_group_size=a_group_size=32。"
        "在此基础上，build_block_input_transforms 会按 Transformer block 的四个输入槽位分别构造搜索任务，"
        "其中 qkv、o 与 gate_up 的输入维度均来自 hidden_size，down 的输入维度来自 intermediate_size。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("T", "g"), " ",
            m_sup(["(", m_sup(m_sub("T", "g"), "-T"), ")"], "T"),
            " = ",
            m_sub("I", "g"),
        ],
        "（4-3）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "对于最终选中的 group-wise 变换，代码并不只保存一个统一矩阵，而是将每个列块对应的前向矩阵与逆转置矩阵堆叠为三维矩阵库。"
        "MixedGroupTransform 在初始化时会显式检查式（4-3）所表示的线性等价条件，以避免由于数值实现偏差导致前向变换与权重折叠不一致。",
        "Normal",
    )
    insert_caption_before(anchor, "表4-1 四类输入槽位的搜索对象与作用范围")
    insert_table_before(
        anchor,
        doc,
        [
            ["槽位", "共享的线性层", "搜索粒度", "输入特征维度来源"],
            ["qkv", "q_proj / k_proj / v_proj", "按 group 列块逐组搜索", "hidden_size"],
            ["o", "o_proj", "按 group 列块逐组搜索", "hidden_size"],
            ["gate_up", "gate_proj / up_proj", "按 group 列块逐组搜索", "hidden_size"],
            ["down", "down_proj", "按 group 列块逐组搜索", "intermediate_size"],
        ],
    )

    insert_heading_before(anchor, "4.2.2 Identity 变换", "Heading 3")
    insert_equation_table_before(
        anchor,
        [m_sub("T", "id"), " = ", m_sub("I", "g")],
        "（4-4）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "Identity 变换对应式（4-4），即不对局部列块施加任何额外旋转。"
        "从实现角度看，将 Identity 纳入候选集合并非仅用于提供基线，而是为了给“本身已较为均匀”的 group 保留不变换选项。"
        "第三章已经表明，并非所有列块都需要被重新组织；因此如果某个 group 在搜索目标下由 Identity 取得最小误差，"
        "当前实现会直接保留原始方向，而不会为了统一形式强行施加旋转。",
        "Normal",
    )

    insert_heading_before(anchor, "4.2.3 Hadamard 类变换", "Heading 3")
    insert_equation_table_before(
        anchor,
        [
            m_sub("T", "had"), " = ", m_frac("1", m_rad("g")), " ", m_sub("H", "g"),
        ],
        "（4-5）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "Hadamard 变换是当前候选集中最基础的结构化正交矩阵之一。代码实现采用按 group_size 分块的局部 Hadamard 变换，"
        "其核心形式如式（4-5）所示。由于 Hadamard 矩阵元素仅取 ±1，且经归一化后满足自逆性质，"
        "因此它既适合表达“将单一峰值扩散到组内多个方向”的作用，也便于在训练态与导出态之间保持数值一致。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("T", "gsr"), " = ", m_frac("1", m_rad("g")), " ", m_sub("H", "g"), " ", m_sub("P", "gsr"),
        ],
        "（4-6）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "在 Hadamard 的基础上，代码还实现了 GSR 变换。其做法是先构造 Hadamard 基，再按照列向量符号变化次数对列进行重排，"
        "相当于在式（4-6）中引入一个重排序矩阵 P_gsr。相比固定顺序的 Hadamard，GSR 仍保持正交与自逆特性，"
        "但对组内方向的排列方式进行了再组织，因而可以视为 Hadamard 类结构化变换的一个扩展候选。",
        "Normal",
    )

    insert_heading_before(anchor, "4.2.4 其他候选正交变换", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "除上述两类结构化变换外，默认候选集合中还包括 DCT、DST 与 Householder。"
        "其中 DCT 与 DST 都是按 group_size 构造的 type-II 正交基矩阵，并在实现中按块重复为 block-diagonal 形式；"
        "它们强调的是频域基方向上的重分解，与 Hadamard 类变换相比具有更平滑的基向量结构。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("T", "hh"), " = ", m_sub("I", "g"), " - ",
            m_frac(["2", " ", "v", " ", m_sup("v", "T")], [m_sup("v", "T"), " ", "v"]),
        ],
        "（4-7）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "Householder 变换是本实现中特别值得强调的一类候选矩阵。"
        "对任意非零向量 v，式（4-7）定义了一个关于超平面的反射矩阵。"
        "与 Hadamard、DCT、DST 等固定基矩阵相比，Householder 只需一个方向向量即可构造出完整变换，"
        "形式紧凑，但仍然能够显式改变局部主导方向。",
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sup(m_sub("T", "hh"), "T"), " ", m_sub("T", "hh"), " = ", m_sub("I", "g"),
            ",    ",
            m_sup(m_sub("T", "hh"), "-1"), " = ", m_sub("T", "hh"),
        ],
        "（4-8）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "由式（4-8）可知，Householder 矩阵既是正交矩阵，又满足自逆性质，因此非常适合与当前“激活侧前向旋转、"
        "权重侧逆变换折叠”的部署方式配合使用。代码中首先随机初始化一个长度为 group_size 的向量 v，"
        "随后在每个 group 上重复同一个反射块并形成 block-diagonal 矩阵。"
        "这一做法虽然不涉及额外优化参数，但比完全固定的基矩阵更容易提供不同方向上的局部几何变化。"
        "此外，项目中还保留了 fast_food 变换实现，不过默认候选集合并未启用该分支，因此本文实验主要围绕默认六种候选变换展开。",
        "Normal",
    )
    insert_caption_before(anchor, "表4-2 候选旋转矩阵及其实现特性")
    insert_table_before(
        anchor,
        doc,
        [
            ["候选变换", "实现类", "主要构造", "实现特点"],
            ["Identity", "IdentityTransform", "单位矩阵", "提供“不旋转”基线，并允许均匀 group 保持原方向"],
            ["Hadamard", "HadamardTransform", "归一化 Hadamard 基", "结构简单、自逆，适合快速扩散峰值"],
            ["GSR", "GSRTransform", "Hadamard 基 + 列重排", "保留正交结构，同时改变基向量顺序"],
            ["DCT / DST", "DCTTransform / DSTransform", "type-II 正交基", "更平滑的频域基，按块重复构造"],
            ["Householder", "HouseholderTransform", "单反射矩阵", "由方向向量生成，正交且自逆"],
            ["FastFood（可选）", "FastFoodTransform", "结构化随机矩阵", "代码保留实现，但默认搜索空间不启用"],
        ],
    )

    insert_heading_before(anchor, "4.3 搜索目标函数与算法流程", "Heading 2")
    insert_heading_before(anchor, "4.3.1 基于 MSE 的目标函数", "Heading 3")
    insert_equation_table_before(
        anchor,
        [
            "Δ", m_sub("W", "g"),
            "(",
            m_sub("T", "g"),
            ") = Q(",
            m_sub("W", "g"), " ", m_sup(m_sub("T", "g"), "-T"),
            ") - ",
            m_sub("W", "g"), " ", m_sup(m_sub("T", "g"), "-T"),
        ],
        "（4-9）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("L", "mse", "(g)"),
            "(",
            m_sub("T", "g"),
            ") = ",
            m_frac(
                ["||", "Δ", m_sub("W", "g"), "(", m_sub("T", "g"), ")", "||", m_sub("F", "2")],
                "o g",
            ),
        ],
        "（4-10）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "MSE 目标直接作用在权重域。对第 g 个列块，先按式（4-9）构造量化残差，再用式（4-10）计算平均重建误差。"
        "由于该目标不依赖额外统计量，RTN 路径在 objective=auto 时会默认回落到这一评分方式。"
        "从实现角度看，MSE 分支最直接、代价最低，适合快速比较不同候选矩阵在局部权重分布上的重建能力。",
        "Normal",
    )

    insert_heading_before(anchor, "4.3.2 基于 COV 的目标函数", "Heading 3")
    insert_equation_table_before(
        anchor,
        [
            "Σ", m_sub("g", "'"), " = ",
            m_sup(m_sub("T", "g"), "T"), " ", m_sub("Σ", "g"), " ", m_sub("T", "g"),
        ],
        "（4-11）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("L", "cov", "(g)"),
            "(",
            m_sub("T", "g"),
            ") = ",
            m_frac(
                ["Tr(", "Δ", m_sub("W", "g"), "(", m_sub("T", "g"), ")", " ", "Σ", m_sub("g", "'"), " ",
                 m_sup(["Δ", m_sub("W", "g"), "(", m_sub("T", "g"), ")"], "T"), ")"],
                "o",
            ),
        ],
        "（4-12）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "仅用重建 MSE 衡量局部变换的优劣，无法反映该残差在真实输入分布下对层输出造成的影响。"
        "为此，GPTQ 路径在 objective=auto 时默认采用协方差加权目标。"
        "实现上首先通过 forward hook 收集 qkv、o、gate_up 与 down 四个槽位的 group 输入协方差，"
        "随后按式（4-11）将协方差旋转到候选变换对应的坐标系，再用式（4-12）评估残差对输出误差的加权影响。"
        "这一形式与 GPTQ 中 Hessian 加权的思想保持一致，但作用粒度从整层进一步缩小到了局部列块。",
        "Normal",
    )

    insert_heading_before(anchor, "4.3.3 基于 act_mse 与 J_tail 的目标函数", "Heading 3")
    insert_equation_table_before(
        anchor,
        [
            m_subsup("L", "act", "(g)"),
            "(",
            m_sub("T", "g"),
            ") = ",
            m_frac(
                ["||Q(", m_sub("X", "g"), " ", m_sub("T", "g"), ") - ", m_sub("X", "g"), " ", m_sub("T", "g"), "||", m_sub("F", "2")],
                "N g",
            ),
        ],
        "（4-13）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("J", "tail", "(g)"),
            "(",
            m_sub("T", "g"),
            ") = ",
            m_subsup("L", "base", "(g)"),
            "(",
            m_sub("T", "g"),
            ") + λ ",
            m_subsup("L", "tail", "(g)"),
            "(",
            m_sub("T", "g"),
            ")",
        ],
        "（4-14）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("L", "tail", "(g)"),
            "(",
            m_sub("T", "g"),
            ") = ",
            m_frac("∑ α_b MSE_b", "∑ α_b"),
        ],
        "（4-15）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "当关注点从权重域进一步转向激活域时，代码提供了 act_mse 目标，如式（4-13）所示。"
        "它直接在校准样本上评估旋转后激活的量化重建误差，因此要求同时开启激活量化，并满足 a_group_size 与 w_group_size 一致。"
        "相比之下，J_tail 目标采用式（4-14）的复合形式，在基础项之外额外加入按分位点加权的尾部误差。"
        "尾部项的实际实现如式（4-15）所示，其中 MSE_b 表示第 b 个分位点区间内的均方误差，α_b 则由 tail_weight_mode 决定。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "当前实现支持的尾部权重模式包括 a_low、b_high、mixed_uniform、mixed_middle 以及兼容旧实验的 two_tail。"
        "其中 a_low 更强调低分位误差，适用于“离群值挤压型”列块；b_high 更强调高分位误差，适用于“大值主导型”列块；"
        "mixed_uniform 与 mixed_middle 则分别对应均匀加权和中间分位更高的折中方式。"
        "此外，代码对 jtail 的基础项与尾部来源做了显式约束：例如 base_loss=act_mse 时，tail_source 必须取 activation；"
        "而 activation-tail 仅允许与 cov 或 act_mse 组合，以保证评分对象的一致性。",
        "Normal",
    )

    insert_heading_before(anchor, "4.3.4 AUTO 自适应选择策略", "Heading 3")
    insert_equation_table_before(
        anchor,
        [
            m_sub("s", "out"), " = ", m_frac(m_sub("q", "99"), [m_sub("q", "90"), " + ε"]),
            ",    ",
            m_sub("r", "bulk"), " = ", m_frac(m_sub("q", "50"), [m_sub("q", "99"), " + ε"]),
            ",    ",
            m_sub("r", "high"), " = ", m_frac(m_sub("q", "90"), [m_sub("q", "99"), " + ε"]),
        ],
        "（4-16）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "当 tail_weight_mode 取 auto_abm 时，系统会先对当前 group 的绝对值分布做启发式分类，再在 a_low、b_high 与 mixed_uniform 之间自动选择权重模式。"
        "具体而言，代码先计算式（4-16）中的三个分布比值。若 s_out 较大而 r_bulk 较小，说明少量尖峰显著高于主体分布，"
        "则归为 A 类；若 r_bulk 与 r_high 都已接近峰值，说明大量数值整体偏高，则归为 B 类；其余情形统一归入 mixed。"
        "当前阈值实现为：A 类满足 s_out≥1.6 且 r_bulk≤0.30；B 类满足 r_bulk≥0.45 且 r_high≥0.80；否则采用 mixed_uniform。"
        "这一策略可以作用于权重分布，也可以作用于激活样本，具体取决于 tail_source 的设定。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "除尾部模式的自动分类外，objective=auto 还会根据量化路径切换基础评分函数。"
        "在 GPTQ 路径中，auto 默认解析为 cov，以便充分利用输入协方差；"
        "在 RTN 路径中，auto 默认解析为 mse，以降低额外统计开销。"
        "这样既保持了两条量化路径的实现一致性，又兼顾了不同路径对统计量的依赖差异。",
        "Normal",
    )

    insert_heading_before(anchor, "4.3.5 group-wise 旋转搜索算法流程", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "在具体实现上，search_best_group_transform 会先对全部候选矩阵进行一次性实例化，并预先提取每个候选的前向矩阵与逆转置矩阵。"
        "若某个候选矩阵不可逆，系统会直接跳过；若候选类给出的 inv_t 与严格意义上的逆转置存在明显偏差，"
        "则实现会以前向矩阵的严格 inverse-transpose 结果为准，并记录相应提示信息。"
        "这一设计保证了后续量化、导出与推理阶段都围绕同一组线性等价矩阵展开。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "随后，算法按列块顺序遍历每个 group。对于当前 group，系统会从共享槽位对应的权重张量中切出局部列块，"
        "并根据目标函数决定是否同时切出激活样本与输入协方差。每个候选矩阵都会在当前 group 上独立打分，"
        "最终取得最小目标值的候选即被记为该列块的最优变换。所有列块搜索结束后，选中的前向矩阵与逆转置矩阵会被堆叠为三维矩阵库，"
        "并连同所选候选名称、尾部权重模式摘要一起封装到 MixedGroupTransform 中。",
        "Normal",
    )
    insert_caption_before(anchor, "表4-3 搜索目标函数与所需统计量对应关系")
    insert_table_before(
        anchor,
        doc,
        [
            ["目标函数", "评分域", "需要的附加统计量", "典型默认路径"],
            ["mse", "权重域", "无", "RTN 下 auto 默认选择"],
            ["cov", "权重域 + 输入分布", "group 输入协方差", "GPTQ 下 auto 默认选择"],
            ["act_mse", "激活域", "槽位输入样本", "显式开启 activation 量化时使用"],
            ["jtail", "权重域或激活域", "分位点 bin、权重模式、必要时的样本/协方差", "用于强化特定误差来源"],
        ],
    )
    insert_text_paragraph_before(
        anchor,
        "【插图占位：此处插入 group-wise 旋转搜索流程图，展示候选矩阵预计算、逐 group 打分、最优变换选择以及 MixedGroupTransform 组装过程。】",
        "Normal",
    )
    insert_caption_before(anchor, "图4-2 group-wise 旋转搜索算法流程图（占位）")

    insert_heading_before(anchor, "4.4 量化导出与等价部署实现", "Heading 2")
    insert_heading_before(anchor, "4.4.1 NVFP 量化导出", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "在 NVFP4 路径中，主程序会首先将 weight 与 activation 的 group 大小统一固定为 16，"
        "并将 scale_precision 固定为 e4m3。对于 E4M3 标度，代码会在权重旋转完成后先调用量化器计算局部尺度，"
        "据此预热并冻结 global scale 的跟踪过程；若启用了 fuse_global_scale，则还会对 qkv 与 gate_up 两组共享槽位分别取最小 global scale 进行融合。"
        "这样做的目的，是在局部 group scale 与全局尺度之间保持稳定的一致关系，避免不同共享线性层之间出现额外漂移。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "导出阶段，realquant 模式会保存打包后的 qweight、按 E4M3 转换后的 scales、前向与逆向变换矩阵、"
        "weight_global_scale 以及 act_global_scale；pseudoquant 模式则保存反量化后的 dqweight 与同样的矩阵和尺度元数据。"
        "虽然导出字段名仍沿用 forward_hadamard_matrix 与 backward_hadamard_matrix 的历史命名，"
        "但其中实际存放的已经是搜索得到的任意候选变换矩阵，而不再局限于 Hadamard。",
        "Normal",
    )

    insert_heading_before(anchor, "4.4.2 MXFP 量化导出", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "在 MXFP4 路径中，主程序会将 weight 与 activation 的 group 大小固定为 32，并将 scale_precision 固定为 e8m0。"
        "相比 NVFP4，这一路径不再依赖 E4M3 风格的局部尺度格式，但导出结构保持一致：realquant 保存 qweight 与 scales，"
        "pseudoquant 保存 dqweight，同时统一附带前向矩阵、逆向矩阵以及 weight_global_scale、act_global_scale。"
        "因此，从检查点结构角度看，两种格式共享同一套 group-wise 旋转导出接口，只是在 group size 与尺度精度上有所区别。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "需要指出的是，pseudoquant 导出允许直接保存完整的 per-group 矩阵库，即形状为 [num_groups, g, g] 的三维张量；"
        "而 realquant 由于后端接口仍以单矩阵为主，当前实现会回退到“选择出现频率最高的代表性块”进行导出。"
        "若某一层实际使用了多种候选矩阵，代码会给出兼容性警告。"
        "这一差异并不影响本文方法在伪量化推理中的完整复现，但会限制 realquant 后端对完全混合变换的直接表达能力。",
        "Normal",
    )

    insert_heading_before(anchor, "4.4.3 激活侧前向旋转与权重侧逆变换折叠", "Heading 3")
    insert_equation_table_before(
        anchor,
        [
            "y = (", m_sub("x", "g"), " ", m_sub("T", "g"), ")",
            m_sup(["(", m_sub("W", "g"), " ", m_sup(m_sub("T", "g"), "-T"), ")"], "T"),
            " + b = ", m_sub("x", "g"), " ", m_sup(m_sub("W", "g"), "T"), " + b",
        ],
        "（4-17）",
        doc,
    )
    insert_text_paragraph_before(
        anchor,
        "式（4-17）给出了当前部署方式的核心等价关系。实现中并不在推理阶段同时显式保留“旋转前权重”与“旋转后激活”，"
        "而是将输入侧旋转与权重侧逆变换折叠成一组配对矩阵。"
        "QLinear 在训练态与 fix_parametrization 过程中都会先按 inv_t=True 对权重做局部折叠，再进行量化；"
        "推理态则只需按照导出的前向矩阵对激活进行相应处理，即可恢复与原始线性映射一致的结果。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "为了避免自定义 inv_t 实现带来的数值偏差，搜索阶段和导出阶段都会优先以“前向矩阵的严格 inverse-transpose”作为 backward matrix。"
        "这也是 MixedGroupTransform 在初始化时要额外执行线性等价检查的原因。"
        "换言之，当前实现并不是仅依赖“候选矩阵理论上应当自逆”这一假设，而是在代码层面对前向矩阵与折叠矩阵之间的配对关系做了显式约束。",
        "Normal",
    )

    insert_heading_before(anchor, "4.4.4 推理框架集成", "Heading 3")
    insert_text_paragraph_before(
        anchor,
        "在推理侧，inference_lib 中的 pseudoquant_linear_fns 已经对 group-wise 矩阵库做了专门兼容。"
        "当加载到的 forward_hadamard_matrix 为三维张量时，forward_pseudoquantize 会先调用 _apply_groupwise_transform，"
        "按 group 对输入或权重做局部矩阵乘法，随后再将已有的 MXFP4 或 NVFP4 kernel 复用到“变换后数据 + 单位块矩阵”的路径上。"
        "这一设计避免了为每一种候选变换单独重写低层量化 kernel，也使 group-wise 搜索结果能够直接进入既有推理框架。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "与此同时，线性模块的状态加载逻辑也被扩展为可接受三维矩阵库。"
        "若检查点中保存的是 pseudoquant 导出的 dqweight，模块在 pre_forward 阶段会直接使用该 dqweight，"
        "而不会再次依赖 master weight 触发伪量化过程；对 forward_hadamard_matrix 与 backward_hadamard_matrix 的缓冲区形状，"
        "加载器也会根据状态字典中的真实尺寸自动调整。"
        "这意味着本文方法不仅完成了搜索与导出，还在推理接口层面打通了从 group-wise 旋转结果到伪量化运行时的完整接入链路。",
        "Normal",
    )

    insert_heading_before(anchor, "4.5 自动化评测与结果汇总", "Heading 2")
    insert_text_paragraph_before(
        anchor,
        "除搜索与导出主体外，项目还将量化、导出与评测组织在统一命令行入口中。"
        "主程序负责模型加载、校准集读取、格式约束、搜索参数解析、量化导出与评测开关管理；"
        "transform_search 开启后，日志中会输出当前 objective、协方差与样本是否启用，以及 qkv、o、gate_up、down 四个槽位对应的候选矩阵计数摘要。"
        "这些统计信息为后续实验章节中的整体效果分析、目标函数对比与消融实验提供了直接可复现的实现基础。",
        "Normal",
    )
    insert_text_paragraph_before(
        anchor,
        "在评测接口方面，主程序同时预留了困惑度评测、OpenLLM 任务评测与 LM-Eval 任务评测选项。"
        "因此，本文的 group-wise 旋转算法并不是停留在离线打分阶段，而是与模型导出、推理接入和任务评测形成了一条贯通链路。"
        "这也是下一章能够围绕整体效果、不同目标函数、不同旋转策略与不同量化格式展开系统对比的实现前提。",
        "Normal",
    )

    insert_heading_before(anchor, "4.6 本章小结", "Heading 2")
    insert_text_paragraph_before(
        anchor,
        "本章围绕面向低比特量化的 group-wise 旋转算法实现展开，首先依据代码给出了四类输入槽位上的搜索粒度与候选矩阵设计，"
        "随后说明了 mse、cov、act_mse 与 jtail 等搜索目标在当前实现中的定义方式、统计量依赖和自动选择策略，"
        "并进一步梳理了 RTN/GPTQ 两条量化路径下的导出方式、等价部署关系与推理框架接入细节。"
        "由此可见，本文提出的方法并非停留在概念层面的局部旋转设想，而是已经在候选矩阵、评分函数、导出协议与运行时接口上形成了完整实现。"
        "下一章将在此基础上对算法的整体效果、目标函数差异、旋转策略消融以及不同量化格式下的表现进行实验分析。",
        "Normal",
    )


def rebuild_chapter4(doc_path: Path):
    doc = Document(str(doc_path))
    paragraphs = list(doc.paragraphs)

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == CH4_TITLE:
            start_idx = idx
        elif start_idx is not None and text.startswith(CH5_TITLE_PREFIX):
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate Chapter 4 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]
    build_chapter4(doc, anchor)

    doc.save(str(doc_path))
    print(f"SAVED:{doc_path}")


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_DOC_PATH
    rebuild_chapter4(target)
