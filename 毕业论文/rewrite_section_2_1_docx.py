from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
SECTION_TITLE = "2.1　后训练低比特量化基础"
NEXT_SECTION_TITLE = "2.2　NVFP4 与 MXFP4 微缩放量化格式"

INTRO = (
    "后训练低比特量化（Post-Training Quantization, PTQ）是指在不重新训练模型参数的前提下，"
    "仅利用少量校准样本或张量统计信息，将高精度权重与激活映射为低比特表示的过程。"
    "其直接目标在于降低模型参数存储、显存访存和推理带宽开销，而更核心的问题则是如何在位宽显著下降后，"
    "尽可能保持原模型的函数行为与任务精度。对于大语言模型而言，量化误差不仅来源于单个数值的舍入，"
    "还会经由残差连接、注意力映射和前馈网络在层间传播与累积。"
    "因此，本节先从量化映射与误差来源出发建立基本分析框架，再对典型 PTQ 方法进行概述，并重点说明 GPTQ 的基本推导过程。"
)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    sub_node = OxmlElement("m:sub")
    append_expr(sub_node, sub)
    sup_node = OxmlElement("m:sup")
    append_expr(sup_node, sup)
    node.append(e)
    node.append(sub_node)
    node.append(sup_node)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def rebuild_section():
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == SECTION_TITLE:
            start_idx = idx
        elif start_idx is not None and text == NEXT_SECTION_TITLE:
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate section 2.1 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]

    insert_text_paragraph_before(anchor, INTRO, "Normal")
    insert_heading_before(anchor, "2.1.1　量化映射与误差来源", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "从数值表示角度看，量化首先需要确定目标码值集合、缩放因子 ",
            math_segment(["s"]),
            " 与零点 ",
            math_segment(["z"]),
            "，再将实值张量投影到可表示格点上。设目标位宽对应的整数码值范围为 ",
            math_segment([m_sub("q", "min")]),
            " 至 ",
            math_segment([m_sub("q", "max")]),
            "，则对标量 ",
            math_segment(["x"]),
            " 的均匀量化映射可表示为式（2-1）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            "q = clip(round(",
            m_frac("x", "s"),
            ") + z, ",
            m_sub("q", "min"),
            ", ",
            m_sub("q", "max"),
            ")",
        ],
        "（2-1）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        ["量化后需通过反量化恢复近似实值，其表达式如式（2-2）所示。"],
        "Normal",
    )
    insert_equation_table_before(anchor, ["x̂ = s(q - z)"], "（2-2）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "其中，",
            math_segment(["q"]),
            " 为量化码值，",
            math_segment([m_sub("q", "min")]),
            " 与 ",
            math_segment([m_sub("q", "max")]),
            " 分别表示目标位宽下的可表示下界与上界。"
            "在对称量化中通常取 ",
            math_segment(["z = 0"]),
            "，以便简化部署；在非对称量化中则通过零点补偿处理分布偏移。"
            "若样本未发生截断，则舍入误差的绝对值满足式（2-3）所示上界。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["|x̂ - x| ≤ ", m_frac("s", "2")],
        "（2-3）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若记单点量化误差为 ",
            math_segment(["e(x)"]),
            "，则其基本定义可写为式（2-4）。",
        ],
        "Normal",
    )
    insert_equation_table_before(anchor, ["e(x) = x̂ - x"], "（2-4）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "在面向大模型推理的低比特实现中，缩放因子通常不会逐元素独立存储，而是按通道、按 token 或按 group 共享。"
            "若以 group 为基本量化粒度，并采用常见的最大值尺度估计方式，则第 ",
            math_segment(["g"]),
            " 个 group 的共享缩放因子可写为式（2-5）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("s", "g"),
            " = ",
            m_frac(["max |", m_sub("x", "g,i"), "|"], m_sub("q", "max")),
        ],
        "（2-5）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "在此基础上，第 ",
            math_segment(["g"]),
            " 个 group 内第 ",
            math_segment(["i"]),
            " 个元素的量化与反量化过程可统一写为式（2-6）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("q", "g,i"),
            " = Q(",
            m_sub("x", "g,i"),
            "; ",
            m_sub("s", "g"),
            "),   ",
            m_sub("x̃", "g,i"),
            " = ",
            m_sub("s", "g"),
            m_sub("q", "g,i"),
        ],
        "（2-6）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        ["若以组内平方误差衡量量化损失，则对应的 group 级重建误差可写为式（2-7）。"],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("E", "g"),
            " = ",
            m_subsup("Σ", "i=1", "m"),
            " (",
            m_sub("x̃", "g,i"),
            " - ",
            m_sub("x", "g,i"),
            ")²",
        ],
        "（2-7）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "式（2-5）至式（2-7）表明，group-wise 量化的关键并不只是单个元素是否被正确舍入，"
            "更在于共享尺度 ",
            math_segment([m_sub("s", "g")]),
            " 能否同时适配该 group 内的整体统计分布。"
            "当组内存在显著重尾、离群值或方差差异时，少量大幅值元素往往会主导 ",
            math_segment([m_sub("s", "g")]),
            " 的取值，从而压缩多数中小值元素的有效分辨率，这正是共享 scale 失配误差的本质来源。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "进一步从线性层输出角度看，量化对模型行为的影响可以通过量化前后输出差异来刻画。"
            "设输入激活为 ",
            math_segment(["X"]),
            "，权重为 ",
            math_segment(["W"]),
            "，量化后的激活与权重分别记为 ",
            math_segment(["X̂"]),
            " 与 ",
            math_segment(["Ŵ"]),
            "，则输出误差的基本分解如式（2-8）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["ΔY = XW - X̂Ŵ = (X - X̂)W + X̂(W - Ŵ)"],
        "（2-8）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        ["若进一步采用 Frobenius 范数衡量整个线性层的重建损失，则可得到式（2-9）。"],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["L = ", m_sup(m_sub("‖ΔY‖", "F"), "2")],
        "（2-9）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "由式（2-8）与式（2-9）可知，量化误差并非单一来源。对于后训练低比特量化，可将其概括为四类："
            "其一是格点舍入带来的离散化误差；其二是超出表示范围造成的截断误差；"
            "其三是 group 共享尺度与局部统计不匹配导致的块内失配误差；其四是单层近似经残差连接和层间传递后形成的传播误差。"
            "此外，缩放因子的估计通常依赖少量校准样本，若样本对长尾 token 或极端激活模式覆盖不足，"
            "则尺度估计偏差会在 FP4 等极低位宽场景下被进一步放大。"
            "因此，后续方法设计的重点并不只是采用更低的位宽表示，还包括如何通过更合适的统计建模与表示变换降低重建损失。",
        ],
        "Normal",
    )

    insert_heading_before(anchor, "2.1.2　典型 PTQ 方法概述", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "从方法机理看，典型 PTQ 方法大致可以分为直接舍入型、误差重建型和尺度重参数化型。"
            "其中，RTN 属于最直接的基线方法，其思想是在统计范围确定后直接执行就近舍入；"
            "SmoothQuant[3] 通过离线尺度迁移缓解激活离群值；AWQ[2] 则利用激活信息识别关键通道并保护重要权重。"
            "相比之下，GPTQ[1] 并不满足于逐元素独立舍入，而是通过二阶信息显式建模量化对层输出的扰动，"
            "因而成为后训练权重量化中最具代表性的误差重建方法之一。下面重点给出 GPTQ 的基本推导过程。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "设某线性层的一行权重为 ",
            math_segment(["w ∈ ", m_sup("ℝ", "d")]),
            "，校准输入矩阵为 ",
            math_segment(["X ∈ ", m_sup("ℝ", "d×N")]),
            "，量化后的近似为 ",
            math_segment(["ŵ"]),
            "。GPTQ 不直接最小化参数级距离 ",
            math_segment(["‖w - ŵ‖"]),
            "，而是最小化该行权重作用于校准输入后的输出重建误差，即式（2-10）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["ŵ* = ", m_sub("arg min", "ŵ∈𝒬"), " ", m_frac(m_sup("‖wX - ŵX‖", "2"), "2")],
        "（2-10）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "令量化残差 ",
            math_segment(["δ = w - ŵ"]),
            "，并记由校准输入诱导的二阶矩阵为 ",
            math_segment(["H"]),
            "，则式（2-10）可在二阶近似下改写为式（2-11）。"
            "在实际实现中，为抑制病态性并提高逆矩阵计算的稳定性，通常在 ",
            math_segment(["H"]),
            " 上加入阻尼项。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["L(δ) ≈ ", m_frac("1", "2"), "δHδᵀ,   H = XXᵀ + λI"],
        "（2-11）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "GPTQ 的关键在于顺序量化。设当前被量化的坐标为 ",
            math_segment(["q"]),
            "，尚未量化的自由变量集合为 ",
            math_segment(["F"]),
            "，并令当前位置量化残差为 ",
            math_segment([m_sub("δ", "q"), " = ", m_sub("w", "q"), " - Q(", m_sub("w", "q"), ")"]),
            "。将变量按 ",
            math_segment(["q"]),
            " 与 ",
            math_segment(["F"]),
            " 进行分块后，可得式（2-12）的二次型展开。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            "L = ",
            m_frac("1", "2"),
            "(",
            m_sub("H", "qq"),
            m_sup(m_sub("δ", "q"), "2"),
            " + 2",
            m_sub("δ", "q"),
            m_sub("H", "qF"),
            m_sub("δ", "F"),
            " + ",
            m_sub("δ", "F"),
            "ᵀ",
            m_sub("H", "FF"),
            m_sub("δ", "F"),
            ")",
        ],
        "（2-12）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "对式（2-12）关于自由变量 ",
            math_segment([m_sub("δ", "F")]),
            " 求偏导并令其为零，可得未量化部分的最优补偿条件，如式（2-13）所示；"
            "进一步即可得到最优补偿解，如式（2-14）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_frac("∂L", m_sub("∂δ", "F")),
            " = ",
            m_sub("H", "Fq"),
            m_sub("δ", "q"),
            " + ",
            m_sub("H", "FF"),
            m_sub("δ", "F"),
            " = 0",
        ],
        "（2-13）",
        doc,
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("δ", "F"),
            "* = -",
            m_sup(m_sub("H", "FF"), "-1"),
            m_sub("H", "Fq"),
            m_sub("δ", "q"),
        ],
        "（2-14）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "式（2-14）说明，当某一权重被量化后，其误差并不需要完全保留在当前位置，"
            "而可以沿着输入相关性所刻画的方向分配到尚未量化的其余权重上。"
            "将式（2-14）代回式（2-12），即可得到当前位置被量化时对应的增量损失，如式（2-15）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("ΔL", "q"),
            " = ",
            m_frac(m_sup(m_sub("δ", "q"), "2"), ["2", m_sub(["(", m_sup("H", "-1"), ")"], "qq")]),
        ],
        "（2-15）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "在此基础上，GPTQ 在得到当前位置量化值 ",
            math_segment(["Q(", m_sub("w", "q"), ")"]),
            " 后，会对剩余未量化权重执行一次闭式更新，"
            "以显式补偿当前量化带来的输出扰动，其更新形式如式（2-16）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("w", "F"),
            " ← ",
            m_sub("w", "F"),
            " - ",
            m_frac(m_sub("δ", "q"), m_sub(["(", m_sup("H", "-1"), ")"], "qq")),
            m_sub(["(", m_sup("H", "-1"), ")"], "Fq"),
        ],
        "（2-16）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "式（2-15）与式（2-16）揭示了 GPTQ 的核心思想：量化不是逐元素独立完成的，"
            "而是借助校准输入所诱导的二阶相关性，对后续尚未量化的权重进行误差回填。"
            "因此，GPTQ 优化的并非简单的参数级距离，而是更接近真实推理行为的层输出重建误差。"
            "在工程实现中，GPTQ 通常按列块逐步处理权重矩阵，并结合阻尼 Hessian、块内更新与逆矩阵递推策略降低显存和计算开销，"
            "从而能够在大规模语言模型上完成离线量化。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "相较而言，RTN 可视为最直接的基线方法，其实现最为简洁，但对离群值与相关性结构缺乏显式建模；"
            "SmoothQuant[3] 通过离线尺度迁移将部分激活动态范围转移到权重侧，以缓解激活侧的量化压力；"
            "AWQ[2] 则利用激活统计识别关键通道，对重要权重给予更高保真度。"
            "总体来看，这些方法分别回答了“如何舍入”“如何重分配尺度”以及“如何保护关键通道”三个问题，"
            "而 GPTQ 所体现的二阶误差重建思想，也为后续更细粒度的低比特格式优化提供了重要参照。",
        ],
        "Normal",
    )

    doc.save(str(DOC_PATH))


if __name__ == "__main__":
    rebuild_section()
