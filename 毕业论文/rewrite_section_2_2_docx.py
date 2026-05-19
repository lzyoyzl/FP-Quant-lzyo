from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
SECTION_TITLE = "2.2　NVFP4 与 MXFP4 微缩放量化格式"
NEXT_SECTION_TITLE = "2.3　旋转辅助量化原理"

INTRO = (
    "与传统整数量化或单尺度 FP4 量化不同，microscaling 格式通过“低位宽元素编码 + 局部共享尺度”的组合方式，"
    "在极低位宽条件下兼顾表示密度与动态范围适配能力。"
    "就本文关注的两类典型 FP4 格式而言，NVFP4 与 MXFP4 都以 FP4 元素编码为主体，"
    "但在共享尺度的编码方式、层级结构以及 group 粒度上存在明显差异。"
    "这些差异不仅决定了两种格式的数值行为，也直接影响后续旋转设计所需面对的误差模式与部署开销。"
)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    sub_node = OxmlElement("m:sub")
    append_expr(sub_node, sub)
    sup_node = OxmlElement("m:sup")
    append_expr(sup_node, sup)
    node.append(e)
    node.append(sub_node)
    node.append(sup_node)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_table_before(anchor, doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    anchor._p.addprevious(table._tbl)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def rebuild_section():
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == SECTION_TITLE:
            start_idx = idx
        elif start_idx is not None and text == NEXT_SECTION_TITLE:
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate section 2.2 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]

    insert_text_paragraph_before(anchor, INTRO, "Normal")

    insert_heading_before(anchor, "2.2.1　NVFP4 量化机制", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "NVFP4 是面向 Blackwell 相关推理与训练栈提出的一种分层微缩放 FP4 格式。"
            "其元素编码采用 ",
            math_segment(["E2M1"]),
            " 形式的 FP4，即 1 位符号位、2 位指数位和 1 位尾数位；"
            "在此基础上，每 ",
            math_segment([m_sub("k", "NV"), " = 16"]),
            " 个连续元素构成一个 micro-block，"
            "共享一个以 ",
            math_segment(["E4M3"]),
            " 编码的局部尺度 ",
            math_segment([m_subsup("s", "g", "blk")]),
            "，同时整个张量还引入一个更高精度的全局尺度 ",
            math_segment([m_sup("s", "glb")]),
            "。"
            "这种“局部尺度 + 全局尺度”的两级结构，使 NVFP4 在保持 4 bit 主体编码的同时，"
            "仍能够对局部动态范围提供较细粒度的补偿[10-12]。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若记第 ",
            math_segment(["g"]),
            " 个 block 中第 ",
            math_segment(["i"]),
            " 个元素的 FP4 码值为 ",
            math_segment([m_subsup("q", "g,i", "nv")]),
            "，则 NVFP4 的数值恢复关系可写为式（2-17）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("x̃", "g,i"),
            " = ",
            m_subsup("q", "g,i", "nv"),
            "·",
            m_subsup("s", "g", "blk"),
            "·",
            m_sup("s", "glb"),
            ",   ",
            m_subsup("q", "g,i", "nv"),
            " ∈ ",
            m_sup("FP4", "E2M1"),
        ],
        "（2-17）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "在实际量化时，通常先用全局尺度 ",
            math_segment([m_sup("s", "glb")]),
            " 对整张量进行粗粒度归一化，再在 block 内依据局部最大幅值确定 ",
            math_segment([m_subsup("s", "g", "blk")]),
            "。"
            "若记第 ",
            math_segment(["g"]),
            " 个 block 的局部最大幅值为 ",
            math_segment([m_sub("a", "g"), " = max |", m_sub("x", "g,i"), "|"]),
            "，则一种常见的局部尺度估计形式如式（2-18）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("s", "g", "blk"),
            " = ",
            m_sub("Q", "E4M3"),
            "(",
            m_frac(
                m_sub("a", "g"),
                [m_sup("s", "glb"), "·", m_subsup("q", "max", "E2M1")],
            ),
            ")",
        ],
        "（2-18）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "据此，NVFP4 的量化与反量化过程可统一表示为式（2-19）。"
            "其中，",
            math_segment([m_sub("Q", "E2M1"), "(·)"]),
            " 表示映射到 FP4 ",
            math_segment(["E2M1"]),
            " 格点集合上的量化算子。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("q", "g,i", "nv"),
            " = ",
            m_sub("Q", "E2M1"),
            "(",
            m_frac(
                m_sub("x", "g,i"),
                [m_subsup("s", "g", "blk"), "·", m_sup("s", "glb")],
            ),
            "),   ",
            m_sub("x̃", "g,i"),
            " = ",
            m_subsup("q", "g,i", "nv"),
            "·",
            m_subsup("s", "g", "blk"),
            "·",
            m_sup("s", "glb"),
        ],
        "（2-19）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "NVFP4 的核心优势在于其局部尺度粒度更细，且 block scale 采用 ",
            math_segment(["FP8 E4M3"]),
            " 编码，相较于仅允许 2 的幂缩放的方案，其尺度表征更为灵活。"
            "对于分布起伏明显、局部离群值频繁出现的大模型权重而言，"
            "较小的 ",
            math_segment(["group_size = 16"]),
            " 能更及时地跟随局部动态范围变化，从而降低同组元素之间的尺度挤压。"
            "但与之相对应，NVFP4 需要同时维护 block 级和 tensor 级两套尺度信息，"
            "在元数据与实现复杂度上也更高。",
        ],
        "Normal",
    )

    insert_heading_before(anchor, "2.2.2　MXFP4 量化机制", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "MXFP4 则对应 OCP Microscaling（MX）规范下的标准化 FP4 微缩放格式。"
            "与 NVFP4 相同，MXFP4 的元素编码同样采用 ",
            math_segment(["E2M1"]),
            " 形式的 FP4；"
            "但其 block 设计更加规整：每 ",
            math_segment([m_sub("k", "MX"), " = 32"]),
            " 个元素共享一个以 ",
            math_segment(["E8M0"]),
            " 编码的尺度。"
            "由于 ",
            math_segment(["E8M0"]),
            " 本质上只保留指数信息，因此对应的共享尺度通常可视作 2 的幂，"
            "这使得 MXFP4 在硬件实现上更容易通过移位或指数调节完成缩放，但也意味着其尺度分辨率相对更粗[8-10]。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若记第 ",
            math_segment(["g"]),
            " 个 block 的共享尺度为 ",
            math_segment([m_subsup("s", "g", "mx")]),
            "，则 MXFP4 的基本恢复关系可写为式（2-20）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("x̃", "g,i"),
            " = ",
            m_subsup("q", "g,i", "mx"),
            "·",
            m_subsup("s", "g", "mx"),
            ",   ",
            m_subsup("q", "g,i", "mx"),
            " ∈ ",
            m_sup("FP4", "E2M1"),
        ],
        "（2-20）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "由于 MXFP4 的共享尺度采用 ",
            math_segment(["E8M0"]),
            " 编码，常用做法是先根据 block 最大幅值确定尺度指数，"
            "再将其映射为对应的 2 的幂。若仍记局部最大幅值为 ",
            math_segment([m_sub("a", "g")]),
            "，则其尺度估计可写为式（2-21）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("s", "g", "mx"),
            " = ",
            m_sup("2", ["⌈log₂(", m_frac(m_sub("a", "g"), m_sub("q", "max")), ")⌉"]),
        ],
        "（2-21）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "相应地，MXFP4 的量化与反量化过程可表示为式（2-22）。"
            "与 NVFP4 相比，这里不再显式引入额外的全局尺度层级，"
            "因此量化路径更短，但局部尺度的表达自由度也更受限制。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_subsup("q", "g,i", "mx"),
            " = ",
            m_sub("Q", "E2M1"),
            "(",
            m_frac(m_sub("x", "g,i"), m_subsup("s", "g", "mx")),
            "),   ",
            m_sub("x̃", "g,i"),
            " = ",
            m_subsup("q", "g,i", "mx"),
            "·",
            m_subsup("s", "g", "mx"),
        ],
        "（2-22）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "从工程角度看，MXFP4 的优势在于结构更规整、尺度元数据更少，且 ",
            math_segment(["E8M0"]),
            " 形式的共享尺度更利于硬件流水线实现。"
            "然而，",
            math_segment(["group_size = 32"]),
            " 也意味着更多元素需要共同适配同一尺度；"
            "当 block 内同时包含大值与小值时，中小值可分辨格点容易被显著压缩。"
            "因此，MXFP4 更依赖 block 内统计的一致性，若局部尺度结构不稳定，则量化误差往往会更快上升。",
        ],
        "Normal",
    )

    insert_heading_before(anchor, "2.2.3　两种格式的差异分析", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "综合来看，NVFP4 与 MXFP4 虽然都属于 microscaling FP4 范式，"
            "但二者在设计空间上的差异主要体现在三个方面。"
            "首先，NVFP4 采用“局部 FP8 block scale + 全局高精度 scale”的两级结构，"
            "而 MXFP4 采用“单级 E8M0 共享 scale”的更简洁结构；"
            "其次，NVFP4 的 ",
            math_segment(["group_size"]),
            " 更小，因而局部适配更细，MXFP4 的 block 更大，更强调块级统一性；"
            "再次，NVFP4 的局部尺度编码精度更高，而 MXFP4 的尺度更偏向硬件友好的幂次缩放。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若从平均元数据开销角度估计两种格式的单元素成本，则可写为式（2-23）。"
            "其中，",
            math_segment([m_sub("N", "T")]),
            " 表示共享同一全局尺度的张量元素总数。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_sub("b", "NV"),
            " ≈ 4 + ",
            m_frac("8", "16"),
            " + ",
            m_frac("32", m_sub("N", "T")),
            ",   ",
            m_sub("b", "MX"),
            " = 4 + ",
            m_frac("8", "32"),
        ],
        "（2-23）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "式（2-23）说明，MXFP4 在元数据摊销上通常更为节省，"
            "而 NVFP4 则以更高的尺度描述能力换取更细粒度的局部适配。"
            "这也决定了两者在精度与开销上的典型取舍：在相同权重量化位宽下，"
            "NVFP4 往往更容易保持局部动态范围，对大模型中的重尾分布和块内尺度波动更为稳健；"
            "而 MXFP4 则更具实现规整性和硬件友好性，但对 block 内统计一致性的要求更高。",
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "对本文后续研究而言，这种差异具有直接的方法学含义。"
            "若某种旋转策略能够有效削弱局部峰值并同时保持 block 内统计稳定，"
            "则其在 NVFP4 和 MXFP4 下都可能带来收益；"
            "但若某种变换虽然降低了全局离群值，却打散了原有局部结构，"
            "则其在 NVFP4 与 MXFP4 下的收益幅度很可能并不一致。"
            "因此，后续旋转算法设计不能脱离具体格式讨论，而应将格式结构本身视为搜索与评估的重要约束。",
        ],
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "表2-2 从元素编码、尺度设计、粒度、元数据开销和部署特性等方面，总结了 NVFP4 与 MXFP4 的主要差异。",
        "Normal",
    )
    insert_caption_before(anchor, "表2-2　NVFP4 与 MXFP4 的格式差异比较")
    insert_table_before(
        anchor,
        doc,
        [
            ["比较维度", "NVFP4", "MXFP4"],
            ["元素编码", "FP4 E2M1", "FP4 E2M1"],
            ["共享尺度编码", "局部 FP8 E4M3 + 全局 FP32", "局部 E8M0（幂次尺度）"],
            ["group_size", "16", "32"],
            ["尺度层级", "两级", "一级"],
            ["单元素摊销元数据", "约 0.5 bit + 全局尺度摊销", "约 0.25 bit"],
            ["数值适配能力", "更强，利于局部动态范围拟合", "较弱，更依赖块内统计一致性"],
            ["实现与部署", "精度更优但元数据和实现更复杂", "结构规整、硬件友好、开销更低"],
        ],
    )

    doc.save(str(DOC_PATH))


if __name__ == "__main__":
    rebuild_section()
