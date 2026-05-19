from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
TEMP_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文_2.3更新待替换.docx")
SECTION_TITLE = "2.3　旋转辅助量化原理"
NEXT_SECTION_TITLE = "2.5　本章小结"

INTRO = (
    "旋转辅助量化的核心思想，是在不改变原始线性映射的前提下，"
    "通过适当的线性变换重塑输入表示或权重表示在坐标空间中的分布形态。"
    "这类方法并不直接修改模型所表达的函数，而是改变量化算子所面对的数值分布，"
    "从而缓解低比特条件下由 outlier、重尾分布和共享尺度失配引起的精度退化。"
    "因此，旋转既是一种分布整形手段，也是一种兼顾理论等价性与工程可部署性的量化增强机制。"
)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def math_segment(expr):
    return {"type": "math", "expr": expr}


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_mixed_paragraph_before(anchor, segments, style: str):
    para = anchor.insert_paragraph_before("")
    para.style = style
    clear_paragraph(para)
    for seg in segments:
        if isinstance(seg, str):
            para._p.append(build_w_run(seg))
        else:
            para._p.append(make_inline_math(seg["expr"]))
    return para


def insert_text_paragraph_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_heading_before(anchor, text: str, style: str):
    para = anchor.insert_paragraph_before(text)
    para.style = style
    return para


def insert_caption_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Caption"
    return para


def insert_center_placeholder_before(anchor, text: str):
    para = anchor.insert_paragraph_before(text)
    para.style = "Normal"
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def insert_table_before(anchor, doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    anchor._p.addprevious(table._tbl)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def insert_equation_table_before(anchor, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    anchor._p.addprevious(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))


def remove_tables_between(start_p, next_p) -> None:
    current = start_p.getnext()
    tables_to_remove = []
    while current is not None and current is not next_p:
        if current.tag == qn("w:tbl"):
            tables_to_remove.append(current)
        current = current.getnext()
    for tbl in tables_to_remove:
        tbl.getparent().remove(tbl)


def rebuild_section():
    doc = Document(str(DOC_PATH))
    paragraphs = doc.paragraphs

    start_idx = None
    next_idx = None
    for idx, para in enumerate(paragraphs):
        text = para.text.strip()
        if text == SECTION_TITLE:
            start_idx = idx
        elif start_idx is not None and text == NEXT_SECTION_TITLE:
            next_idx = idx
            break

    if start_idx is None or next_idx is None or next_idx <= start_idx:
        raise RuntimeError("Failed to locate section 2.3 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    next_p = doc.paragraphs[next_idx]._p
    remove_tables_between(start_p, next_p)

    for idx in range(next_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]

    insert_text_paragraph_before(anchor, INTRO, "Normal")

    insert_heading_before(anchor, "2.3.1　旋转辅助量化的基本原理", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "设某线性层输入为 ",
            math_segment(["x"]),
            "，权重矩阵为 ",
            math_segment(["W"]),
            "，输出为 ",
            math_segment(["y"]),
            "，则其原始映射关系可表示为式（2-24）。",
        ],
        "Normal",
    )
    insert_equation_table_before(anchor, ["y = xW"], "（2-24）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "在经典旋转量化中，常用的变换矩阵为正交矩阵 ",
            math_segment(["R"]),
            "，其满足式（2-25）中的正交条件。",
        ],
        "Normal",
    )
    insert_equation_table_before(anchor, ["RᵀR = RRᵀ = I"], "（2-25）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "此时原始映射可改写为式（2-26）。"
            "由于正交变换保持范数与内积结构，因此它是最常见、也最容易控制数值稳定性的旋转形式。",
        ],
        "Normal",
    )
    insert_equation_table_before(anchor, ["y = xW = (xR)(RᵀW)"], "（2-26）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "不过，从线性映射等价性的角度看，旋转矩阵并不一定要严格限制为正交阵。"
            "只要变换矩阵 ",
            math_segment(["R"]),
            " 是可逆的，即存在 ",
            math_segment(["R⁻¹"]),
            "，就可以通过对激活施加前向变换、对权重折叠逆变换来保持映射等价。"
            "相应的可逆性条件可写为式（2-27）。",
        ],
        "Normal",
    )
    insert_equation_table_before(anchor, ["det(R) ≠ 0"], "（2-27）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "在更一般的情形下，原始映射可表示为式（2-28）。"
            "可见，正交变换只是可逆线性变换的一种特殊情形；"
            "当 ",
            math_segment(["R⁻¹ = Rᵀ"]),
            " 时，式（2-28）便退化为式（2-26）。",
        ],
        "Normal",
    )
    insert_equation_table_before(anchor, ["y = xW = (xR)(R⁻¹W)"], "（2-28）", doc)

    insert_mixed_paragraph_before(
        anchor,
        [
            "从工程实现角度看，式（2-28）对应“激活侧前向变换 + 权重侧逆变换折叠”的部署方式。"
            "在量化场景中，之所以仍然更偏好正交或近似正交变换，"
            "主要是因为这类变换通常具有更好的数值稳定性、能量保持性与硬件实现友好性；"
            "但从理论上说，只要可逆性成立，广义旋转或一般可逆线性变换同样可以作为量化前的分布整形工具。"
        ],
        "Normal",
    )

    insert_heading_before(anchor, "2.3.2　旋转改善量化误差的作用机理", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "在共享尺度的低比特量化中，outlier 之所以会显著放大量化误差，"
            "根本原因在于其会主导同一 group 的尺度估计。"
            "若记第 ",
            math_segment(["g"]),
            " 个 group 中元素为 ",
            math_segment([m_sub("x", "g,1"), ", … , ", m_sub("x", "g,m")]),
            "，则共享尺度的一种典型估计方式可写为式（2-29）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [m_sub("s", "g"), " = ", m_frac(["max |", m_sub("x", "g,i"), "|"], m_sub("q", "max"))],
        "（2-29）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若采用就近舍入且暂不考虑截断，则组内第 ",
            math_segment(["i"]),
            " 个元素的绝对误差上界满足式（2-30）。",
            "该式表明，所有元素共享同一个尺度时，误差上界直接由组内最大幅值决定。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["|", m_sub("e", "g,i"), "| = |", m_sub("x̃", "g,i"), " - ", m_sub("x", "g,i"), "| ≤ ", m_frac(m_sub("s", "g"), "2")],
        "（2-30）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "若记该 group 中的 outlier 幅值为 ",
            math_segment([m_sub("M", "g"), " = max |", m_sub("x", "g,i"), "|"]),
            "，则对非 outlier 元素而言，其相对误差上界可进一步写为式（2-31）。"
            "当普通元素幅值显著小于 ",
            math_segment([m_sub("M", "g")]),
            " 时，该上界会迅速增大，这正是 outlier 挤压有效分辨率的直接体现。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [
            m_frac(["|", m_sub("e", "g,i"), "|"], ["|", m_sub("x", "g,i"), "|"]),
            " ≤ ",
            m_frac(m_sub("M", "g"), ["2", m_sub("q", "max"), "|", m_sub("x", "g,i"), "|"]),
        ],
        "（2-31）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "旋转缓解误差的作用机理，就在于它能够在不改变总体映射关系的前提下重新分配能量。"
            "若记旋转后的表示为 ",
            math_segment(["x'", m_sub("", "g"), " = x", m_sub("", "g"), "R"]),
            "，则对应的共享尺度变为式（2-32）。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        [m_sub("s'", "g"), " = ", m_frac(["max |", m_sub("x'", "g,i"), "|"], m_sub("q", "max"))],
        "（2-32）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "当所选旋转能够削弱峰值主导、使旋转后最大幅值下降，即 ",
            math_segment([m_sub("s'", "g"), " < ", m_sub("s", "g")]),
            " 时，组内误差上界将相应收缩，如式（2-33）所示。",
        ],
        "Normal",
    )
    insert_equation_table_before(
        anchor,
        ["|", m_sub("e'", "g,i"), "| ≤ ", m_frac(m_sub("s'", "g"), "2"), " < ", m_frac(m_sub("s", "g"), "2")],
        "（2-33）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "因此，旋转并不是直接“消除” outlier，而是通过扩散局部尖峰、降低块内峰值与普通值之间的失衡程度，"
            "来改善共享尺度的匹配质量。"
            "当然，这种收益并非无条件成立：若旋转虽然降低了某一局部峰值，却同时破坏了原有块级结构，"
            "使高幅值信息跨 group 扩散，则也可能削弱 microscaling 格式依赖的局部尺度隔离优势。"
            "因此，旋转效果最终取决于“峰值抑制能力”与“局部结构保持能力”之间的平衡。"
        ],
        "Normal",
    )

    insert_heading_before(anchor, "2.3.3　常见旋转量化方法及适用场景", "Heading 3")

    insert_mixed_paragraph_before(
        anchor,
        [
            "从方法形式看，旋转辅助量化并非单一算法，而是一类围绕“候选变换如何构造与选择”的方法集合。"
            "按照变换来源与代价特征，可将常见方法概括为三类："
            "零旋转基线、固定结构旋转，以及可学习或结构增强旋转。"
            "其中，前者主要用于建立对照，后两者分别对应低成本部署和更强分布适配两种典型路线。"
        ],
        "Normal",
    )

    insert_center_placeholder_before(anchor, "【插图占位：此处插入常见旋转量化方法分类与代表算法示意图】")
    insert_caption_before(anchor, "图2-3　常见旋转量化方法分类与代表算法示意（占位）")

    insert_mixed_paragraph_before(
        anchor,
        [
            "其中，",
            math_segment(["identity"]),
            " 可视为零成本基线，不会引入额外数值扰动，但也无法主动改善 outlier 分布。"
            "Hadamard 类结构化正交变换则具有实现简单、计算开销低和硬件友好等优点，"
            "因此适合作为固定候选。QuaRot[5] 即属于这一类方法：其利用固定旋转削弱隐藏状态中的 outlier，"
            "从而使权重、激活乃至 KV cache 的 4 bit 量化更易实现，适合强调端到端低比特部署的场景。"
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "相比之下，SpinQuant[6] 引入可学习旋转，利用少量校准数据直接优化旋转矩阵。"
            "这类方法能够更充分地贴合模型分布和量化目标，在难量化模型或极低位宽场景下通常具有更高精度；"
            "但相应地，也需要额外的优化过程与更高的搜索成本。"
            "因此，它更适合对精度较为敏感、且允许一定离线优化开销的应用场景。"
        ],
        "Normal",
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "另一类代表方法是结构增强型训练自由旋转。GSR[7] 在 Walsh-Hadamard 变换基础上引入 sequency 排列和分组块对角结构，"
            "通过更细粒度地隔离 outlier 影响，提升极低位宽下的量化稳定性。"
            "这类方法不依赖显式训练或复杂学习过程，但通过改进旋转结构本身获得比标准 Hadamard 更好的局部误差行为，"
            "适合希望兼顾较低部署成本与更强局部适配能力的场景。"
        ],
        "Normal",
    )

    insert_text_paragraph_before(anchor, "表2-3 对常见旋转量化方法的主要性质、适用场景与典型局限进行了归纳。", "Normal")
    insert_caption_before(anchor, "表2-3　常见旋转量化方法的性质比较")
    insert_table_before(
        anchor,
        doc,
        [
            ["方法类别", "代表形式", "主要性质", "适用场景", "典型局限"],
            ["零旋转基线", "Identity", "无额外计算与参数开销，可作为是否需要旋转的参照", "需要建立对比基线或某些层旋转收益不确定时", "无法主动改善分布"],
            ["固定结构旋转", "Hadamard、QuaRot", "结构规则、实现简单、硬件友好、训练自由", "强调部署效率、希望低成本覆盖大规模层时", "适配性有限，难以针对特定层精细调整"],
            ["可学习旋转", "SpinQuant", "与数据分布适配更强，精度上限较高", "难量化模型、极低位宽、精度敏感场景", "需要额外优化过程与更高搜索成本"],
            ["结构增强旋转", "GSR", "保持训练自由特征，同时改进固定旋转的局部误差行为", "希望兼顾低开销与更强局部适配能力时", "结构设计较复杂，收益依赖具体分组机制"],
        ],
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "总体而言，旋转量化方法之间并不存在对所有层、所有格式都普适最优的单一选择。"
            "固定结构旋转的优势在于成本低、便于部署；可学习旋转的优势在于适配性强；结构增强旋转则提供了介于二者之间的折中。"
            "因此，在后续方法设计中，更合理的做法不是预设某一种旋转必然优于其他候选，"
            "而是将多种可逆变换统一纳入候选集合，再结合量化格式、层角色与误差模式进行针对性选择。"
        ],
        "Normal",
    )

    try:
        doc.save(str(DOC_PATH))
        print(f"SAVED:{DOC_PATH}")
    except PermissionError:
        doc.save(str(TEMP_PATH))
        print(f"TEMP_SAVED:{TEMP_PATH}")


if __name__ == "__main__":
    rebuild_section()
