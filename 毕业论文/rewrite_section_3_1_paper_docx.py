from pathlib import Path

from docx import Document


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_TITLE = "3.1 全层旋转与 group 粒度分块旋转的影响对比"
END_TITLE = "3.2 局部列块分布统计与选择性旋转动机"


def find_paragraph(doc: Document, text: str):
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Paragraph not found: {text}")


def delete_between(start_para, end_para):
    body = start_para._parent._element
    children = list(body.iterchildren())
    start_idx = children.index(start_para._element)
    end_idx = children.index(end_para._element)
    for child in children[start_idx + 1:end_idx]:
        body.remove(child)


def collect_styles(doc: Document, start_para, end_para):
    body_style = None
    fig_style = None
    table_style = None
    passed = False
    for para in doc.paragraphs:
        if para is start_para:
            passed = True
            continue
        if para is end_para:
            break
        text = para.text.strip()
        if not text:
            continue
        if body_style is None and not text.startswith("图") and not text.startswith("表") and not text.startswith("【"):
            body_style = para.style
        if fig_style is None and text.startswith("图"):
            fig_style = para.style
        if table_style is None and text.startswith("表"):
            table_style = para.style
    if body_style is None:
        for para in doc.paragraphs:
            if para.text.strip() == END_TITLE:
                break
            if para.text.strip():
                body_style = para.style
                break
    if fig_style is None:
        for para in doc.paragraphs:
            if para.text.strip().startswith("图"):
                fig_style = para.style
                break
    if table_style is None:
        for para in doc.paragraphs:
            if para.text.strip().startswith("表"):
                table_style = para.style
                break
    return body_style, fig_style or body_style, table_style or body_style


def insert_paragraph_before(anchor_para, text, style):
    para = anchor_para.insert_paragraph_before(text)
    if style is not None:
        para.style = style
    return para


def main():
    doc = Document(str(DOC_PATH))
    start_para = find_paragraph(doc, START_TITLE)
    end_para = find_paragraph(doc, END_TITLE)

    body_style, fig_style, table_style = collect_styles(doc, start_para, end_para)
    delete_between(start_para, end_para)

    blocks = [
        (
            "已有研究已经较为清晰地说明了，在 MXFP4 这类以 group 为共享尺度单元的微缩放格式中，"
            "直接沿用全层统一旋转并不合适。Shao 等在《Block Rotation is All You Need for MXFP4 "
            "Quantization》中对多种量化格式进行了系统比较，指出随机 Hadamard 等全局旋转在 "
            "INT4 场景中通常能够带来收益，但在 BINT4、BFP4、MXINT4 尤其是 MXFP4 这类 "
            "group-wise 格式下却往往导致性能下降。由此可见，旋转策略的有效性与量化格式的尺度"
            "组织方式密切相关，不能简单将面向 INT4 的经验直接迁移到 MXFP4 场景。"
        ),
        (
            "该文进一步指出，问题的关键并不在于旋转本身完全失效，而在于全层旋转会将原本集中在"
            "少数通道上的离群能量重新分配到全层范围内。对于 MXFP4 而言，各个 group 独立共享 "
            "PoT 标度，跨 group 的能量扩散会使原本较为平稳的 regular block 也出现更多中等偏"
            "大的数值，从而抬升局部 scale 并放大量化误差。换言之，全层旋转缓解了极少数最强离群"
            "值，却可能同时污染大量原本正常的量化块。"
        ),
        (
            "从该文图 5 给出的统计结果可以看到，全层旋转虽然显著削弱了极端大值，但并没有消除总"
            "体能量，而是把一部分大值转化为更大范围内的中等偏大值。具体而言，阈值为 3 的极端激"
            "活占比明显下降，但阈值在 1.5 附近时的数据占比却由约 5% 上升到约 11%。这一现象说"
            "明，全层旋转在削弱少数 outlier 的同时，也扩大了受影响通道的覆盖范围。"
        ),
        "【插图占位：建议插入 Shao 等《Block Rotation is All You Need for MXFP4 Quantization》中的 Figure 7】",
        "图3-1 全层旋转与分块旋转在 MXFP4 中的作用差异示意（引自 Shao 等《Block Rotation is All You Need for MXFP4 Quantization》Figure 7）",
        (
            "更重要的是，该文图 6 表明，全层旋转之后 regular blocks 的平均量化损失会明显升高。"
            "由于实际模型中 regular blocks 的数量远多于 outlier blocks，这部分误差在累积后会主"
            "导总体精度退化。因此，在 MXFP4 场景下，问题并不只是“是否削弱了最大的离群值”，而"
            "是“是否把原本局部化的异常能量扩散到了大量普通 group 中”。"
        ),
        (
            "基于上述分析，Shao 等提出采用与量化 block 对齐的分块旋转策略。其基本思想是在每个"
            "局部 block 内独立进行旋转，使数值重分配被限制在 group 内部，而不是扩散到整个层空"
            "间。这样做一方面能够在局部范围内缓解 outlier 对共享尺度的主导，另一方面又能避免跨 "
            "group 污染，从而保持 regular blocks 的尺度稳定性。就本文关注的问题而言，这一结论直"
            "接说明了 group 粒度分块旋转的必要性。"
        ),
        "表3-1 全层旋转与分块旋转在 MXFP4 中的关键现象比较（根据 Shao 等《Block Rotation is All You Need for MXFP4 Quantization》中的 Figure 2、Figure 5 和 Figure 6 整理）",
        "【表格占位：此处可整理文献中关于全层旋转导致性能下降、较大值扩散以及 regular blocks 误差增加的关键现象比较表】",
        (
            "因此，本节更关注“旋转粒度是否合适”这一前提问题，而不是直接讨论具体旋转矩阵的设"
            "计。已有研究已经表明，在 MXFP4 这类 group-wise 微缩放格式下，应优先采用与量化 group"
            " 对齐的局部分块旋转。至于哪些局部列块真正需要旋转、是否应对不同 group 使用不同变"
            "换，则还需要结合局部数据分布与量化误差来源继续分析，这也正是下一节展开讨论的重点。"
        ),
    ]

    for block in blocks:
        if block.startswith("图3-1"):
            style = fig_style
        elif block.startswith("表3-1"):
            style = table_style
        else:
            style = body_style
        insert_paragraph_before(end_para, block, style)

    doc.save(str(DOC_PATH))
    print(f"SAVED:{DOC_PATH}")


if __name__ == "__main__":
    main()
