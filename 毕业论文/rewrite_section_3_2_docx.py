from pathlib import Path

from docx import Document


DOC_PATH = Path(r"\\wsl.localhost\Ubuntu-22.04\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_TITLE = "3.2 局部列块分布统计与选择性旋转"
END_TITLE = "3.3 group 粒度下的量化误差来源分析"


def find_paragraph(doc: Document, text: str):
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Paragraph not found: {text}")


def delete_between(start_para, end_para):
    body = start_para._parent._element
    children = list(body.iterchildren())
    start_idx = children.index(start_para._element)
    end_idx = children.index(end_para._element)
    for child in children[start_idx + 1:end_idx]:
        body.remove(child)


def collect_styles(doc: Document, start_para, end_para):
    body_style = None
    fig_style = None
    for para in doc.paragraphs:
        if para is start_para:
            continue
        if para is end_para:
            break
        text = para.text.strip()
        if not text:
            continue
        if body_style is None and not text.startswith("图") and not text.startswith("【"):
            body_style = para.style
        if fig_style is None and text.startswith("图"):
            fig_style = para.style
    if body_style is None:
        body_style = start_para.style
    if fig_style is None:
        fig_style = body_style
    return body_style, fig_style


def insert_paragraph_before(anchor_para, text, style):
    para = anchor_para.insert_paragraph_before(text)
    if style is not None:
        para.style = style
    return para


def main():
    doc = Document(str(DOC_PATH))
    start_para = find_paragraph(doc, START_TITLE)
    end_para = find_paragraph(doc, END_TITLE)
    body_style, fig_style = collect_styles(doc, start_para, end_para)
    delete_between(start_para, end_para)

    blocks = [
        (
            "为进一步判断哪些局部列块值得旋转，本文选取 llama-2-7b-hf、llama-3-8b-hf 和 Qwen3-8B "
            "作为统计对象，在多个抽样 Transformer block 上对 qkv、o、gate_up 和 down 槽位输入的 "
            "group-wise 激活分布进行统计。统计时分别采用 group size 为 16 和 32 的两种粒度，以对应更细和更粗的局部列块划分方式，"
            "从而观察不同 group 粒度下局部异常是否具有一致的分布规律。"
        ),
        (
            "统计结果表明，强 outlier 富集 group 在三类模型的不同抽样 block 中均可观测到。无论在 group size=16 "
            "还是 group size=32 的设置下，图中“强 outlier 富集 group 比例”始终保持非零，且不同模型在多个抽样 block 上都出现了"
            "相近的长尾现象。这说明局部异常并非只局限于个别模型或少数特殊层，而是广泛存在于大语言模型的局部激活表示之中。"
        ),
        "【插图占位：此处插入 group size=16 的抽样层 group-wise 激活分布统计综合图】",
        "【插图占位：此处插入 group size=32 的抽样层 group-wise 激活分布统计综合图】",
        "【插图占位：此处插入 group size=16 的槽位平均比例统计图】",
        "【插图占位：此处插入 group size=32 的槽位平均比例统计图】",
        "图3-3 不同 group 粒度下的局部列块激活分布统计结果（占位）",
        (
            "但与此同时，统计图也显示出另一条更重要的信息：大多数 group 列块的局部分布其实相对均匀。首先，从 outlier-density "
            "multiplier 的 CDF 曲线可以看出，三类模型的曲线都在靠近 1 的区域快速上升，并在较小的 multiplier 范围内迅速逼近 1，"
            "说明绝大多数 group 的 outlier 密度并未显著高于全局基线。其次，在代表性 qkv 槽位的排序曲线中，只有极少数排名靠前的 group "
            "呈现出明显偏高的 outlier 密度，而其余绝大部分 group 的曲线很快回落到接近基线甚至低于基线的区域。这表明局部异常主要集中于少数"
            "列块，而不是在整层范围内普遍扩散。"
        ),
        (
            "这一结论还可由“低尾 group 比例”进一步印证。无论 group size 为 16 还是 32，三类模型在多数抽样 block 中的低尾 "
            "group 比例都保持在较高水平，通常处于约 0.75 到 0.90 的区间；分槽位的均值统计也表明，qkv、o、gate_up 和 down "
            "等不同槽位虽然存在一定差异，但其大多数 group 仍然落在相对平滑、低尾的分布区域。换言之，局部异常的确广泛存在，但真正需要重点处理的"
            "列块只占少数，常规列块仍然构成了层内激活分布的主体。"
        ),
        (
            "同时还应注意到，不同槽位之间的统计形态并不完全一致。例如，qkv 和 gate_up 槽位中强 outlier 富集 group 的平均比例相对"
            "更高，而 down 槽位的低尾比例相对更低，说明不同结构功能的线性层在局部激活分布上具有差异性。这意味着，即便都采用 group 粒度建模，"
            "也不能简单假定所有槽位都适合同一种处理方式。"
        ),
        (
            "因此，局部列块分布统计给出的并不是“是否需要旋转”的简单二元结论，而是更细致的选择性启示：一方面，outlier 在局部 group 中的广泛"
            "存在说明旋转具有现实必要性；另一方面，大多数列块的相对均匀性又进一步说明旋转不能采取“一刀切”的方式。更合理的思路应当是以 group 为"
            "粒度，在局部分布统计基础上有选择地对确有必要的列块施加旋转，而对于原本已经较为平滑的多数 group，则应尽量避免无差别变换对其局部结构造成"
            "扰动。后续的 group-wise 旋转算法设计，正是建立在这一统计观察基础之上的。"
        ),
    ]

    for text in blocks:
        style = fig_style if text.startswith("图3-3") else body_style
        insert_paragraph_before(end_para, text, style)

    doc.save(str(DOC_PATH))
    print(f"SAVED:{DOC_PATH}")


if __name__ == "__main__":
    main()
