import sys
from pathlib import Path

from docx import Document

from rewrite_chapter4_docx import (
    delete_paragraph,
    insert_caption_before,
    insert_equation_table_before,
    insert_mixed_paragraph_before,
    insert_text_paragraph_before,
    m_sub,
    m_subsup,
    m_sup,
    remove_tables_between,
)


DEFAULT_DOC_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_TITLE = "4.1 问题定义与总体思路"
END_TITLE = "4.2 group-wise 旋转搜索空间设计"


def rebuild_section_4_1(doc: Document, anchor):
    insert_text_paragraph_before(
        anchor,
        "在完成第三章对局部列块分布统计与量化误差来源的分析之后，本章进一步转入 group-wise 旋转算法的具体实现。"
        "本节主要关注如何在低比特量化约束下，围绕 Transformer block 内不同输入槽位的 group 列块，从候选变换集合中搜索更合适的局部旋转矩阵，"
        "并将搜索结果贯穿量化、导出与推理部署全流程。相较于整层统一施加单一旋转，本实现将搜索粒度固定在 activation/weight group 上，"
        "允许同一槽位内不同列块采用不同变换。",
        "Normal",
    )

    insert_equation_table_before(
        anchor,
        [
            m_sub("x", "g"), "'", " = ", m_sub("x", "g"), " ", m_sub("T", "g"),
            ",    ",
            m_sub("W", "g"), "'", " = ", m_sub("W", "g"), " ", m_sup(m_sub("T", "g"), "-T"),
        ],
        "（4-1）",
        doc,
    )

    insert_text_paragraph_before(
        anchor,
        "式（4-1）给出了实现中统一采用的等价变换关系。前向阶段对激活侧施加 group-wise 输入变换，"
        "权重侧则在量化前折叠对应的逆转置矩阵，从而保持线性映射不变。代码层面并未对每个线性层完全独立搜索，"
        "而是按照功能相近的输入槽位共享搜索，这样既保留了局部搜索的灵活性，也控制了实际搜索与导出的复杂度。",
        "Normal",
    )

    insert_equation_table_before(
        anchor,
        [
            m_subsup("T", "g", "*"), " = arg min_{T ∈ C} ", m_sub("S", "g"), "(", m_sub("T", "g"), ")",
            ",    C = {", m_sub("T", "id"), ", ", m_sub("T", "had"), ", ", m_sub("T", "dct"),
            ", ", m_sub("T", "dst"), ", ", m_sub("T", "gsr"), ", ", m_sub("T", "hh"), "}",
            ",    ", m_sub("S", "g"), " ∈ {",
            m_subsup("L", "mse", "(g)"), ", ",
            m_subsup("L", "cov", "(g)"), ", ",
            m_subsup("L", "act", "(g)"), ", ",
            m_subsup("J", "tail", "(g)"),
            "}",
        ],
        "（4-2）",
        doc,
    )

    insert_mixed_paragraph_before(
        anchor,
        [
            "对每个 group 列块，搜索过程都可写成式（4-2）的形式。其中，候选集合 ",
            "C",
            " 由当前实现中实际参与搜索的局部变换矩阵构成，默认包括 Identity、Hadamard、DCT、DST、GSR 与 Householder；评分函数 ",
            m_sub("S", "g"),
            "(·) 则由量化目标决定，可取基于权重重建误差的 ",
            m_subsup("L", "mse", "(g)"),
            "、基于输入协方差加权的 ",
            m_subsup("L", "cov", "(g)"),
            "、基于激活侧量化误差的 ",
            m_subsup("L", "act", "(g)"),
            "，以及在基础项之外进一步引入尾部误差约束的 ",
            m_subsup("J", "tail", "(g)"),
            "。也就是说，当前实现并不是在固定旋转下做被动适配，而是在“候选变换集合”和“评分函数集合”双重约束下，为每个局部列块选择更合适的变换方向。",
        ],
        "Normal",
    )

    insert_text_paragraph_before(
        anchor,
        "【插图占位：此处插入第4章总体流程图，展示校准数据收集、group-wise 旋转搜索、量化导出与推理接入的完整链路。】",
        "Normal",
    )
    insert_caption_before(anchor, "图4-1 第4章 group-wise 旋转算法总体实现流程图（占位）")


def main():
    doc_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_DOC_PATH
    if not doc_path.is_absolute():
        doc_path = (Path.cwd() / doc_path).resolve()
    doc = Document(str(doc_path))
    paragraphs = list(doc.paragraphs)

    start_idx = None
    end_idx = None
    for i, p in enumerate(paragraphs):
        text = p.text.strip()
        if text == START_TITLE:
            start_idx = i
        elif start_idx is not None and text == END_TITLE:
            end_idx = i
            break

    if start_idx is None or end_idx is None or end_idx <= start_idx:
        raise RuntimeError("Failed to locate section 4.1 boundaries.")

    start_p = doc.paragraphs[start_idx]._p
    end_p = doc.paragraphs[end_idx]._p
    remove_tables_between(start_p, end_p)

    for idx in range(end_idx - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx + 1]
    rebuild_section_4_1(doc, anchor)

    doc.save(str(doc_path))
    print(f"SAVED:{doc_path}")


if __name__ == "__main__":
    main()
