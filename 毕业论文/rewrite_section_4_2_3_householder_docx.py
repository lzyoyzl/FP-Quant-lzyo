from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.2.3 group-wise自适应Householder 变换实现"
END_HEADING = "4.2.4 Hadamard 类变换实现"


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def m_rad(expr):
    node = OxmlElement("m:rad")
    deg_hide = OxmlElement("m:degHide")
    deg_hide.set(qn("m:val"), "1")
    node.append(deg_hide)
    e = OxmlElement("m:e")
    append_expr(e, expr)
    node.append(e)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def make_inline_math(expr):
    node = OxmlElement("m:oMath")
    append_expr(node, expr)
    return node


def insert_paragraph_after(paragraph, text: str, style_name: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.add_run(text)
    return new_para


def insert_mixed_paragraph_after(paragraph, segments, style_name: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    clear_paragraph(new_para)
    for seg in segments:
        if isinstance(seg, str):
            new_para._p.append(build_w_run(seg))
        else:
            new_para._p.append(make_inline_math(seg))
    return new_para


def insert_equation_table_after(paragraph, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    paragraph._p.addnext(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))

    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def insert_caption_after(paragraph, text: str, style_name: str = "Caption"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style_name
    clear_paragraph(new_para)
    new_para._p.append(build_w_run(text))
    return new_para


def insert_table_after(paragraph, doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    paragraph._p.addnext(table._tbl)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def main():
    doc = Document(str(DOCX_PATH))

    start_idx = None
    end_idx = None
    start_para = None
    body_style = None

    for idx, para in enumerate(doc.paragraphs):
        if para.text.strip() == START_HEADING:
            start_idx = idx
            start_para = para
        elif para.text.strip() == END_HEADING and start_idx is not None:
            end_idx = idx
            break

    if start_idx is None or end_idx is None or start_para is None:
        raise RuntimeError("Unable to locate section 4.2.3 boundaries.")

    paragraphs = list(doc.paragraphs)
    for para in paragraphs[start_idx + 1:end_idx]:
        if para.style is not None and body_style is None and para.text.strip():
            body_style = para.style.name
        delete_paragraph(para)

    if body_style is None:
        body_style = "Normal"

    p1 = (
        "Householder 反射的基本形式由式（4-6）给出，且式（4-7）所表示的正交、自逆性质在当前实现中直接成立。"
        "在项目代码里，householder 被归入 ADAPTIVE_GROUP_TRANSFORMS，因此它并不与 Hadamard、DCT、DST、GSR 等候选一起在搜索开始前统一实例化，"
        "而是作为局部自适应候选，在逐 group 搜索过程中按当前列块的统计特征动态构造。"
    )
    cursor = insert_paragraph_after(start_para, p1, body_style)

    cursor = insert_mixed_paragraph_after(
        cursor,
        [
            "从具体实现看，当前 Householder 候选并不是先固定一个全局反射块，再对所有列块重复使用，"
            "而是对每个 group 单独构造 per-group data-adaptive Householder candidate。"
            "对第 ",
            m_sub("g", ""),
            " 个 group，代码先依据当前搜索目标自动选择局部数据源，提取能量向量 ",
            m_sub("e", "g"),
            "。当目标为 act_mse，或 jtail 且 tail_source = activation 时，",
            m_sub("e", "g"),
            " 由当前 group 激活样本 ",
            m_sub("X", "g"),
            " 的均方根统计给出；当目标为 cov 时，",
            m_sub("e", "g"),
            " 由当前 group 输入协方差 ",
            m_sub("Σ", "g"),
            " 对角线的均方根统计给出；其余情况下，",
            m_sub("e", "g"),
            " 由共享槽位权重列块 ",
            m_sub("W", "g"),
            " 的均方根统计给出。这样一来，Householder 反射方向并不是预设的，而是随当前 group 和当前目标函数共同变化的。"
        ],
        body_style,
    )

    cursor = insert_caption_after(cursor, "表4-3 Householder 变换中局部能量向量的数据源选择")
    cursor = insert_table_after(
        cursor,
        doc,
        [
            ["当前目标", "Householder 构造来源"],
            ["mse", "当前 group 的权重块"],
            ["cov", "当前 group 的输入协方差对角线"],
            ["act_mse", "当前 group 的激活样本"],
            ["jtail + tail_source=weight", "当前 group 的权重块"],
            ["jtail + tail_source=activation", "当前 group 的激活样本"],
        ],
    )

    cursor = insert_mixed_paragraph_after(
        cursor,
        [
            "为了更清晰地说明能量向量与反射方向的计算过程，可将当前实现分解为“数据源选择、方向归一化、反射矩阵生成”三个步骤。"
            "首先依据表4-3选定当前 group 的局部数据源，并计算能量向量 ",
            m_sub("e", "g"),
            "；随后将其归一化为局部主导方向 ",
            m_sub("u", "g"),
            "；最后以均匀方向 ",
            "t",
            " 作为参考，构造反射方向 ",
            m_sub("r", "g"),
            "。由此，Householder 变换不再是固定模板，而是由当前列块的局部统计直接诱导。"
        ],
        body_style,
    )

    cursor = insert_equation_table_after(
        cursor,
        [
            m_sub("e", "g"), " = source_rms,   ",
            m_sub("u", "g"), " = ", m_frac(m_sub("e", "g"), ["∥", m_sub("e", "g"), "∥", m_sub("", "2")]),
            ",   ",
            "t = ", m_frac("1", m_rad("k")), "1",
            ",   ",
            m_sub("r", "g"), " = ", m_sub("u", "g"), " - t",
        ],
        "（4-7a）",
        doc,
    )

    cursor = insert_mixed_paragraph_after(
        cursor,
        [
            "得到 ",
            m_sub("e", "g"),
            " 后，代码会先将其归一化为 source direction ",
            m_sub("u", "g"),
            "，并以均匀方向 ",
            "t",
            " 作为 target direction，其中 ",
            "k",
            " 表示当前 group 的维度，",
            "1",
            " 表示全 1 向量。随后由 ",
            m_sub("r", "g"),
            " = ",
            m_sub("u", "g"),
            " - t",
            " 构造反射方向；若当前局部能量已接近零向量，或 ",
            m_sub("u", "g"),
            " 与 ",
            "t",
            " 已近似一致，则直接返回单位阵 ",
            "I",
            "；否则按式（4-7b）构造局部 Householder 矩阵。"
        ],
        body_style,
    )

    cursor = insert_equation_table_after(
        cursor,
        [
            m_sub("H", "g"), " = I - 2",
            m_frac(
                [m_sub("r", "g"), " ", m_sup(m_sub("r", "g"), "T")],
                [m_sup(m_sub("r", "g"), "T"), " ", m_sub("r", "g")]
            ),
            ",   ",
            "if ",
            ["∥", m_sub("e", "g"), "∥", m_sub("", "2")],
            " > ε and ",
            m_sup(m_sub("r", "g"), "T"),
            " ",
            m_sub("r", "g"),
            " > ε;   otherwise ",
            m_sub("H", "g"),
            " = I",
        ],
        "（4-7b）",
        doc,
    )

    cursor = insert_paragraph_after(
        cursor,
        "在 search_best_group_transform 的逐 group 循环中，代码会为每个 group_idx 单独调用上述自适应构造函数，并将得到的 "
        "householder_matrix 写入 group_candidate_forward 与 group_candidate_backward。由于 Householder 反射同时满足正交与对称性质，"
        "当前实现直接令前向矩阵与逆转置矩阵相同。随后，householder 与其他候选矩阵在同一评分函数下参与比较；若其目标值最小，"
        "则当前 group 的前向矩阵与逆转置矩阵会被追加到 selected_forward 和 selected_backward，并最终按 group 维堆叠为 MixedGroupTransform。"
        "因此，当前项目中的 Householder 实际上是一个面向局部统计、自适应构造、逐 group 选择的低成本正交候选。这种实现更契合第三章所揭示的误差异质性："
        "对于主导方向明显、能量分布不均衡的列块，它能够沿局部能量结构自适应调整反射方向；而对本身较为平滑的列块，则可能自然退化为接近单位阵的温和变换。",
        body_style,
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
