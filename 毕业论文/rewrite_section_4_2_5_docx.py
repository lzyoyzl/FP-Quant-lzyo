from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.2.5 其他候选正交变换实现"
END_HEADING = "4.3 group-wise旋转搜索目标函数设计"


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def build_w_run(text: str):
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def build_m_run(text: str):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def append_expr(container, expr) -> None:
    if isinstance(expr, (list, tuple)):
        for item in expr:
            append_expr(container, item)
    elif isinstance(expr, str):
        container.append(build_m_run(expr))
    else:
        container.append(expr)


def m_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sub")
    append_expr(s, sub)
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    append_expr(e, base)
    s = OxmlElement("m:sup")
    append_expr(s, sup)
    node.append(e)
    node.append(s)
    return node


def m_frac(num, den):
    node = OxmlElement("m:f")
    num_node = OxmlElement("m:num")
    den_node = OxmlElement("m:den")
    append_expr(num_node, num)
    append_expr(den_node, den)
    node.append(num_node)
    node.append(den_node)
    return node


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_table_no_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def insert_paragraph_after(paragraph, text: str, style_name: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.add_run(text)
    return new_para


def insert_equation_table_after(paragraph, equation_parts, number_text: str, doc: Document):
    table = doc.add_table(rows=1, cols=2)
    paragraph._p.addnext(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell, 14.6)
    set_cell_width(right_cell, 2.2)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left_para = left_cell.paragraphs[0]
    clear_paragraph(left_para)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    append_expr(omath, equation_parts)
    omath_para.append(omath)
    left_para._p.append(omath_para)

    right_para = right_cell.paragraphs[0]
    clear_paragraph(right_para)
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_para._p.append(build_w_run(number_text))

    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def main():
    doc = Document(str(DOCX_PATH))

    start_idx = None
    end_idx = None
    start_para = None
    body_style = None

    for idx, para in enumerate(doc.paragraphs):
        if para.text.strip() == START_HEADING:
            start_idx = idx
            start_para = para
        elif para.text.strip() == END_HEADING and start_idx is not None:
            end_idx = idx
            break

    if start_idx is None or end_idx is None or start_para is None:
        raise RuntimeError("Unable to locate section 4.2.5 boundaries.")

    paragraphs = list(doc.paragraphs)
    for para in paragraphs[start_idx + 1:end_idx]:
        if para.style is not None and body_style is None and para.text.strip():
            body_style = para.style.name
        delete_paragraph(para)

    if body_style is None:
        body_style = "Normal"

    cursor = insert_paragraph_after(
        start_para,
        "除 Householder 和 Hadamard 类变换外，默认候选集合中还包括 DCT 与 DST；项目中同时保留了 FastFood 的实现，"
        "但默认搜索空间并未启用该分支。相比 Hadamard 这类以符号翻转和能量扩散为主的结构化变换，"
        "DCT 与 DST 更强调基向量的平滑频率结构，因此更适合处理那些“并非由单一极端峰值主导，而是存在稳定局部相关性”的列块。",
        body_style,
    )

    cursor = insert_paragraph_after(
        cursor,
        "在实现上，DCTTransform 与 DSTransform 都是先构造 group_size 维的 type-II 正交基矩阵，再按 group 重复为 block-diagonal 形式。"
        "对当前实现而言，局部 DCT 基和 DST 基分别由 scipy.fftpack.dct(I_g, type=2, norm='ortho')"
        "与 scipy.fftpack.dst(I_g, type=2, norm='ortho') 直接作用于单位阵得到。其对应的单 group 基矩阵可写为式（4-8a）与式（4-8b）。",
        body_style,
    )

    cursor = insert_equation_table_after(
        cursor,
        [
            "[", m_sub("C", "g"), "]",
            m_sub("p,q", ""),
            " = ",
            m_sub("α", "p"),
            " cos",
            "[",
            m_frac("π", "g"),
            " (q + ",
            m_frac("1", "2"),
            ") p",
            "]",
            ",   ",
            m_sub("α", "0"),
            " = ",
            m_frac("1", "√g"),
            ",   ",
            m_sub("α", "p"),
            " = ",
            m_frac("√2", "√g"),
            " , p > 0",
        ],
        "（4-8a）",
        doc,
    )

    cursor = insert_equation_table_after(
        cursor,
        [
            "[", m_sub("S", "g"), "]",
            m_sub("p,q", ""),
            " = ",
            m_sub("β", "p"),
            " sin",
            "[",
            m_frac("π", "g"),
            " (q + ",
            m_frac("1", "2"),
            ") (p + 1)",
            "]",
            ",   ",
            m_sub("β", "p"),
            " = ",
            m_frac("√2", "√g"),
            " (p < g - 1),   ",
            m_sub("β", "g-1"),
            " = ",
            m_frac("1", "√g"),
        ],
        "（4-8b）",
        doc,
    )

    cursor = insert_equation_table_after(
        cursor,
        [
            m_sub("T", "dct"), " = blkdiag(",
            m_sub("C", "g"), ", ", m_sub("C", "g"), ", … , ", m_sub("C", "g"),
            "),   ",
            m_sub("T", "dst"), " = blkdiag(",
            m_sub("S", "g"), ", ", m_sub("S", "g"), ", … , ", m_sub("S", "g"),
            ")",
        ],
        "（4-8c）",
        doc,
    )

    cursor = insert_paragraph_after(
        cursor,
        "式（4-8a）和式（4-8b）说明，DCT 与 DST 的局部基向量都由平滑变化的频率分量构成，"
        "区别在于前者以余弦基组织局部相关方向，后者以正弦基组织局部相关方向；式（4-8c）则对应当前代码在 forward 阶段"
        "将单 group 基矩阵沿对角线重复，并通过 torch.matmul(x, mat) 作用于整层输入的实现方式。"
        "也就是说，DCT/DST 的实际部署仍然遵循前文给出的 group-wise block-diagonal 结构，只是局部块由显式频率基矩阵填充。",
        body_style,
    )

    cursor = insert_paragraph_after(
        cursor,
        "对第三章中提到的大值主导型或混合型误差而言，DCT 与 DST 的意义在于："
        "它们不一定像 Hadamard 那样强烈地打散峰值，而是更倾向于改变局部相关方向与坐标轴之间的匹配关系。"
        "如果某个 group 的误差并不是来自单点离群值，而是来自多个较大分量在原始基下共同形成的不利结构，"
        "那么这种平滑基变换往往更有可能在不造成过度扩散的前提下改善量化分布。",
        body_style,
    )

    cursor = insert_paragraph_after(
        cursor,
        "FastFood 变换在代码中也给出了结构化随机矩阵实现，但由于其依赖更复杂的随机因子与多重矩阵乘法，"
        "当前默认搜索空间仍然聚焦于 Identity、Hadamard、DCT、DST、GSR 与 Householder 这六类候选。"
        "因此，在本章的实现分析中，DCT 与 DST 主要作为“平滑频率基”这一类候选的代表。",
        body_style,
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
