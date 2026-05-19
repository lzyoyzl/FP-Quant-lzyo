from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.3.2 基于 COV 的目标函数"
END_HEADING = "4.3.3 基于J_tail 的目标函数"


def paragraph_text(p):
    return p.text.strip()


def find_paragraph_index(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if paragraph_text(p) == text:
            return i
    raise ValueError(f"Paragraph not found: {text}")


def set_style(paragraph, style_name):
    try:
        paragraph.style = paragraph.part.document.styles[style_name]
    except Exception:
        paragraph.style = style_name


def insert_paragraph_after(paragraph, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def insert_paragraph_after_table(table, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def move_table_after(table, paragraph):
    tbl = table._tbl
    body = paragraph._p.getparent()
    if tbl.getparent() is not None:
        tbl.getparent().remove(tbl)
    body.insert(body.index(paragraph._p) + 1, tbl)


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def m_run(text):
    mr = m_el("r")
    mt = m_el("t")
    mt.text = text
    mr.append(mt)
    return mr


def m_sub(base, sub):
    node = m_el("sSub")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sub")
    s.append(m_run(sub))
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = m_el("sSup")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sup")
    s.append(m_run(sup))
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = m_el("sSubSup")
    e = m_el("e")
    e.append(m_run(base))
    sub_el = m_el("sub")
    sub_el.append(m_run(sub))
    sup_el = m_el("sup")
    sup_el.append(m_run(sup))
    node.append(e)
    node.append(sub_el)
    node.append(sup_el)
    return node


def m_frac(num_children, den_children):
    node = m_el("f")
    num = m_el("num")
    den = m_el("den")
    for child in num_children:
        num.append(child)
    for child in den_children:
        den.append(child)
    node.append(num)
    node.append(den)
    return node


def append_omath(paragraph, children):
    omp = m_el("oMathPara")
    om = m_el("oMath")
    for child in children:
        om.append(child)
    omp.append(om)
    paragraph._p.append(omp)


def equation_table_after(doc, paragraph, omath_builders, eq_no):
    table = doc.add_table(rows=1, cols=2)
    set_table_no_borders(table)
    set_cell_width(table.cell(0, 0), 13.8)
    set_cell_width(table.cell(0, 1), 2.0)

    left_para = table.cell(0, 0).paragraphs[0]
    right_para = table.cell(0, 1).paragraphs[0]
    set_style(left_para, "Normal")
    set_style(right_para, "Normal")
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for idx, builder in enumerate(omath_builders):
        target = left_para if idx == 0 else table.cell(0, 0).add_paragraph(style="Normal")
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_omath(target, builder())

    right_para.add_run(eq_no)
    move_table_after(table, paragraph)
    return table


def eq_417_line1():
    return [
        m_sub("Δy", "g,ℓ"),
        m_run(" = "),
        m_sub("x", "g"),
        m_subsup("ΔW", "g,ℓ", "⊤"),
    ]


def eq_417_line2():
    return [
        m_run("𝔼‖"),
        m_sub("Δy", "g,ℓ"),
        m_run("‖"),
        m_sub("2", ""),
        m_sup("", "2"),
        m_run(" = Tr("),
        m_sub("ΔW", "g,ℓ"),
        m_sub("Σ", "g,ℓ"),
        m_subsup("ΔW", "g,ℓ", "⊤"),
        m_run(")")
    ]


def eq_417_line3():
    return [
        m_sub("Σ", "g,ℓ"),
        m_run(" = 𝔼["),
        m_subsup("x", "g", "⊤"),
        m_sub("x", "g"),
        m_run("]")
    ]


def eq_418_line1():
    return [
        m_sub("x′", "g"),
        m_run(" = "),
        m_sub("x", "g"),
        m_sub("T", "g"),
        m_run(" , "),
        m_sub("W′", "g,ℓ"),
        m_run(" = "),
        m_sub("W", "g,ℓ"),
        m_subsup("T", "g", "−⊤")
    ]


def eq_418_line2():
    return [
        m_sub("ΔW′", "g,ℓ"),
        m_run(" = "),
        m_sub("Q", "w"),
        m_run("("),
        m_sub("W′", "g,ℓ"),
        m_run(") − "),
        m_sub("W′", "g,ℓ")
    ]


def eq_418_line3():
    return [
        m_sub("Σ′", "g,ℓ"),
        m_run("("),
        m_sub("T", "g"),
        m_run(") = "),
        m_subsup("T", "g", "⊤"),
        m_sub("Σ", "g,ℓ"),
        m_sub("T", "g")
    ]


def eq_418_line4():
    return [
        m_subsup("L", "cov", "(g)"),
        m_run("("),
        m_sub("T", "g"),
        m_run(") = "),
        m_run("Σ"),
        m_subsup("", "ℓ=1", "m"),
        m_frac(
            [
                m_run("Tr("),
                m_sub("ΔW′", "g,ℓ"),
                m_sub("Σ′", "g,ℓ"),
                m_subsup("ΔW′", "g,ℓ", "⊤"),
                m_run(")")
            ],
            [m_sub("d", "out,ℓ")]
        )
    ]


def main():
    doc = Document(str(DOCX_PATH))
    start_idx = find_paragraph_index(doc, START_HEADING)
    end_idx = find_paragraph_index(doc, END_HEADING)

    start_para = doc.paragraphs[start_idx]
    end_para = doc.paragraphs[end_idx]

    body = start_para._p.getparent()
    start_body_idx = body.index(start_para._p)
    end_body_idx = body.index(end_para._p)
    for idx in range(end_body_idx - 1, start_body_idx, -1):
        body.remove(body[idx])

    anchor = start_para
    anchor = insert_paragraph_after(
        anchor,
        "仅以重建 MSE 衡量局部变换的优劣，并不能反映该残差在真实输入分布下对层输出的影响。为此，当前实现提供了基于输入协方差加权的 COV 目标。其基本思想是：如果某一方向在真实输入中出现频繁或能量较高，那么该方向上的量化残差即使幅值不大，也可能对输出误差造成更明显的影响。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "对线性映射 y = xW^⊤ 而言，设第 g 个 group、共享槽位内第 ℓ 个权重张量在量化后的局部残差为 ΔW_{g,ℓ}。若只考察这一局部残差引起的输出扰动，则可先写出局部输出误差，再将其平方期望改写为与输入协方差相关的迹形式，如式（4-17）所示。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_417_line1, eq_417_line2, eq_417_line3], "（4-17）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-17）给出了 COV 目标的数学来源：当输入协方差 Σ_{g,ℓ} 在某些方向上具有更大能量时，同样大小的权重量化残差会产生更高的输出误差代价。因此，COV 路径并不是直接最小化局部权重残差的均方，而是最小化该残差在真实输入分布下诱导的平均输出误差。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "在当前实现中，搜索实际评估的并不是原始权重 W_{g,ℓ} 的误差，而是旋转折叠后的局部权重 W′_{g,ℓ} 的量化误差。由于激活侧同时施加了输入变换，协方差也需要在同一坐标系下同步变换。于是，COV 路径最终比较的是旋转后残差与旋转后协方差之间的加权二次型，其具体形式如式（4-18）所示。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_418_line1, eq_418_line2, eq_418_line3, eq_418_line4], "（4-18）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-18）与代码中的实现一一对应：其中 x′_g = x_gT_g 对应激活侧旋转，W′_{g,ℓ} = W_{g,ℓ}T_g^{-⊤} 对应权重侧逆变换折叠，Σ′_{g,ℓ}(T_g)=T_g^⊤Σ_{g,ℓ}T_g 对应 rotated_covariance = T^T Σ T，而最终的 L_cov^(g)(T_g) 则对应 _compute_gptq_consistent_group_error 中的加权误差计算。与普通 MSE 相比，COV 目标的含义更加明确：它惩罚的不是残差本身，而是残差是否恰好落在高能量输入方向上。因此，当某些方向虽然残差不大、但在真实输入中被频繁激活时，COV 目标会比 MSE 更敏感。",
        style="Normal",
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
