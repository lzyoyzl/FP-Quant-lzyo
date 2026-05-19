from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.3.3 基于J_tail 的目标函数"
END_HEADING = "4.3.4 AUTO 自适应选择策略"


def paragraph_text(p):
    return p.text.strip()


def find_paragraph_index(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if paragraph_text(p) == text:
            return i
    raise ValueError(f"Paragraph not found: {text}")


def set_style(paragraph, style_name):
    try:
        paragraph.style = paragraph.part.document.styles[style_name]
    except Exception:
        paragraph.style = style_name


def insert_paragraph_after(paragraph, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def insert_paragraph_after_table(table, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def move_table_after(table, paragraph):
    tbl = table._tbl
    body = paragraph._p.getparent()
    if tbl.getparent() is not None:
        tbl.getparent().remove(tbl)
    body.insert(body.index(paragraph._p) + 1, tbl)


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def m_run(text):
    mr = m_el("r")
    mt = m_el("t")
    mt.text = text
    mr.append(mt)
    return mr


def m_sub(base, sub):
    node = m_el("sSub")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sub")
    s.append(m_run(sub))
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = m_el("sSup")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sup")
    s.append(m_run(sup))
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = m_el("sSubSup")
    e = m_el("e")
    e.append(m_run(base))
    sub_el = m_el("sub")
    sub_el.append(m_run(sub))
    sup_el = m_el("sup")
    sup_el.append(m_run(sup))
    node.append(e)
    node.append(sub_el)
    node.append(sup_el)
    return node


def m_frac(num_children, den_children):
    node = m_el("f")
    num = m_el("num")
    den = m_el("den")
    for child in num_children:
        num.append(child)
    for child in den_children:
        den.append(child)
    node.append(num)
    node.append(den)
    return node


def append_omath(paragraph, children):
    omp = m_el("oMathPara")
    om = m_el("oMath")
    for child in children:
        om.append(child)
    omp.append(om)
    paragraph._p.append(omp)


def equation_table_after(doc, paragraph, omath_builders, eq_no):
    table = doc.add_table(rows=1, cols=2)
    set_table_no_borders(table)
    set_cell_width(table.cell(0, 0), 13.8)
    set_cell_width(table.cell(0, 1), 2.0)

    left_para = table.cell(0, 0).paragraphs[0]
    right_para = table.cell(0, 1).paragraphs[0]
    set_style(left_para, "Normal")
    set_style(right_para, "Normal")
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for idx, builder in enumerate(omath_builders):
        target = left_para if idx == 0 else table.cell(0, 0).add_paragraph(style="Normal")
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_omath(target, builder())

    right_para.add_run(eq_no)
    move_table_after(table, paragraph)
    return table


def insert_plain_table_after(doc, paragraph, rows, col_widths_cm):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                set_style(p, "Normal")
                if r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_cell_width(cell, col_widths_cm[c_idx])
    move_table_after(table, paragraph)
    return table


def eq_419_line1():
    return [
        m_subsup("L", "base", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") ∈ {"),
        m_subsup("L", "mse", "(g)"), m_run("("), m_sub("T", "g"), m_run("), "),
        m_subsup("L", "cov", "(g)"), m_run("("), m_sub("T", "g"), m_run("), "),
        m_subsup("L", "act", "(g)"), m_run("("), m_sub("T", "g"), m_run(")}")
    ]


def eq_419_line2():
    return [
        m_sub("q", "b"),
        m_run("(R) = Quantile(|R|, "),
        m_frac([m_run("b")], [m_run("B")]),
        m_run("),  b = 0,1,…,B")
    ]


def eq_419_line3():
    return [
        m_sub("ℬ", "b"),
        m_run("(R) = { i : "),
        m_sub("q", "b-1"),
        m_run("(R) ≤ |"),
        m_sub("R", "i"),
        m_run("| < "),
        m_sub("q", "b"),
        m_run("(R) }")
    ]


def eq_420_line1():
    return [
        m_subsup("mse", "b", ""),
        m_run("(Δ,R) = mean{"),
        m_sub("δ", "i"),
        m_run("² : i ∈ "),
        m_sub("ℬ", "b"),
        m_run("(R)}")
    ]


def eq_420_line2():
    return [
        m_subsup("L", "tail", ""),
        m_run("(Δ,R; mode) = "),
        m_frac(
            [
                m_run("Σ"), m_subsup("", "b=1", "B"),
                m_sub("α", "b"),
                m_run("(mode) "),
                m_subsup("mse", "b", ""),
                m_run("(Δ,R)")
            ],
            [
                m_run("Σ"), m_subsup("", "b=1", "B"),
                m_sub("α", "b"),
                m_run("(mode)")
            ]
        )
    ]


def eq_420_line3():
    return [
        m_run("mode ∈ {a_low, b_high, mixed_uniform, mixed_middle, two_tail, auto_abm}")
    ]


def eq_421_line1():
    return [
        m_subsup("L", "jtail", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") = "),
        m_subsup("L", "base", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") + λ "),
        m_subsup("L", "tail", ""),
        m_run("(Δ,R; mode)")
    ]


def eq_421_line2():
    return [
        m_run("(Δ,R) = "),
        m_run("{("),
        m_sub("ΔX", "g"), m_run(","), m_sub("X′", "g"), m_run("), if tail_source = activation; "),
        m_run("("), m_sub("ΔW′", "g,ℓ"), m_run(","), m_sub("W′", "g,ℓ"), m_run("), if tail_source = weight }")
    ]


def main():
    doc = Document(str(DOCX_PATH))
    start_idx = find_paragraph_index(doc, START_HEADING)
    end_idx = find_paragraph_index(doc, END_HEADING)

    start_para = doc.paragraphs[start_idx]
    end_para = doc.paragraphs[end_idx]

    body = start_para._p.getparent()
    start_body_idx = body.index(start_para._p)
    end_body_idx = body.index(end_para._p)
    for idx in range(end_body_idx - 1, start_body_idx, -1):
        body.remove(body[idx])

    anchor = start_para
    anchor = insert_paragraph_after(
        anchor,
        "J_tail 在当前实现中不是独立于基础损失之外的全新目标，而是“基础项 + 尾部项”的复合目标。其核心思想是：在保持基础重建能力的同时，引入一个针对局部分位点误差分布的附加约束，从而让搜索过程不仅考虑整体误差大小，还进一步区分误差集中在哪一段幅值区间上。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "从代码实现看，J_tail 的第一部分是基础目标 L_base^(g)。该基础项并不是固定的，而是由参数 base_loss 指定，可在权重域 MSE、协方差加权的 COV，以及激活域 ACT_MSE 三者之间选择。第二部分是尾部项 L_tail，它以参考张量 |R| 的分位点划分为若干 quantile bins，再对不同区间内的均方误差做加权聚合。J_tail 的组成及其在代码中的对应关系如表4-4所示。",
        style="Normal",
    )
    caption = insert_paragraph_after(anchor, "表4-4 J_tail目标函数的组成与主要参数", style="Caption")
    tbl = insert_plain_table_after(
        doc,
        caption,
        [
            ["组成项", "代码参数/实现", "含义"],
            ["基础项", "base_loss ∈ {mse, cov, act_mse}", "决定以何种统计准则衡量局部基础误差"],
            ["尾部来源", "tail_source ∈ {weight, activation}", "决定尾部项作用在权重侧还是激活侧"],
            ["尾部模式", "a_low, b_high, mixed_uniform, mixed_middle, two_tail, auto_abm", "决定各分位点区间的权重分配方式"],
            ["权重强度", "tail_lambda, tail_weight_power", "分别控制尾部项总强度与分位点权重变化幅度"],
        ],
        [2.8, 5.7, 6.0],
    )
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "为便于说明，记第 g 个 group 在当前候选变换下的基础项为 L_base^(g)(T_g)，其中 T_g 表示当前局部候选矩阵；记参考张量的绝对值分位点为 q_b(R)，相应的第 b 个分位点区间记为 𝔅_b(R)。在当前实现中，基础项的可选集合与分位点区间的定义可写为式（4-19）。其中，R 用于刻画“误差应该围绕哪一份统计对象来划分分位点”，当尾部来源取权重时，R 对应旋转后的局部权重；当尾部来源取激活时，R 对应旋转后的局部激活样本。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_419_line1, eq_419_line2, eq_419_line3], "（4-19）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "在式（4-19）的基础上，先在每个分位点区间内计算局部均方误差，再使用与模式 mode 对应的权重 α_b(mode) 对各区间误差做加权平均，即得到尾部项 L_tail。这里，Δ 表示当前统计链路下的量化残差；δ_i 表示该残差在第 i 个样本位置上的标量分量；B 表示分位点区间总数。由此，尾部项的定义可写为式（4-20）。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_420_line1, eq_420_line2, eq_420_line3], "（4-20）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-20）中的 α_b(mode) 对应 _build_tail_bin_weights 的实现。当前代码支持 a_low、b_high、mixed_uniform、mixed_middle、two_tail 和 auto_abm 六种模式：其中 a_low 对低分位区间赋予更大权重，用于突出小值区域误差；b_high 对高分位区间赋予更大权重，用于突出大值区域误差；mixed_uniform 对各 bin 等权处理；mixed_middle 强调中间分位；two_tail 同时强调两端区间；而 auto_abm 则只作为自适应占位模式，其具体判定逻辑在下一节再进一步展开。因此，J_tail 的尾部项并不是单一固定函数，而是一族由 tail_weight_mode 控制的加权误差函数。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "在此基础上，J_tail 的最终目标函数可写为式（4-21）。式中 λ 表示尾部项的整体权重，决定在总目标中基础项与尾部项的相对重要性；(Δ,R) 则由 tail_source 指定其来源。若 tail_source = weight，则尾部项作用在权重侧；若 tail_source = activation，则尾部项作用在激活侧。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_421_line1, eq_421_line2], "（4-21）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-21）与 search_best_group_transform 中的实现完全一致：当 tail_source = weight 时，尾部项作用在旋转后的局部权重及其量化残差上；当 tail_source = activation 时，尾部项作用在旋转后的激活样本及其量化残差上。需要注意的是，当前代码还对基础项与尾部来源的组合做了约束：若 base_loss = act_mse，则尾部来源必须取 activation；而当 tail_source = activation 时，基础项只能取 cov 或 act_mse。这样做的原因，是为了保证尾部项与基础项作用于同一统计链路，避免出现权重域与激活域混合计量所带来的解释歧义。",
        style="Normal",
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
