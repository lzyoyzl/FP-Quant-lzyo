from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.3.4 AUTO 自适应选择策略"
END_HEADING = "4.4 group-wise旋转搜索算法流程"


def paragraph_text(p):
    return p.text.strip()


def find_paragraph_index(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if paragraph_text(p) == text:
            return i
    raise ValueError(f"Paragraph not found: {text}")


def set_style(paragraph, style_name):
    try:
        paragraph.style = paragraph.part.document.styles[style_name]
    except Exception:
        paragraph.style = style_name


def insert_paragraph_after(paragraph, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def insert_paragraph_after_table(table, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def move_table_after(table, paragraph):
    tbl = table._tbl
    body = paragraph._p.getparent()
    if tbl.getparent() is not None:
        tbl.getparent().remove(tbl)
    body.insert(body.index(paragraph._p) + 1, tbl)


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def m_run(text):
    mr = m_el("r")
    mt = m_el("t")
    mt.text = text
    mr.append(mt)
    return mr


def m_sub(base, sub):
    node = m_el("sSub")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sub")
    s.append(m_run(sub))
    node.append(e)
    node.append(s)
    return node


def m_frac(num_children, den_children):
    node = m_el("f")
    num = m_el("num")
    den = m_el("den")
    for child in num_children:
        num.append(child)
    for child in den_children:
        den.append(child)
    node.append(num)
    node.append(den)
    return node


def append_omath(paragraph, children):
    omp = m_el("oMathPara")
    om = m_el("oMath")
    for child in children:
        om.append(child)
    omp.append(om)
    paragraph._p.append(omp)


def equation_table_after(doc, paragraph, omath_builders, eq_no):
    table = doc.add_table(rows=1, cols=2)
    set_table_no_borders(table)
    set_cell_width(table.cell(0, 0), 13.8)
    set_cell_width(table.cell(0, 1), 2.0)

    left_para = table.cell(0, 0).paragraphs[0]
    right_para = table.cell(0, 1).paragraphs[0]
    set_style(left_para, "Normal")
    set_style(right_para, "Normal")
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for idx, builder in enumerate(omath_builders):
        target = left_para if idx == 0 else table.cell(0, 0).add_paragraph(style="Normal")
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_omath(target, builder())

    right_para.add_run(eq_no)
    move_table_after(table, paragraph)
    return table


def eq_422_line1():
    return [
        m_sub("z", "g"),
        m_run(" = "),
        m_run("{ vec(|"),
        m_sub("X", "g"),
        m_run("|),  if tail_source = activation ; "),
        m_run("vec(|"),
        m_sub("W", "g"),
        m_run("|),  if tail_source = weight }")
    ]


def eq_422_line2():
    return [
        m_sub("q", "50"),
        m_run(" = Quantile("), m_sub("z", "g"), m_run(",0.50) , "),
        m_sub("q", "90"),
        m_run(" = Quantile("), m_sub("z", "g"), m_run(",0.90) , "),
        m_sub("q", "99"),
        m_run(" = Quantile("), m_sub("z", "g"), m_run(",0.99)")
    ]


def eq_422_line3():
    return [
        m_sub("s", "out"), m_run(" = "),
        m_frac([m_sub("q", "99")], [m_run("("), m_sub("q", "90"), m_run(" + ε)")]),
        m_run(" , "),
        m_sub("r", "bulk"), m_run(" = "),
        m_frac([m_sub("q", "50")], [m_run("("), m_sub("q", "99"), m_run(" + ε)")]),
        m_run(" , "),
        m_sub("r", "high"), m_run(" = "),
        m_frac([m_sub("q", "90")], [m_run("("), m_sub("q", "99"), m_run(" + ε)")]),
    ]


def eq_423_line1():
    return [
        m_run("mode"), m_sub("", "g"),
        m_run(" = a_low , if "),
        m_sub("s", "out"), m_run(" ≥ 1.6 ∧ "),
        m_sub("r", "bulk"), m_run(" ≤ 0.30")
    ]


def eq_423_line2():
    return [
        m_run("mode"), m_sub("", "g"),
        m_run(" = b_high , if "),
        m_sub("r", "bulk"), m_run(" ≥ 0.45 ∧ "),
        m_sub("r", "high"), m_run(" ≥ 0.80")
    ]


def eq_423_line3():
    return [
        m_run("mode"), m_sub("", "g"),
        m_run(" = mixed_uniform , otherwise")
    ]


def main():
    doc = Document(str(DOCX_PATH))
    start_idx = find_paragraph_index(doc, START_HEADING)
    end_idx = find_paragraph_index(doc, END_HEADING)

    start_para = doc.paragraphs[start_idx]
    end_para = doc.paragraphs[end_idx]

    body = start_para._p.getparent()
    start_body_idx = body.index(start_para._p)
    end_body_idx = body.index(end_para._p)
    for idx in range(end_body_idx - 1, start_body_idx, -1):
        body.remove(body[idx])

    anchor = start_para
    anchor = insert_paragraph_after(
        anchor,
        "在 J_tail 路径中，AUTO 的核心并不是自动切换基础目标，而是通过 auto_abm 为每个 group 自适应选择尾部权重模式。也就是说，系统不会对整层统一指定 a_low、b_high 或 mixed_uniform，而是先读取当前 group 的局部分布统计，再根据该列块的尾部形态决定应当强调低分位、高分位还是采用较为均衡的尾部加权方式。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "从实现流程看，auto_abm 的判定发生在逐 group 搜索循环内部，并且先于各候选矩阵的打分过程。若 tail_source = activation 且当前 group 已收集到激活样本，则判定统计量取自该 group 的激活样本；否则，判定统计量取自共享槽位权重在当前列块上的局部权重块。为统一表述，记这一步得到的绝对值统计向量为 z_g。随后，代码从 z_g 中提取三个分位点 q50、q90 与 q99，并构造三个无量纲统计量，如式（4-22）所示。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_422_line1, eq_422_line2, eq_422_line3], "（4-22）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-22）中的 q50、q90 与 q99 分别表示 z_g 在 50%、90% 与 99% 位置上的分位点；其中 q99 反映局部峰值上界，q90 反映高值主体水平，q50 则反映中位附近的整体体量。基于这三个分位点，代码进一步定义了三个比值：s_out = q99/(q90+ε) 用于衡量尾部尖锐程度，r_bulk = q50/(q99+ε) 用于衡量中等幅值相对峰值的占比，r_high = q90/(q99+ε) 用于衡量高值主体相对峰值的接近程度。这里 ε 是数值稳定项，用于避免分母过小。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "在得到上述统计量后，auto_abm 采用一个显式的分段规则，将当前 group 映射到具体尾部模式。其判定条件与 `_select_tail_weight_mode_for_group` 中的实现一致，可写为式（4-23）。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_423_line1, eq_423_line2, eq_423_line3], "（4-23）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-23）的含义可以直接对应代码中的三类判断。若 s_out 较大而 r_bulk 较小，说明 q99 明显高于 q90，同时中位水平相对峰值较低，即当前 group 只包含少量尖锐离群值，这时选择 a_low，以便在后续尾部项中强化低分位误差约束；若 r_bulk 与 r_high 都较高，则说明 q50、q90 已经与 q99 较为接近，表明多数分量本身就处于高值区域，这时选择 b_high，以便突出高分位误差；其余情况则退化到 mixed_uniform，表示该 group 不具有极端单峰或整体高值主导特征，尾部项采用较为平衡的加权方式。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "因此，auto_abm 的“自适应”体现在两个层面：第一，它是逐 group 而不是逐层生效，不同列块可以得到不同的尾部模式；第二，它的判断依据来自当前 group 的真实局部分布，而不是预设标签。待 mode_g 被确定后，后续 `_build_tail_bin_weights` 才会据此生成对应的 α_b(mode_g)，并把该模式真正传递到 J_tail 的尾部项计算中。",
        style="Normal",
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
