from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.3 group-wise旋转搜索目标函数设计"
END_HEADING = "4.4 group-wise旋转搜索算法流程"


def paragraph_text(p):
    return p.text.strip()


def find_paragraph_index(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if paragraph_text(p) == text:
            return i
    raise ValueError(f"Paragraph not found: {text}")


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def set_style(paragraph, style_name):
    try:
        paragraph.style = paragraph.part.document.styles[style_name]
    except Exception:
        paragraph.style = style_name


def insert_paragraph_after(paragraph, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def insert_paragraph_after_table(table, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def move_table_after(doc, table, paragraph):
    tbl = table._tbl
    body = paragraph._p.getparent()
    if tbl.getparent() is not None:
        tbl.getparent().remove(tbl)
    body.insert(body.index(paragraph._p) + 1, tbl)


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def w_el(tag):
    return OxmlElement(f"w:{tag}")


def m_run(text):
    mr = m_el("r")
    mt = m_el("t")
    mt.text = text
    mr.append(mt)
    return mr


def m_sub(base, sub):
    node = m_el("sSub")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sub")
    s.append(m_run(sub))
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = m_el("sSup")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sup")
    s.append(m_run(sup))
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = m_el("sSubSup")
    e = m_el("e")
    e.append(m_run(base))
    sub_el = m_el("sub")
    sub_el.append(m_run(sub))
    sup_el = m_el("sup")
    sup_el.append(m_run(sup))
    node.append(e)
    node.append(sub_el)
    node.append(sup_el)
    return node


def m_frac(num_children, den_children):
    node = m_el("f")
    num = m_el("num")
    den = m_el("den")
    for child in num_children:
        num.append(child)
    for child in den_children:
        den.append(child)
    node.append(num)
    node.append(den)
    return node


def m_rad(children):
    node = m_el("rad")
    deg = m_el("deg")
    e = m_el("e")
    node.append(deg)
    for child in children:
        e.append(child)
    node.append(e)
    return node


def append_omath(paragraph, children):
    omp = m_el("oMathPara")
    om = m_el("oMath")
    for child in children:
        om.append(child)
    omp.append(om)
    paragraph._p.append(omp)


def equation_table_after(doc, paragraph, omath_builders, eq_no):
    table = doc.add_table(rows=1, cols=2)
    set_table_no_borders(table)
    set_cell_width(table.cell(0, 0), 13.8)
    set_cell_width(table.cell(0, 1), 2.0)

    left = table.cell(0, 0)
    right = table.cell(0, 1)

    # clear default paras
    left_para = left.paragraphs[0]
    right_para = right.paragraphs[0]
    set_style(left_para, "Normal")
    set_style(right_para, "Normal")
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for idx, builder in enumerate(omath_builders):
        target = left_para if idx == 0 else left.add_paragraph(style="Normal")
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_omath(target, builder())

    right_para.add_run(eq_no)
    move_table_after(doc, table, paragraph)
    return table


def eq_411():
    children = [
        m_sub("W′", "g,ℓ"),
        m_run("("),
        m_sub("T", "g"),
        m_run(") = "),
        m_sub("W", "g,ℓ"),
        m_subsup("T", "g", "−T"),
        m_run(" , "),
        m_sub("ΔW", "g,ℓ"),
        m_run("("),
        m_sub("T", "g"),
        m_run(") = "),
        m_sub("Q", "w"),
        m_run("("),
        m_sub("W′", "g,ℓ"),
        m_run("("),
        m_sub("T", "g"),
        m_run(")) − "),
        m_sub("W′", "g,ℓ"),
        m_run("("),
        m_sub("T", "g"),
        m_run(")")
    ]
    return children


def eq_412():
    children = [
        m_sub("X′", "g"),
        m_run("("),
        m_sub("T", "g"),
        m_run(") = "),
        m_sub("X", "g"),
        m_sub("T", "g"),
        m_run(" , "),
        m_sub("ΔX", "g"),
        m_run("("),
        m_sub("T", "g"),
        m_run(") = "),
        m_sub("Q", "a"),
        m_run("("),
        m_sub("X′", "g"),
        m_run("("),
        m_sub("T", "g"),
        m_run(")) − "),
        m_sub("X′", "g"),
        m_run("("),
        m_sub("T", "g"),
        m_run(")")
    ]
    return children


def eq_413a():
    return [
        m_subsup("L", "mse", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") = "),
        m_run("Σ"),
        m_subsup("", "ℓ=1", "m"),
        m_run(" mean("),
        m_sub("ΔW", "g,ℓ"),
        m_run("("), m_sub("T", "g"), m_run(")"),
        m_run("² )")
    ]


def eq_413b():
    return [
        m_subsup("L", "act", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") = mean("),
        m_sub("ΔX", "g"),
        m_run("("), m_sub("T", "g"), m_run(")² )")
    ]


def eq_414():
    return [
        m_sub("Σ′", "g,ℓ"),
        m_run("("), m_sub("T", "g"), m_run(") = "),
        m_subsup("T", "g", "⊤"),
        m_sub("Σ", "g,ℓ"),
        m_sub("T", "g")
    ]


def eq_415():
    return [
        m_subsup("L", "cov", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") = "),
        m_run("Σ"),
        m_subsup("", "ℓ=1", "m"),
        m_frac(
            [
                m_run("Tr("),
                m_sub("ΔW", "g,ℓ"), m_run("("), m_sub("T", "g"), m_run(") "),
                m_sub("Σ′", "g,ℓ"), m_run("("), m_sub("T", "g"), m_run(") "),
                m_subsup("ΔW", "g,ℓ", "⊤"), m_run("("), m_sub("T", "g"), m_run("))")
            ],
            [m_sub("d", "out,ℓ")]
        )
    ]


def eq_416a():
    return [
        m_subsup("mse", "b", ""),
        m_run("(Δ,R) = mean{"),
        m_sub("δ", "i"),
        m_run("² : |"),
        m_sub("R", "i"),
        m_run("| ∈ "),
        m_sub("B", "b"),
        m_run("(R)}")
    ]


def eq_416b():
    return [
        m_subsup("L", "tail", ""),
        m_run("(Δ,R) = "),
        m_frac(
            [
                m_run("Σ"), m_subsup("", "b=1", "B"),
                m_sub("α", "b"),
                m_subsup("mse", "b", ""),
                m_run("(Δ,R)")
            ],
            [
                m_run("Σ"), m_subsup("", "b=1", "B"),
                m_sub("α", "b")
            ]
        )
    ]


def eq_417():
    return [
        m_subsup("L", "jtail", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") = "),
        m_subsup("L", "base", "(g)"),
        m_run("("), m_sub("T", "g"), m_run(") + "),
        m_run("λ "),
        m_subsup("L", "tail", ""),
        m_run("(Δ,R)")
    ]


def eq_418():
    return [
        m_run("obj_resolved = auto_default_objective , if obj = auto ; obj , otherwise")
    ]


def eq_419():
    return [
        m_sub("s", "out"), m_run(" = "),
        m_frac([m_sub("q", "0.99")], [m_run("("), m_sub("q", "0.90"), m_run(" + ε)")]),
        m_run(" , "),
        m_sub("r", "bulk"), m_run(" = "),
        m_frac([m_sub("q", "0.50")], [m_run("("), m_sub("q", "0.99"), m_run(" + ε)")]),
        m_run(" , "),
        m_sub("r", "high"), m_run(" = "),
        m_frac([m_sub("q", "0.90")], [m_run("("), m_sub("q", "0.99"), m_run(" + ε)")]),
    ]


def eq_420():
    return [
        m_run("mode"), m_sub("", "g"),
        m_run(" = a_low , if "),
        m_sub("s", "out"), m_run(" ≥ 1.6 ∧ "),
        m_sub("r", "bulk"), m_run(" ≤ 0.30 ; "),
        m_run("b_high , if "),
        m_sub("r", "bulk"), m_run(" ≥ 0.45 ∧ "),
        m_sub("r", "high"), m_run(" ≥ 0.80 ; "),
        m_run("mixed_uniform , otherwise")
    ]


def main():
    doc = Document(str(DOCX_PATH))
    start_idx = find_paragraph_index(doc, START_HEADING)
    end_idx = find_paragraph_index(doc, END_HEADING)

    start_para = doc.paragraphs[start_idx]
    end_para = doc.paragraphs[end_idx]

    body = start_para._p.getparent()
    start_body_idx = body.index(start_para._p)
    end_body_idx = body.index(end_para._p)
    for idx in range(end_body_idx - 1, start_body_idx, -1):
        body.remove(body[idx])

    anchor = start_para

    anchor = insert_paragraph_after(
        anchor,
        "当前实现中的 group-wise 旋转搜索，本质上是在每个候选变换下，比较局部量化残差在不同统计准则中的代价。为统一后续表述，先定义第 g 个 group、共享槽位内第 ℓ 个权重张量的权重量化残差与激活量化残差。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_411], "（4-14）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    tbl = equation_table_after(doc, anchor, [eq_412], "（4-15）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")

    anchor = insert_paragraph_after(anchor, "4.3.1 基于 MSE 的目标函数", style="Heading 3")
    anchor = insert_paragraph_after(
        anchor,
        "在当前代码中，MSE 与 ACT_MSE 分别对应权重域和激活域的均方重建误差。其中，MSE 分支对共享同一槽位的全部权重列块逐一计算残差平方均值并求和；ACT_MSE 分支则直接对旋转后的激活样本计算量化重建误差。二者都以局部重建能力为核心，不依赖额外的二阶统计量，因此是实现中最直接的一类评分函数。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_413a, eq_413b], "（4-16）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-16）与代码中的实现完全对应：当 resolved_objective = act_mse 时，系统直接以 mean(delta_x.pow(2)) 作为候选打分；而在权重域 MSE 路径下，每个共享权重张量都会先计算局部残差 ΔW，再将其平方均值累加到 total_error 中。因此，MSE 更关注当前候选矩阵能否降低权重量化残差，而 ACT_MSE 更关注当前候选矩阵能否降低激活量化残差。",
        style="Normal",
    )

    anchor = insert_paragraph_after(anchor, "4.3.2 基于 COV 的目标函数", style="Heading 3")
    anchor = insert_paragraph_after(
        anchor,
        "仅以重建 MSE 衡量局部变换的优劣，并不能反映该残差在真实输入分布下对层输出的影响。为此，当前实现提供了基于输入协方差加权的 COV 目标，并在 GPTQ 路径下作为 auto 的默认解析结果。这一目标首先将当前 group 的输入协方差旋转到候选变换对应的坐标系，再以加权二次型度量量化残差。因此，它本质上是 GPTQ 风格 Hessian 加权思想在局部列块上的实现化版本。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_414], "（4-17）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    tbl = equation_table_after(doc, anchor, [eq_415], "（4-18）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-17）对应代码中 rotated_covariance = T^T Σ T 的实现；式（4-18）则对应 _compute_gptq_consistent_group_error 中的加权误差计算。与普通 MSE 相比，COV 目标不只关注残差本身的大小，还关注该残差是否恰好落在高能量输入方向上。因此，当某些方向虽然残差不大、但在真实输入中出现频繁时，COV 目标会比 MSE 更敏感。",
        style="Normal",
    )

    anchor = insert_paragraph_after(anchor, "4.3.3 基于J_tail 的目标函数", style="Heading 3")
    anchor = insert_paragraph_after(
        anchor,
        "J_tail 在当前实现中不是独立于基础损失之外的全新目标，而是“基础项 + 尾部项”的复合目标。基础项可以取 MSE、COV 或 ACT_MSE；尾部项则根据分位点统计，对大值区域、低值区域或中间区域给予不同权重。因此，J_tail 的核心思想不是替代原有损失，而是在保持基础重建能力的同时，进一步强化对特定误差模式的约束。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_416a, eq_416b], "（4-19）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    tbl = equation_table_after(doc, anchor, [eq_417], "（4-20）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-19）对应 _compute_tail_quantile_error 的实现：代码先依据参考张量 |R| 的分位点构造 quantile bins，再在每个区间内计算局部均方误差，并用 tail_weight_mode 决定的权重做加权平均。式（4-20）则对应 search_best_group_transform 中的复合打分过程。其中，当 base_loss = act_mse 时，尾部来源必须取 activation；当 tail_source = activation 时，代码进一步要求基础项只能取 cov 或 act_mse。这样做的原因，是为了保证尾部项与基础项作用在同一激活统计链路上，避免出现基础损失与尾部损失统计对象不一致的问题。",
        style="Normal",
    )

    anchor = insert_paragraph_after(anchor, "4.3.4 AUTO 自适应选择策略", style="Heading 3")
    anchor = insert_paragraph_after(
        anchor,
        "当前实现中的 AUTO 机制包含两层含义。第一层用于解析基础目标：当 objective = auto 时，系统并不再逐候选显式比较多种损失，而是根据调用路径给出默认结果，其中 RTN 路径默认解析为 MSE，GPTQ 路径默认解析为 COV。第二层则体现在 J_tail 的 auto_abm 策略上，即根据当前 group 的局部分位点统计，自适应选择尾部权重模式。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_418], "（4-21）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    tbl = equation_table_after(doc, anchor, [eq_419], "（4-22）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    tbl = equation_table_after(doc, anchor, [eq_420], "（4-23）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-22）和式（4-23）对应 _select_tail_weight_mode_for_group 的实现逻辑。代码先由 q50、q90 与 q99 构造局部尾部分数，再据此将当前 group 划分为 a_low、b_high 或 mixed_uniform 三类模式，以便在尾部损失中对不同幅值区间施加不同强调。因而，AUTO 的作用并不是替代前述目标函数，而是在既有搜索框架上，为不同量化路径和不同局部误差形态提供更合适的默认选择。",
        style="Normal",
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
