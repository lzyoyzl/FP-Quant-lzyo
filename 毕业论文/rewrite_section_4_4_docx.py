from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL


DOCX_PATH = Path(r"Z:\root\graduation_work\FP-Quant-lzyo\毕业论文\欧阳照林-毕业论文.docx")
START_HEADING = "4.4 group-wise旋转搜索算法流程"
END_HEADING = "4.5 本章小结"


def paragraph_text(p):
    return p.text.strip()


def find_paragraph_index(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if paragraph_text(p) == text:
            return i
    raise ValueError(f"Paragraph not found: {text}")


def set_style(paragraph, style_name):
    try:
        paragraph.style = paragraph.part.document.styles[style_name]
    except Exception:
        paragraph.style = style_name


def insert_paragraph_after(paragraph, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def insert_paragraph_after_table(table, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    set_style(new_para, style)
    if text:
        run = new_para.add_run(text)
        run.font.name = "宋体"
    return new_para


def move_table_after(table, paragraph):
    tbl = table._tbl
    body = paragraph._p.getparent()
    if tbl.getparent() is not None:
        tbl.getparent().remove(tbl)
    body.insert(body.index(paragraph._p) + 1, tbl)


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def insert_plain_table_after(doc, paragraph, rows, col_widths_cm):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                set_style(p, "Normal")
                if r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_cell_width(cell, col_widths_cm[c_idx])
    move_table_after(table, paragraph)
    return table


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def m_run(text):
    mr = m_el("r")
    mt = m_el("t")
    mt.text = text
    mr.append(mt)
    return mr


def m_sub(base, sub):
    node = m_el("sSub")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sub")
    s.append(m_run(sub))
    node.append(e)
    node.append(s)
    return node


def m_sup(base, sup):
    node = m_el("sSup")
    e = m_el("e")
    e.append(m_run(base))
    s = m_el("sup")
    s.append(m_run(sup))
    node.append(e)
    node.append(s)
    return node


def m_subsup(base, sub, sup):
    node = m_el("sSubSup")
    e = m_el("e")
    e.append(m_run(base))
    sub_el = m_el("sub")
    sub_el.append(m_run(sub))
    sup_el = m_el("sup")
    sup_el.append(m_run(sup))
    node.append(e)
    node.append(sub_el)
    node.append(sup_el)
    return node


def append_omath(paragraph, children):
    omp = m_el("oMathPara")
    om = m_el("oMath")
    for child in children:
        om.append(child)
    omp.append(om)
    paragraph._p.append(omp)


def equation_table_after(doc, paragraph, omath_builders, eq_no):
    table = doc.add_table(rows=1, cols=2)
    set_table_no_borders(table)
    set_cell_width(table.cell(0, 0), 13.8)
    set_cell_width(table.cell(0, 1), 2.0)

    left_para = table.cell(0, 0).paragraphs[0]
    right_para = table.cell(0, 1).paragraphs[0]
    set_style(left_para, "Normal")
    set_style(right_para, "Normal")
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for idx, builder in enumerate(omath_builders):
        target = left_para if idx == 0 else table.cell(0, 0).add_paragraph(style="Normal")
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_omath(target, builder())

    right_para.add_run(eq_no)
    move_table_after(table, paragraph)
    return table


def eq_424():
    return [
        m_sub("S", "blk"),
        m_run(" = {"),
        m_sub("Task", "qkv"), m_run(", "),
        m_sub("Task", "o"), m_run(", "),
        m_sub("Task", "gu"), m_run(", "),
        m_sub("Task", "d"), m_run("}")
    ]


def eq_425():
    return [
        m_subsup("T", "s,g", "*"),
        m_run(" = arg min "),
        m_subsup("T", "", ""),
        m_run("∈ "),
        m_sub("C", "s"),
        m_subsup("J", "s", "(g)"),
        m_run("(T)")
    ]


def eq_426():
    return [
        m_sub("F", "s"),
        m_run(" = Stack("),
        m_subsup("F", "s,1", "*"), m_run(", … , "),
        m_subsup("F", "s,G_s", "*"), m_run("), "),
        m_sub("B", "s"),
        m_run(" = Stack("),
        m_subsup("B", "s,1", "*"), m_run(", … , "),
        m_subsup("B", "s,G_s", "*"), m_run(")")
    ]


def eq_427():
    return [
        m_sub("T", "s"),
        m_run(" = MixedGroupTransform("),
        m_sub("F", "s"), m_run(", "),
        m_sub("B", "s"), m_run(", g, "),
        m_sub("N", "s"), m_run(")")
    ]


def main():
    doc = Document(str(DOCX_PATH))
    start_idx = find_paragraph_index(doc, START_HEADING)
    end_idx = find_paragraph_index(doc, END_HEADING)

    start_para = doc.paragraphs[start_idx]
    end_para = doc.paragraphs[end_idx]

    body = start_para._p.getparent()
    start_body_idx = body.index(start_para._p)
    end_body_idx = body.index(end_para._p)
    for idx in range(end_body_idx - 1, start_body_idx, -1):
        body.remove(body[idx])

    anchor = start_para
    anchor = insert_paragraph_after(
        anchor,
        "在完成候选变换集合与搜索目标函数的设计之后，当前项目中的 group-wise 旋转搜索可以概括为一个自顶向下的局部优化流程：先围绕 Transformer block 的四类输入槽位构造搜索任务，再在每个槽位内按 group 列块逐一比较候选矩阵，最后将最优局部变换组装为可部署的 MixedGroupTransform。这样做的目的，是把第四章前面给出的搜索空间设计、目标函数设计与实际代码中的量化导出流程统一起来。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "从任务组织方式看，build_block_input_transforms 并不直接对单个线性层逐一发起搜索，而是先按照功能相关性构造四类槽位任务：q_proj、k_proj、v_proj 共享 qkv 槽位，o_proj 单独对应 o 槽位，gate_proj 与 up_proj 共享 gate_up 槽位，down_proj 对应 down 槽位。记当前 block 的槽位任务集合为式（4-24）所示形式。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_424], "（4-24）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "在得到任务集合之后，算法会先对固定候选矩阵做一次性实例化，并预提取前向矩阵与逆转置矩阵；对 Householder 这类自适应候选，则在逐 group 搜索阶段按局部统计动态构造。随后，算法对每个槽位按输入维切分局部列块，并在每个列块上独立比较候选矩阵。当前列块的最优变换由式（4-25）给出。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_425], "（4-25）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "式（4-25）对应 search_best_group_transform 中的核心循环。具体而言，算法会先切出当前 group 的局部权重块；若目标函数需要激活或协方差统计，还会同步切出对应的激活样本与协方差块。随后，对每个候选矩阵分别计算局部评分，其中评分函数可以是上一节定义的 MSE、COV、ACT_MSE 或 J_tail。所有候选比较结束后，目标值最小的变换即被记为当前列块的最优候选。",
        style="Normal",
    )

    caption = insert_paragraph_after(anchor, "表4-5 group-wise旋转搜索算法的主要步骤与输入输出", style="Caption")
    tbl2 = insert_plain_table_after(
        doc,
        caption,
        [
            ["阶段", "主要输入", "核心操作", "输出"],
            ["槽位任务构造", "block、group_size、槽位划分规则", "围绕 qkv / o / gate_up / down 组织搜索对象", "四类槽位任务"],
            ["候选预处理", "候选变换集合、group_size", "固定候选实例化并提取 forward / backward 矩阵", "可评分候选库"],
            ["逐 group 评分", "局部权重块、激活样本、协方差、目标函数", "对每个候选矩阵计算局部目标值", "当前 group 的最优候选"],
            ["结果组装", "各 group 的最优前向矩阵、逆转置矩阵、候选名称", "堆叠为三维矩阵库并封装", "MixedGroupTransform"],
            ["后续接入", "四个槽位的 MixedGroupTransform", "写回 block 并进入 RTN/GPTQ 量化与导出流程", "可部署量化模型"],
        ],
        [2.5, 4.1, 6.0, 3.4],
    )
    anchor = insert_paragraph_after_table(tbl2, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "当某一槽位的全部 group 搜索结束后，代码不会直接构造整层稠密大矩阵，而是将各列块选中的前向矩阵和逆转置矩阵分别按顺序堆叠，形成式（4-26）所示的前向矩阵库与逆转置矩阵库。其中，上标“*”表示对应列块最终选中的最优候选。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_426], "（4-26）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "在此基础上，系统进一步把矩阵库、group 大小以及各列块选中的候选名称一起封装到 MixedGroupTransform 中，得到当前槽位的最终输入变换，如式（4-27）所示。由于 MixedGroupTransform 在初始化时显式检查前向矩阵与逆转置矩阵之间的线性等价关系，因此后续无论是在激活侧前向旋转，还是在权重侧做逆变换折叠，都能够保持映射一致。",
        style="Normal",
    )
    tbl = equation_table_after(doc, anchor, [eq_427], "（4-27）")
    anchor = insert_paragraph_after_table(tbl, "", style="Normal")
    anchor = insert_paragraph_after(
        anchor,
        "至此，group-wise 旋转搜索算法的主流程可以概括为：先构造四类槽位任务，再在每个槽位内逐 group 搜索最优候选，最后将结果组装成可部署的局部混合变换，并交由后续 RTN 或 GPTQ 路径完成量化、导出与推理接入。换言之，第四章前面介绍的变换粒度、候选空间与目标函数，并不是彼此孤立的模块，而是在这一统一流程中被串联起来，共同构成完整的 group-wise 旋转搜索算法。",
        style="Normal",
    )
    anchor = insert_paragraph_after(
        anchor,
        "【伪代码占位：此处插入 group-wise 旋转搜索算法总体伪代码，展示槽位任务构造、候选矩阵预处理、逐 group 打分、最优变换选择以及 MixedGroupTransform 组装流程。】",
        style="Body Text",
    )
    anchor = insert_paragraph_after(
        anchor,
        "算法5 group-wise旋转搜索算法总体流程（占位）",
        style="Body Text",
    )

    doc.save(str(DOCX_PATH))
    print(f"SAVED:{DOCX_PATH}")


if __name__ == "__main__":
    main()
